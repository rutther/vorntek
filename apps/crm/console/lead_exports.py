from __future__ import annotations

import csv
from datetime import timedelta
from tempfile import SpooledTemporaryFile

from django.conf import settings
from django.db.models import Q, QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from leads.models import LeadSubmission

from .access import active_lead_queryset, lead_attribution_queryset_for_user


MAX_LEAD_EXPORT_ROWS = 50000
MAX_LEAD_EXPORT_BYTES = 64 * 1024 * 1024
LEAD_EXPORT_MEMORY_SPOOL_BYTES = 2 * 1024 * 1024


class LeadExportLimitError(ValueError):
    """The requested lead export exceeds a declared resource boundary."""


class _BoundedUtf8Writer:
    def __init__(self, stream, max_bytes: int):
        self.stream = stream
        self.max_bytes = max_bytes
        self.bytes_written = 0

    def write(self, text: str) -> int:
        encoded = text.encode('utf-8')
        next_size = self.bytes_written + len(encoded)
        if next_size > self.max_bytes:
            raise LeadExportLimitError(
                f'导出文件大小超过 {self.max_bytes // (1024 * 1024) or 1} MB，请缩小筛选范围。'
            )
        self.stream.write(encoded)
        self.bytes_written = next_size
        return len(text)


EXPORT_COLUMNS = [
    ('id', 'ID'),
    ('submitted_at', '提交时间'),
    ('stage', '阶段'),
    ('assignee', '负责人'),
    ('full_name', '姓名'),
    ('company', '公司'),
    ('country', '国家/地区'),
    ('email', '邮箱'),
    ('phone', '电话/WhatsApp'),
    ('product_category', '产品类别'),
    ('capacity', '产能目标'),
    ('message', '项目需求'),
    ('source_channel', '来源渠道'),
    ('source_detail', '来源说明'),
    ('source_url', '提交页面'),
    ('referrer_url', '引荐页面'),
    ('utm_source', 'UTM Source'),
    ('utm_medium', 'UTM Medium'),
    ('utm_campaign', 'UTM Campaign'),
    ('utm_content', 'UTM Content'),
    ('locale', '语言'),
    ('next_follow_up_at', '下次跟进时间'),
    ('follow_up_notes', '跟进备注'),
    ('notification_status', '通知状态'),
    ('duplicate_of_id', '重复线索 ID'),
]


def lead_export_queryset(request, site) -> QuerySet[LeadSubmission]:
    qs = active_lead_queryset(
        lead_attribution_queryset_for_user(
            LeadSubmission.objects.filter(site=site).select_related(
                'form', 'locale', 'assignee', 'team'
            ),
            request.user,
        )
    ).order_by('-submitted_at', '-id')
    stage = request.GET.get('stage', 'all')
    locale = request.GET.get('lead_locale', 'all')
    source = request.GET.get('source_channel', 'all')
    assignee = request.GET.get('assignee', 'all')
    follow_up = request.GET.get('follow_up', 'all')
    query = (request.GET.get('q') or '').strip()
    now = timezone.now()

    if stage != 'all':
        qs = qs.filter(stage=stage)
    if locale == 'shared':
        qs = qs.filter(locale__isnull=True)
    elif locale != 'all':
        qs = qs.filter(locale__locale_code=locale)
    if source != 'all':
        qs = qs.filter(source_channel=source)
    if assignee == 'unassigned':
        qs = qs.filter(assignee__isnull=True)
    elif assignee.isdigit():
        qs = qs.filter(assignee_id=int(assignee))
    if follow_up == 'overdue':
        cutoff = now - timedelta(hours=settings.SITEOS_LEAD_REMINDER_HOURS)
        qs = qs.filter(stage='new', submitted_at__lte=cutoff)
    elif follow_up == 'due':
        qs = qs.filter(stage__in=['new', 'contacted', 'qualified'], next_follow_up_at__lte=now)
    elif follow_up == 'unassigned':
        qs = qs.filter(assignee__isnull=True, stage__in=['new', 'contacted', 'qualified'])
    elif follow_up == 'notification_failed':
        qs = qs.filter(notification_status='failed')
    if query:
        qs = qs.filter(
            Q(full_name__icontains=query)
            | Q(company__icontains=query)
            | Q(email__icontains=query)
            | Q(phone__icontains=query)
            | Q(message__icontains=query)
            | Q(follow_up_notes__icontains=query)
            | Q(source_detail__icontains=query)
            | Q(form__name__icontains=query)
            | Q(form__code__icontains=query)
        )
    return qs


def _safe_text(value) -> str:
    text = '' if value is None else str(value)
    if text.startswith(('=', '+', '-', '@')):
        return f"'{text}"
    return text


def _format_datetime(value) -> str:
    if not value:
        return ''
    return timezone.localtime(value).strftime('%Y-%m-%d %H:%M:%S')


def submission_export_row(submission: LeadSubmission) -> list[str | int]:
    extra = submission.payload_json or {}
    utm = submission.utm_json or {}
    values = {
        'id': submission.id,
        'submitted_at': _format_datetime(submission.submitted_at),
        'stage': submission.stage,
        'assignee': submission.assignee.get_username() if submission.assignee else '',
        'full_name': submission.full_name,
        'company': submission.company,
        'country': submission.country,
        'email': submission.email,
        'phone': submission.phone,
        'product_category': extra.get('product_category', ''),
        'capacity': extra.get('capacity', ''),
        'message': submission.message,
        'source_channel': submission.source_channel,
        'source_detail': submission.source_detail,
        'source_url': submission.source_url,
        'referrer_url': submission.referrer_url,
        'utm_source': utm.get('utm_source', ''),
        'utm_medium': utm.get('utm_medium', ''),
        'utm_campaign': utm.get('utm_campaign', ''),
        'utm_content': utm.get('utm_content', ''),
        'locale': submission.locale.locale_code if submission.locale else '',
        'next_follow_up_at': _format_datetime(submission.next_follow_up_at),
        'follow_up_notes': submission.follow_up_notes,
        'notification_status': submission.notification_status,
        'duplicate_of_id': submission.duplicate_of_id or '',
    }
    return [values[key] if isinstance(values[key], int) else _safe_text(values[key]) for key, _label in EXPORT_COLUMNS]


def _export_count(queryset, max_rows: int) -> int:
    count = queryset.count()
    if count > max_rows:
        raise LeadExportLimitError(f'单次最多导出 {max_rows} 条线索，请缩小筛选范围。')
    return count


def _spooled_export_stream(memory_bytes: int = LEAD_EXPORT_MEMORY_SPOOL_BYTES):
    return SpooledTemporaryFile(max_size=memory_bytes, mode='w+b')


def build_leads_csv(
    queryset: QuerySet[LeadSubmission],
    *,
    max_rows: int = MAX_LEAD_EXPORT_ROWS,
    max_bytes: int = MAX_LEAD_EXPORT_BYTES,
    spool_memory_bytes: int = LEAD_EXPORT_MEMORY_SPOOL_BYTES,
):
    count = _export_count(queryset, max_rows)
    stream = _spooled_export_stream(spool_memory_bytes)
    sink = _BoundedUtf8Writer(stream, max_bytes)
    try:
        sink.write('\ufeff')
        writer = csv.writer(sink, lineterminator='\r\n')
        writer.writerow([label for _key, label in EXPORT_COLUMNS])
        written = 0
        for submission in queryset.iterator(chunk_size=500):
            writer.writerow(submission_export_row(submission))
            written += 1
        if written != count:
            raise LeadExportLimitError('导出期间线索数量发生变化，请刷新后重试。')
        stream.seek(0)
        return stream, written
    except Exception:
        stream.close()
        raise


def build_leads_xlsx(
    queryset: QuerySet[LeadSubmission],
    *,
    max_rows: int = MAX_LEAD_EXPORT_ROWS,
    max_bytes: int = MAX_LEAD_EXPORT_BYTES,
    spool_memory_bytes: int = LEAD_EXPORT_MEMORY_SPOOL_BYTES,
):
    count = _export_count(queryset, max_rows)
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet()
    worksheet.title = 'Leads'
    worksheet.freeze_panes = 'A2'
    headers = [label for _key, label in EXPORT_COLUMNS]
    fill = PatternFill('solid', fgColor='1246D8')
    header_cells = []
    for label in headers:
        cell = WriteOnlyCell(worksheet, value=label)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = fill
        header_cells.append(cell)
    worksheet.append(header_cells)

    written = 0
    for submission in queryset.iterator(chunk_size=500):
        worksheet.append(submission_export_row(submission))
        written += 1
    if written != count:
        workbook.close()
        raise LeadExportLimitError('导出期间线索数量发生变化，请刷新后重试。')

    worksheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{written + 1}'
    widths = [10, 20, 14, 16, 18, 24, 16, 28, 22, 22, 18, 36, 16, 28, 36, 36, 18, 18, 22, 22, 10, 20, 40, 14, 14]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    output = _spooled_export_stream(spool_memory_bytes)
    try:
        workbook.save(output)
        workbook.close()
        if output.tell() > max_bytes:
            raise LeadExportLimitError(
                f'导出文件大小超过 {max_bytes // (1024 * 1024) or 1} MB，请缩小筛选范围。'
            )
        output.seek(0)
        return output, written
    except Exception:
        workbook.close()
        output.close()
        raise
