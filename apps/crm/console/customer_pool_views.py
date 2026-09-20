from __future__ import annotations

import uuid
from io import BytesIO

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db import DatabaseError, connection
from django.db.models import BigIntegerField, Count, DecimalField, F, Func, OrderBy, Q, Value
from django.db.models.functions import Cast, NullIf, Replace
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from django.utils.http import url_has_allowed_host_and_scheme
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from leads.customer_pool_services import (
    archive_reviewed_customer,
    assign_customer,
    claim_customer,
    claim_customers_bulk,
    create_manual_customer,
    publish_reviewed_customer,
    public_pool_queryset,
    release_customer,
    restore_archived_customer,
)
from leads.customer_import_services import (
    MAX_IMPORT_BYTES,
    MAX_RESEARCH_IMPORT_BYTES,
    RESEARCH_SNAPSHOT_PROFILES,
    commit_customer_import,
    preview_customer_import,
    preview_research_snapshot_import,
)
from leads.customer_pool_publish import publish_reviewed_customers_bulk
from leads.customer_standard21_import import STANDARD21_HEADERS, has_hard_restriction
from leads.customer_value_engine import VALUE_CEILING
from leads.models import (
    Company,
    CompanyPoolState,
    CustomerExportJob,
    CustomerImportBatch,
    CustomerImportRow,
    CustomerPoolRow,
    SalesTeam,
    SalesTeamMember,
)

from .access import crm_owned_queryset_for_user, primary_role_label
from .capabilities import DataScope, SalesCapability, can_sales, require_sales, sales_capability_scope
from .customer_pool_exports import (
    EXPORT_FIELD_GROUPS,
    EXPORT_FIELD_LABELS,
    STANDARD21_COLUMN_LABELS,
    STANDARD21_REQUIRED_COLUMNS,
)
from .customer_pool_search import resolve_sensitive_search, store_sensitive_search
from .customer_pool_queries import (
    POOL_STATE_LABELS,
    all_customers_queryset,
    country_condition,
    filter_import_batch,
    filter_pool_row_state,
    filter_pool_state,
    normalize_source_filters,
    parse_country_filter,
    parse_value_bound,
)
from .navigation import build_navigation
from .payloads import PUBLIC_PREVIEW_BASE_URL, default_site_locale


SOURCE_LABELS = {
    'manual': '手动录入',
    'research': '研究开发',
    'website_form': '网站表单',
    'meta_native': 'Meta 原生表单',
    'other': '其他来源',
}
INTAKE_LABELS = {
    'manual_create': '手动创建',
    'file_import': '文件导入',
    'automatic_receive': '自动接收',
}
PAGE_SIZES = (100, 500, 1000)
DEFAULT_PAGE_SIZE = 100
EVIDENCE_OPTIONS = (('pending', '待核验'), ('declared', '已声明'), ('verified', '已核验'), ('rejected', '已拒绝'))
POOL_VIEWS = (('dense', '高密度视图（一行一条联系方式）'), ('company', '企业视图（一行一家企业）'))
# 高密度视图的列（拖拽键, 排序字段, 表头文字）；顺序与表格列序一致。
DENSE_COLUMNS = (
    ('phone', 'phone', '电话'),
    ('email', 'email', '邮箱'),
    ('country', 'country_name', '国家'),
    ('country_code', 'country_code', '国家代码'),
    ('person', 'person_name', '姓名'),
    ('company', 'company_name', '企业'),
    ('value', 'value', '价值'),
    ('email_2', 'email_2', '邮箱2'),
    ('email_3', 'email_3', '邮箱3'),
    ('phone_2', 'phone_2', '电话2'),
    ('phone_3', 'phone_3', '电话3'),
    ('route_type', 'route_type', '路线类型'),
    ('route_tier', 'route_tier', '路线层级'),
    ('account_id', 'account_id', '账户ID'),
    ('whatsapp', 'whatsapp_confirmed', 'WhatsApp'),
    ('evidence_v', 'evidence_v', '证据V'),
    ('identity_i', 'identity_i', '身份I'),
    ('tech_t', 'tech_t', '技术T'),
    ('priority_p', 'priority_p', '优先级P'),
    ('restriction_note', 'restriction_note', '限制注记'),
    ('source_channel', 'source_channel', '来源渠道'),
)
# 电话列按「+ 号之后的数字」数值排序，而不是按字符串。
DENSE_PHONE_SORT_FIELDS = ('phone', 'phone_2', 'phone_3')
DENSE_SORT_DEFAULT_DIRECTION = {'value': 'desc'}
DENSE_SORT_KEYS = {'value', 'file'} | {sort_key for _key, sort_key, _label in DENSE_COLUMNS}
DENSE_SORT_CHOICES = (
    ('value', '价值分高到低（打电话用）'),
    ('file', '原始文件顺序'),
) + tuple((sort_key, label) for _key, sort_key, label in DENSE_COLUMNS if sort_key != 'value')
DENSE_COLUMN_WIDTHS = {
    'phone': 160,
    'email': 210,
    'country': 120,
    'country_code': 96,
    'person': 120,
    'company': 210,
    'value': 76,
    'email_2': 190,
    'email_3': 190,
    'phone_2': 150,
    'phone_3': 150,
    'route_type': 140,
    'route_tier': 140,
    'account_id': 140,
    'whatsapp': 110,
    'evidence_v': 80,
    'identity_i': 80,
    'tech_t': 80,
    'priority_p': 80,
    'restriction_note': 240,
    'source_channel': 150,
}
COMPANY_SORT_ORDERS = {
    'updated': ('-updated_at', '-id'),
    'value': ('-value', 'name', 'id'),
    'name': ('name', 'id'),
    'country': ('country', 'name', 'id'),
}
COMPANY_SORT_CHOICES = (
    ('updated', '最近更新优先'),
    ('value', '价值分高到低'),
    ('name', '企业名称'),
    ('country', '国家 / 企业名称'),
)


class _PhoneDigits(Func):
    """Postgres：剥掉非数字字符后按数字排序；空值交给 NULLS LAST。"""

    arity = 1
    output_field = DecimalField(max_digits=32, decimal_places=0)
    template = "NULLIF(REGEXP_REPLACE(COALESCE(%(expressions)s, ''), '[^0-9]', '', 'g'), '')::numeric"


def _digits_sort_expression(field):
    if connection.vendor == 'postgresql':
        return _PhoneDigits(F(field))
    stripped = F(field)
    for token in ('+', ' ', '-', '(', ')', '.', '/', '#'):
        stripped = Replace(stripped, Value(token), Value(''))
    return Cast(NullIf(stripped, Value('')), output_field=BigIntegerField())


def _dense_sort_order(sort_key, direction):
    if sort_key == 'file':
        return ('company_id', 'row_number', 'row_key')
    descending = direction == 'desc'
    if sort_key in DENSE_PHONE_SORT_FIELDS:
        expression = _digits_sort_expression(sort_key)
    elif sort_key == 'value':
        expression = F('value')
    else:
        expression = NullIf(F(sort_key), Value(''))
    return (OrderBy(expression, descending=descending, nulls_last=True), 'company_id', 'row_number', 'row_key')


def _dense_header_columns(request, *, sort, direction):
    columns = []
    for col_key, sort_key, label in DENSE_COLUMNS:
        active = sort_key == sort
        if active:
            next_direction = 'desc' if direction == 'asc' else 'asc'
        else:
            next_direction = DENSE_SORT_DEFAULT_DIRECTION.get(sort_key, 'asc')
        columns.append({
            'key': col_key,
            'label': label,
            'active': active,
            'aria_sort': ('ascending' if direction == 'asc' else 'descending') if active else '',
            'indicator': ('▲' if direction == 'asc' else '▼') if active else '',
            'sort_url': _query_url(request, sort=sort_key, dir=next_direction, page=None),
        })
    return columns


def _detail_route_columns():
    return [
        {
            'key': key,
            'field': field,
            'label': label,
            'width': DENSE_COLUMN_WIDTHS[key],
        }
        for key, field, label in DENSE_COLUMNS
    ]


def _detail_route_cells(pool_row, *, company_name):
    cells = []
    for key, field, _label in DENSE_COLUMNS:
        value = getattr(pool_row, field, '')
        if field == 'value':
            value = f'{value:.1f}' if value is not None else '—'
        elif field == 'company_name' and not value:
            value = company_name
        else:
            value = value or '—'
        cells.append({'key': key, 'value': value})
    return cells


def _context(request, *, workspace, locale, active_key='customer_pool'):
    return {
        'workspace': workspace,
        'nav_items': build_navigation(active_key, request.user),
        'console_role_label': primary_role_label(request.user),
        'preview_site_url': f'{PUBLIC_PREVIEW_BASE_URL}/?lang={locale.locale_code}',
        'create_actions': [],
    }


def _error_text(error: Exception) -> str:
    if isinstance(error, ValidationError):
        return '；'.join(error.messages)
    return str(error)


def _safe_next(request) -> str:
    candidate = str(request.POST.get('next') or '')
    if candidate and url_has_allowed_host_and_scheme(
        candidate,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return candidate
    return reverse('console:customer_pool')


def _safe_website(value: str) -> str:
    website = str(value or '').strip()
    return website if website.lower().startswith(('https://', 'http://')) else ''


def _single_active_team_for_actor(actor):
    memberships = list(
        SalesTeamMember.objects.select_related('team')
        .filter(user=actor, team__enabled=True)
        .order_by('team_id')[:2]
    )
    return memberships[0].team if len(memberships) == 1 else None


def _pool_queryset(*, request, site, scope: str):
    if scope == 'all':
        return all_customers_queryset(site=site, user=request.user).select_related(
            'pool_state', 'owner_user', 'team',
        ).prefetch_related('customer_sources', 'contact_points')
    if scope == 'mine':
        return (
            Company.objects.filter(site=site, pool_state__state='owned', owner_user=request.user)
            .select_related('pool_state', 'owner_user', 'team')
            .prefetch_related('customer_sources', 'contact_points')
        )
    if scope == 'team':
        return (
            crm_owned_queryset_for_user(
                Company.objects.filter(site=site, pool_state__state='owned'), request.user
            )
            .select_related('pool_state', 'owner_user', 'team')
            .prefetch_related('customer_sources', 'contact_points')
        )
    if scope in {'review', 'archived'}:
        if not can_sales(request.user, SalesCapability.POOL_REVIEW):
            return Company.objects.none()
        return (
            crm_owned_queryset_for_user(
                Company.objects.filter(site=site, pool_state__state=scope), request.user
            )
            .select_related('pool_state', 'owner_user', 'team')
            .prefetch_related('customer_sources', 'contact_points')
        )
    return public_pool_queryset(site=site, actor=request.user).select_related('owner_user', 'team')


def _pool_row_queryset(*, request, site, scope: str):
    """高密度视图：一行 = 一条联系方式（21 列原样行），权限与公司视图同源。"""

    rows = CustomerPoolRow.objects.filter(site=site).select_related(
        'company', 'company__pool_state', 'company__owner_user', 'company__team',
    )
    if scope == 'all':
        return rows.filter(
            company_id__in=all_customers_queryset(site=site, user=request.user).values('pk')
        )
    if scope == 'mine':
        return rows.filter(company__pool_state__state='owned', company__owner_user=request.user)
    if scope == 'team':
        return rows.filter(
            company_id__in=crm_owned_queryset_for_user(
                Company.objects.filter(site=site, pool_state__state='owned'), request.user
            ).values('pk')
        )
    if scope in {'review', 'archived'}:
        if not can_sales(request.user, SalesCapability.POOL_REVIEW):
            return rows.none()
        return rows.filter(
            company_id__in=crm_owned_queryset_for_user(
                Company.objects.filter(site=site, pool_state__state=scope), request.user
            ).values('pk')
        )
    return rows.filter(
        company_id__in=public_pool_queryset(site=site, actor=request.user).values('pk')
    )


def _pool_row_country_groups(rows):
    return (
        (item['country_code'], item['country_name'], item['total'])
        for item in rows.values('country_code', 'country_name')
        .annotate(total=Count('id'))
        .order_by('-total', 'country_code')
    )


def _company_country_groups(companies):
    return (
        (item['country_code'], '', item['total'])
        for item in companies.values('country_code')
        .annotate(total=Count('id'))
        .order_by('-total', 'country_code')
    )


def _country_name_map(site, codes):
    names = {}
    if not codes:
        return names
    for code, name in (
        CustomerPoolRow.objects.filter(site=site, country_code__in=codes)
        .exclude(country_name='')
        .values_list('country_code', 'country_name')
    ):
        names.setdefault(str(code or '').upper(), name)
    return names


def _country_options(groups, *, selected_codes, names=None):
    merged = {}
    for code, name, count in groups:
        token = str(code or '').strip().upper()
        if not token:
            continue
        entry = merged.setdefault(token, {'code': token, 'name': '', 'count': 0})
        entry['count'] += count
        if not entry['name'] and name:
            entry['name'] = str(name)
    options = sorted(merged.values(), key=lambda item: item['code'])
    for option in options:
        if not option['name'] and names:
            option['name'] = names.get(option['code'], '')
        option['label'] = f"{option['code']} {option['name']}".strip()
        option['selected'] = option['code'] in selected_codes
    return options


def _filter_pool_rows(
    rows,
    *,
    query,
    contact_search,
    source_type,
    intake_method,
    country_codes,
    country_terms,
    evidence_status,
    owner_id,
    team_id,
    batch_id,
    value_min=None,
    value_max=None,
):
    if query:
        rows = rows.filter(
            Q(phone__icontains=query)
            | Q(phone_2__icontains=query)
            | Q(phone_3__icontains=query)
            | Q(email__icontains=query)
            | Q(email_2__icontains=query)
            | Q(email_3__icontains=query)
            | Q(person_name__icontains=query)
            | Q(company_name__icontains=query)
            | Q(country_name__icontains=query)
            | Q(account_id__icontains=query)
        )
    if contact_search:
        rows = rows.filter(
            Q(phone__icontains=contact_search)
            | Q(phone_2__icontains=contact_search)
            | Q(phone_3__icontains=contact_search)
            | Q(email__icontains=contact_search)
            | Q(email_2__icontains=contact_search)
            | Q(email_3__icontains=contact_search)
        )
    if source_type in SOURCE_LABELS:
        rows = rows.filter(
            company_id__in=Company.objects.filter(
                customer_sources__source_type=source_type
            ).values('pk')
        )
    if intake_method in INTAKE_LABELS:
        rows = rows.filter(
            company_id__in=Company.objects.filter(
                customer_sources__intake_method=intake_method
            ).values('pk')
        )
    if country_codes or country_terms:
        rows = rows.filter(
            country_condition(
                codes=country_codes, terms=country_terms,
                code_field='country_code', name_field='country_name',
            )
        )
    if evidence_status in {'pending', 'declared', 'verified', 'rejected'}:
        rows = rows.filter(company__pool_state__evidence_status=evidence_status)
    if owner_id:
        rows = rows.filter(company__owner_user_id=owner_id)
    if team_id:
        rows = rows.filter(company__team_id=team_id)
    if batch_id:
        rows = rows.filter(import_batch_id=batch_id)
    if value_min is not None:
        rows = rows.filter(value__gte=value_min)
    if value_max is not None:
        rows = rows.filter(value__lte=value_max)
    return rows


def _review_teams_for_actor(*, actor, site):
    teams = SalesTeam.objects.filter(site=site, enabled=True)
    if sales_capability_scope(actor, SalesCapability.POOL_REVIEW) is DataScope.ALL:
        return teams.order_by('name', 'id')
    return teams.filter(
        memberships__user=actor,
        memberships__membership_role='manager',
    ).distinct().order_by('name', 'id')


def _assignment_candidates_for_company(*, actor, company):
    if (
        company.team_id is None
        or getattr(getattr(company, 'pool_state', None), 'state', None) not in {'available', 'owned'}
        or not can_sales(actor, SalesCapability.ASSIGN, record=company)
    ):
        return get_user_model().objects.none()
    return (
        get_user_model().objects.filter(
            is_active=True,
            sales_team_memberships__team_id=company.team_id,
            sales_team_memberships__team__enabled=True,
        )
        .exclude(pk=company.owner_user_id)
        .distinct()
        .order_by('username', 'id')
    )


def _query_url(request, **updates) -> str:
    params = request.GET.copy()
    for key, value in updates.items():
        if value in ('', None) or (value == 'all' and key != 'scope'):
            params.pop(key, None)
        else:
            params[key] = value
    query = params.urlencode()
    return f'{reverse("console:customer_pool")}{"?" + query if query else ""}'


def _owner_options(companies):
    return (
        companies.exclude(owner_user__isnull=True)
        .values('owner_user_id', 'owner_user__username')
        .distinct()
        .order_by('owner_user__username', 'owner_user_id')
    )


def _team_options(companies):
    return (
        companies.exclude(team__isnull=True)
        .values('team_id', 'team__name')
        .distinct()
        .order_by('team__name', 'team_id')
    )


@login_required
@require_GET
def customer_pool(request):
    require_sales(request.user, SalesCapability.POOL_READ)
    site, locale = default_site_locale(None)
    create_draft = request.session.pop('customer_pool_create_draft', {})
    bulk_claim_result = request.session.pop('customer_pool_bulk_claim_result', None)
    bulk_publish_result = request.session.pop('customer_pool_bulk_publish_result', None)
    global_reader = sales_capability_scope(request.user, SalesCapability.POOL_READ) is DataScope.ALL
    scope = str(request.GET.get('scope') or ('all' if global_reader else 'available'))
    can_review = can_sales(request.user, SalesCapability.POOL_REVIEW)
    if scope not in {'all', 'available', 'mine', 'team', 'review', 'archived'}:
        scope = 'available'
    if scope in {'review', 'archived'} and not can_review:
        scope = 'available'
    query = ' '.join(str(request.GET.get('q') or '').split())[:120]
    contact_search_token = str(request.GET.get('contact_search') or '')[:80]
    contact_search = resolve_sensitive_search(request, contact_search_token)
    source_type = str(request.GET.get('source') or 'all')
    intake_method = str(request.GET.get('intake') or 'all')
    source_type, intake_method = normalize_source_filters(source_type, intake_method)
    country_codes, country_terms = parse_country_filter(request.GET.getlist('country'))
    evidence_status = str(request.GET.get('evidence') or 'all')
    owner_id = int(request.GET.get('owner')) if str(request.GET.get('owner') or '').isdigit() else 0
    team_id = int(request.GET.get('team')) if str(request.GET.get('team') or '').isdigit() else 0
    batch_id = int(request.GET.get('batch')) if str(request.GET.get('batch') or '').isdigit() else 0
    value_min = parse_value_bound(request.GET.get('value_min'))
    value_max = parse_value_bound(request.GET.get('value_max'))
    if value_min is not None and value_max is not None and value_min > value_max:
        value_min, value_max = value_max, value_min
    state_filter = str(request.GET.get('state') or 'all')
    view = str(request.GET.get('view') or 'dense')
    if view not in dict(POOL_VIEWS):
        view = 'dense'
    dense_view = view == 'dense'
    sort_choices = DENSE_SORT_CHOICES if dense_view else COMPANY_SORT_CHOICES
    sort_keys = DENSE_SORT_KEYS if dense_view else set(COMPANY_SORT_ORDERS)
    sort = str(request.GET.get('sort') or sort_choices[0][0])
    if sort not in sort_keys:
        sort = sort_choices[0][0]
    direction = str(request.GET.get('dir') or '')
    if direction not in {'asc', 'desc'}:
        direction = DENSE_SORT_DEFAULT_DIRECTION.get(sort, 'asc') if dense_view else 'asc'
    try:
        claim_team = _single_active_team_for_actor(request.user)
        can_export_capability = can_sales(request.user, SalesCapability.POOL_EXPORT)
        rows = []
        records = []
        page_company_count = 0
        if dense_view:
            queryset = _pool_row_queryset(request=request, site=site, scope=scope)
            scope_count = queryset.count()
            option_companies = Company.objects.filter(
                site=site, pk__in=queryset.values('company_id')
            )
            owner_options = list(_owner_options(option_companies))
            team_options = list(_team_options(option_companies))
            country_options = _country_options(
                _pool_row_country_groups(
                    CustomerPoolRow.objects.filter(
                        site=site, company_id__in=option_companies.values('pk')
                    )
                ),
                selected_codes=country_codes,
            )
            queryset = _filter_pool_rows(
                queryset,
                query=query,
                contact_search=contact_search,
                source_type=source_type,
                intake_method=intake_method,
                country_codes=country_codes,
                country_terms=country_terms,
                evidence_status=evidence_status,
                owner_id=owner_id,
                team_id=team_id,
                batch_id=batch_id,
                value_min=value_min,
                value_max=value_max,
            )
            queryset = filter_pool_row_state(queryset, state_filter).order_by(*_dense_sort_order(sort, direction))
            page_company_count = queryset.values('company_id').distinct().count()
        else:
            queryset = _pool_queryset(request=request, site=site, scope=scope)
            scope_count = queryset.count()
            owner_options = list(_owner_options(queryset))
            team_options = list(_team_options(queryset))
            country_groups = list(_company_country_groups(queryset))
            country_options = _country_options(
                country_groups,
                selected_codes=country_codes,
                names=_country_name_map(
                    site, [code for code, _name, _count in country_groups if code]
                ),
            )
            if query:
                queryset = queryset.filter(
                    Q(name__icontains=query)
                    | Q(website__icontains=query)
                    | Q(industry__icontains=query)
                    | Q(customer_sources__source_detail__icontains=query)
                )
            if contact_search:
                queryset = queryset.filter(contact_points__raw_value__icontains=contact_search)
            if source_type in SOURCE_LABELS:
                queryset = queryset.filter(customer_sources__source_type=source_type)
            if intake_method in INTAKE_LABELS:
                queryset = queryset.filter(customer_sources__intake_method=intake_method)
            if country_codes or country_terms:
                queryset = queryset.filter(
                    country_condition(
                        codes=country_codes, terms=country_terms,
                        code_field='country_code', name_field='country',
                    )
                )
            if evidence_status in {'pending', 'declared', 'verified', 'rejected'}:
                queryset = queryset.filter(pool_state__evidence_status=evidence_status)
            if owner_id:
                queryset = queryset.filter(owner_user_id=owner_id)
            if team_id:
                queryset = queryset.filter(team_id=team_id)
            if batch_id:
                queryset = filter_import_batch(queryset, batch_id)
            if value_min is not None:
                queryset = queryset.filter(value__gte=value_min)
            if value_max is not None:
                queryset = queryset.filter(value__lte=value_max)
            queryset = filter_pool_state(queryset, state_filter).distinct().order_by(*COMPANY_SORT_ORDERS[sort])
        try:
            page_size = int(request.GET.get('page_size') or DEFAULT_PAGE_SIZE)
        except (TypeError, ValueError):
            page_size = DEFAULT_PAGE_SIZE
        if page_size not in PAGE_SIZES:
            page_size = DEFAULT_PAGE_SIZE
        paginator = Paginator(queryset, page_size)
        page = paginator.get_page(request.GET.get('page'))
        page.object_list = list(page.object_list)
        if dense_view:
            seen_companies = set()
            for record in page.object_list:
                company = record.company
                state = getattr(company, 'pool_state', None)
                records.append({
                    'record': record,
                    'company': company,
                    'state': state,
                    'state_label': POOL_STATE_LABELS.get(getattr(state, 'state', 'unmanaged'), '未知状态'),
                    'value_label': f'{record.value:.1f}' if record.value is not None else '—',
                    'restricted': has_hard_restriction(record.restriction_note),
                    'first_of_company': company.pk not in seen_companies,
                    'detail_url': reverse('console:customer_pool_detail', args=[company.pk]),
                    'claim_url': reverse('console:customer_pool_claim', args=[company.pk]),
                    'release_url': reverse('console:customer_pool_release', args=[company.pk]),
                    'claim_token': f'claim-{uuid.uuid4().hex}',
                })
                seen_companies.add(company.pk)
        else:
            batch_ids_by_company = {}
            for company_id, import_batch_id in CustomerImportRow.objects.filter(
                company_id__in=[company.pk for company in page.object_list],
                status__in=('imported', 'unchanged'), batch__status__in=('succeeded', 'partial'),
            ).values_list('company_id', 'batch_id').distinct():
                batch_ids_by_company.setdefault(company_id, set()).add(import_batch_id)
            route_counts = dict(
                CustomerPoolRow.objects.filter(
                    site=site, company_id__in=[company.pk for company in page.object_list]
                )
                .values_list('company_id')
                .annotate(total=Count('id'))
            )
            for company in page.object_list:
                sources = list(company.customer_sources.all())
                contact_points = list(company.contact_points.all())
                first_source = sources[0] if sources else None
                rows.append({
                    'company': company,
                    'state': getattr(company, 'pool_state', None),
                    'value_label': f'{company.value:.1f}' if company.value is not None else '—',
                    'route_count': route_counts.get(company.pk, 0),
                    'state_label': POOL_STATE_LABELS.get(getattr(getattr(company, 'pool_state', None), 'state', 'unmanaged'), '未知状态'),
                    'contact_count': len(contact_points),
                    'evidence_label': dict(EVIDENCE_OPTIONS).get(getattr(getattr(company, 'pool_state', None), 'evidence_status', ''), '证据待补充'),
                    'batch_ids': sorted(batch_ids_by_company.get(company.pk, set()) | {int(source.evidence_json['batch_id']) for source in sources if isinstance(source.evidence_json, dict) and str(source.evidence_json.get('batch_id', '')).isdigit()}),
                    'source_count': len(sources),
                    'source_label': SOURCE_LABELS.get(getattr(first_source, 'source_type', ''), '来源待补充'),
                    'intake_label': INTAKE_LABELS.get(getattr(first_source, 'intake_method', ''), '进入方式待补充'),
                    'source_detail': getattr(first_source, 'source_detail', '') or '—',
                    'contact_point': contact_points[0] if contact_points else None,
                    'detail_url': reverse('console:customer_pool_detail', args=[company.pk]),
                    'claim_url': reverse('console:customer_pool_claim', args=[company.pk]),
                    'release_url': reverse('console:customer_pool_release', args=[company.pk]),
                    'claim_token': f'claim-{uuid.uuid4().hex}',
                })
        export_jobs = []
        has_active_exports = False
        export_status_snapshot = ''
        if can_export_capability:
            now = timezone.now()
            for job in CustomerExportJob.objects.filter(
                site=site,
                created_by=request.user,
            ).order_by('-created_at', '-id')[:10]:
                effective_status = (
                    'expired'
                    if job.status == 'ready' and job.expires_at and job.expires_at <= now
                    else job.status
                )
                has_active_exports = has_active_exports or effective_status in {'pending', 'processing'}
                export_jobs.append({
                    'record': job,
                    'status': effective_status,
                    'field_labels': '、'.join(
                        EXPORT_FIELD_LABELS.get(field, field)
                        for field in (job.fields_json or [])
                    ),
                    'format_label': str(job.format or '').upper(),
                    'download_url': reverse(
                        'console:customer_pool_export_download', args=[job.pk]
                    ),
                    'retry_url': reverse(
                        'console:customer_pool_export_retry', args=[job.pk]
                    ),
                    'cancel_url': reverse(
                        'console:customer_pool_export_cancel', args=[job.pk]
                    ),
                })
            export_status_snapshot = '|'.join(
                f'{item["record"].pk}:{item["status"]}:{item["record"].updated_at.isoformat()}'
                for item in export_jobs
            )
        workspace = {
            'error': False,
            'view': view,
            'dense_view': dense_view,
            'view_options': POOL_VIEWS,
            'view_dense_url': _query_url(request, view='dense', page=None),
            'view_company_url': _query_url(request, view='company', page=None),
            'search_label': '搜索电话 / 邮箱 / 姓名 / 企业' if dense_view else '企业搜索',
            'search_placeholder': '电话、邮箱、姓名、企业、账户号或国家' if dense_view else '企业、行业、网站或来源详情',
            'rows': rows,
            'records': records,
            'page_company_count': page_company_count,
            'page': page,
            'previous_url': _query_url(request, page=page.previous_page_number()) if page.has_previous() else '',
            'next_url': _query_url(request, page=page.next_page_number()) if page.has_next() else '',
            'first_url': _query_url(request, page=1) if page.has_previous() else '',
            'last_url': _query_url(request, page=page.paginator.num_pages) if page.has_next() else '',
            'scope': scope,
            'query': query,
            'contact_search_token': contact_search_token if contact_search else '',
            'contact_search_active': bool(contact_search),
            'clear_contact_search_url': _query_url(request, contact_search=None, page=None),
            'contact_search_url': reverse('console:customer_pool_contact_search'),
            'page_size': page_size,
            'page_size_options': PAGE_SIZES,
            'source_type': source_type,
            'intake_method': intake_method,
            'country_options': country_options,
            'country_selected': country_codes,
            'country_selected_count': len(country_codes),
            'country_filter_active': bool(country_codes or country_terms),
            'country_clear_url': _query_url(request, country=None, page=None),
            'value_min': '' if value_min is None else f'{value_min:g}',
            'value_max': '' if value_max is None else f'{value_max:g}',
            'value_range_max': int(VALUE_CEILING),
            'export_field_groups': [
                {'value': field, 'label': EXPORT_FIELD_LABELS.get(field, field)}
                for field in EXPORT_FIELD_GROUPS
            ],
            'standard21_columns': [
                {
                    'value': header,
                    'label': STANDARD21_COLUMN_LABELS.get(header, header),
                    'required': header in STANDARD21_REQUIRED_COLUMNS,
                }
                for header in STANDARD21_HEADERS
            ],
            'evidence_status': evidence_status,
            'evidence_options': EVIDENCE_OPTIONS,
            'owner_id': owner_id,
            'owner_options': owner_options,
            'team_id': team_id,
            'team_options': team_options,
            'batch_id': batch_id,
            'source_options': SOURCE_LABELS.items(),
            'intake_options': INTAKE_LABELS.items(),
            'clear_url': f'{reverse("console:customer_pool")}?view={view}&scope={scope}',
            'all_url': _query_url(request, scope='all', state=None, page=None),
            'scope_count': scope_count,
            'state_filter': state_filter,
            'state_options': POOL_STATE_LABELS.items(),
            'advanced_active': (
                state_filter != 'all'
                or evidence_status != 'all'
                or bool(owner_id or team_id or batch_id)
                or page_size != DEFAULT_PAGE_SIZE
            ),
            'sort': sort,
            'direction': direction,
            'dense_columns': _dense_header_columns(request, sort=sort, direction=direction) if dense_view else [],
            'page_form_pairs': [
                (key, value)
                for key, values in request.GET.lists()
                if key not in {'page', 'page_size'}
                for value in values
            ],
            'sort_options': sort_choices,
            'sort_value_url': _query_url(request, sort='value', page=None),
            'sort_alt_url': _query_url(request, sort='file' if dense_view else 'updated', page=None),
            'sort_alt_label': '按原始文件顺序' if dense_view else '按最近更新排序',
            'available_url': _query_url(request, scope='available', state=None, page=None),
            'mine_url': _query_url(request, scope='mine', state=None, page=None),
            'team_url': _query_url(request, scope='team', state=None, page=None),
            'review_url': _query_url(request, scope='review', state=None, page=None),
            'archived_url': _query_url(request, scope='archived', state=None, page=None),
            'can_create': can_sales(request.user, SalesCapability.POOL_CREATE),
            'can_claim': claim_team is not None and can_sales(request.user, SalesCapability.POOL_CLAIM),
            'can_bulk_claim': scope == 'available' and claim_team is not None and can_sales(
                request.user, SalesCapability.POOL_CLAIM
            ),
            'can_assign_self': claim_team is not None,
            'can_review': can_review,
            'can_bulk_review': scope == 'review' and can_review,
            'review_team_options': _review_teams_for_actor(actor=request.user, site=site) if can_review else [],
            'bulk_review_url': reverse('console:customer_pool_bulk_review'),
            'bulk_review_token': f'bulk-review-{uuid.uuid4().hex}',
            'bulk_publish_result': bulk_publish_result,
            'can_import': can_sales(request.user, SalesCapability.POOL_IMPORT),
            'can_export': scope != 'available' and can_export_capability and (
                scope != 'all' or sales_capability_scope(request.user, SalesCapability.POOL_EXPORT) is DataScope.ALL
            ),
            'can_view_export_history': can_export_capability,
            'export_jobs': export_jobs,
            'has_active_exports': has_active_exports,
            'export_status_snapshot': export_status_snapshot,
            'export_status_url': reverse('console:customer_pool_export_status'),
            'template_url': reverse('console:customer_pool_template'),
            'import_url': reverse('console:customer_pool_import'),
            'export_url': reverse('console:customer_pool_export'),
            'export_standard21_url': reverse('console:customer_pool_export_standard21'),
            'bulk_claim_url': reverse('console:customer_pool_bulk_claim'),
            'bulk_claim_token': f'bulk-claim-{uuid.uuid4().hex}',
            'bulk_claim_result': bulk_claim_result,
            'manual_create_url': reverse('console:customer_pool_manual_create'),
            'manual_create_token': f'manual-create-{uuid.uuid4().hex}',
            'create_draft': create_draft,
            'reopen_create_modal': bool(create_draft),
        }
    except DatabaseError:
        workspace = {'error': True, 'clear_url': reverse('console:customer_pool')}
        return render(request, 'console/v2/pages/customer_pool.html', _context(request, workspace=workspace, locale=locale), status=503)
    return render(request, 'console/v2/pages/customer_pool.html', _context(request, workspace=workspace, locale=locale))


@login_required
@require_GET
def customer_pool_detail(request, company_id: int):
    require_sales(request.user, SalesCapability.POOL_READ)
    site, locale = default_site_locale(None)
    visible = all_customers_queryset(site=site, user=request.user)
    company = (
        Company.objects.filter(pk__in=visible.values('pk'), pk=company_id)
        .select_related('pool_state', 'owner_user', 'team')
        .prefetch_related('customer_sources', 'contact_points', 'contacts')
        .first()
    )
    if company is None:
        raise Http404('当前数据范围中没有这条客户记录。')
    claim_team = _single_active_team_for_actor(request.user)
    can_review_record = can_sales(request.user, SalesCapability.POOL_REVIEW, record=company)
    assignment_candidates = list(
        _assignment_candidates_for_company(actor=request.user, company=company)
    )
    can_assign_record = bool(assignment_candidates)
    contact_points = list(company.contact_points.all())
    channel_labels = {
        'email': '邮箱',
        'phone': '电话',
        'whatsapp': 'WhatsApp',
        'website': '网站',
        'other': '其他',
    }
    status_labels = {
        'active': '有效',
        'unverified': '待核验',
        'do_not_contact': '禁止联系',
        'invalid': '无效',
    }
    usage_labels = {
        'permitted': '允许使用',
        'restricted': '限制使用',
        'unknown': '用途待确认',
    }
    for point in contact_points:
        point.channel_label = channel_labels.get(point.channel, '未知类型')
        point.status_label = status_labels.get(point.status, '未知状态')
        point.usage_label = usage_labels.get(point.usage_status, '用途待确认')
    routes = [
        {
            'record': pool_row,
            'cells': _detail_route_cells(pool_row, company_name=company.name),
            'restricted': has_hard_restriction(pool_row.restriction_note),
        }
        for pool_row in CustomerPoolRow.objects.filter(site=site, company=company).order_by(
            '-value', 'row_number', 'row_key'
        )
    ]
    workspace = {
        'company': company,
        'state': getattr(company, 'pool_state', None),
        'value_label': f'{company.value:.1f}' if company.value is not None else '—',
        'route_columns': _detail_route_columns(),
        'routes': routes,
        'sources': [
            {
                'record': source,
                'source_label': SOURCE_LABELS.get(source.source_type, source.source_type),
                'intake_label': INTAKE_LABELS.get(source.intake_method, source.intake_method),
                'evidence_label': dict(EVIDENCE_OPTIONS).get(
                    source.evidence_status, '证据待补充'
                ),
            }
            for source in company.customer_sources.all()
        ],
        'contact_points': contact_points,
        'claim_url': reverse('console:customer_pool_claim', args=[company.pk]),
        'release_url': reverse('console:customer_pool_release', args=[company.pk]),
        'claim_token': f'claim-{uuid.uuid4().hex}',
        'can_claim': claim_team is not None and can_sales(request.user, SalesCapability.POOL_CLAIM),
        'can_assign': can_assign_record,
        'assignment_candidates': assignment_candidates,
        'assignment_url': reverse('console:customer_pool_assign', args=[company.pk]),
        'assignment_token': f'assign-{uuid.uuid4().hex}',
        'can_release': can_sales(request.user, SalesCapability.POOL_RELEASE, record=company),
        'can_review': can_review_record,
        'review_teams': _review_teams_for_actor(actor=request.user, site=site) if can_review_record else [],
        'review_url': reverse('console:customer_pool_review', args=[company.pk]),
        'review_token': f'review-{uuid.uuid4().hex}',
        'archive_url': reverse('console:customer_pool_archive', args=[company.pk]),
        'archive_token': f'archive-{uuid.uuid4().hex}',
        'restore_url': reverse('console:customer_pool_restore', args=[company.pk]),
        'restore_token': f'restore-{uuid.uuid4().hex}',
        'release_token': f'release-{uuid.uuid4().hex}',
        'website_url': _safe_website(company.website),
        'company_workspace_url': (
            reverse('console:company_workspace_detail_v2', args=[company.pk])
            if getattr(getattr(company, 'pool_state', None), 'state', None) == 'owned'
            and can_sales(request.user, SalesCapability.READ, record=company)
            else ''
        ),
        'back_url': reverse('console:customer_pool'),
    }
    return render(request, 'console/v2/pages/customer_pool_detail.html', _context(request, workspace=workspace, locale=locale))


@login_required
@require_POST
def customer_pool_manual_create(request):
    site, _locale = default_site_locale(None)
    try:
        result = create_manual_customer(
            site=site,
            actor=request.user,
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
            company_name=str(request.POST.get('company_name') or ''),
            country=str(request.POST.get('country') or ''),
            city=str(request.POST.get('city') or ''),
            industry=str(request.POST.get('industry') or ''),
            website=str(request.POST.get('website') or ''),
            contact_name=str(request.POST.get('contact_name') or ''),
            email=str(request.POST.get('email') or ''),
            phone=str(request.POST.get('phone') or ''),
            notes=str(request.POST.get('notes') or ''),
            source_type=str(request.POST.get('source_type') or 'manual'),
            source_detail=str(request.POST.get('source_detail') or 'CRM 手动创建'),
            assignment_mode=str(request.POST.get('assignment_mode') or 'review'),
        )
    except (ValidationError, PermissionDenied) as error:
        draft_limits = {
            'company_name': 300,
            'country': 120,
            'city': 120,
            'industry': 160,
            'website': 500,
            'contact_name': 200,
            'email': 320,
            'phone': 120,
            'source_type': 32,
            'source_detail': 1000,
            'notes': 4000,
            'assignment_mode': 24,
        }
        request.session['customer_pool_create_draft'] = {
            key: str(request.POST.get(key) or '')[:maximum]
            for key, maximum in draft_limits.items()
        }
        messages.error(request, _error_text(error))
        return redirect('console:customer_pool')
    messages.success(request, f'已创建客户“{result.company.name}”。')
    return redirect('console:customer_pool_detail', company_id=result.company.pk)


@login_required
@require_POST
def customer_pool_contact_search(request):
    require_sales(request.user, SalesCapability.POOL_READ)
    query = ' '.join(str(request.POST.get('contact_query') or '').split())[:120]
    next_url = _safe_next(request)
    if not query:
        messages.error(request, '请输入要查找的电话或邮箱。')
        return redirect(next_url)
    token = store_sensitive_search(request, query)
    separator = '&' if '?' in next_url else '?'
    return redirect(f'{next_url}{separator}contact_search={token}')


@login_required
@require_POST
def customer_pool_claim(request, company_id: int):
    site, _locale = default_site_locale(None)
    try:
        result = claim_customer(
            company_id=company_id,
            site=site,
            actor=request.user,
            expected_version=int(request.POST.get('expected_version') or 0),
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
        )
    except (ValueError, ValidationError, PermissionDenied, Company.DoesNotExist, CompanyPoolState.DoesNotExist) as error:
        messages.error(request, _error_text(error))
    else:
        messages.success(request, f'已领取“{result.company.name}”。')
    return redirect(_safe_next(request))


@login_required
@require_POST
def customer_pool_bulk_claim(request):
    site, _locale = default_site_locale(None)
    try:
        selections = []
        for raw in request.POST.getlist('items'):
            company_id, expected_version = str(raw).split(':', 1)
            selections.append((int(company_id), int(expected_version)))
        result = claim_customers_bulk(
            site=site,
            actor=request.user,
            selections=selections,
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
        )
    except (ValueError, ValidationError, PermissionDenied) as error:
        messages.error(request, _error_text(error))
    else:
        request.session['customer_pool_bulk_claim_result'] = {
            'succeeded': result.succeeded,
            'failed': result.failed,
            'items': [
                {
                    'position': index,
                    'company_id': item.company_id,
                    'status': item.status,
                    'message': item.message,
                    'workflow_code': item.workflow_code,
                }
                for index, item in enumerate(result.items, start=1)
            ],
        }
        if result.failed:
            messages.warning(
                request,
                f'批量领取完成：{result.succeeded} 家成功，{result.failed} 家未领取。',
            )
        else:
            messages.success(request, f'已领取 {result.succeeded} 家客户。')
    return redirect(_safe_next(request))


@login_required
@require_POST
def customer_pool_bulk_review(request):
    site, _locale = default_site_locale(None)
    try:
        company_ids = [int(value) for value in request.POST.getlist('company_ids')]
        team = SalesTeam.objects.filter(
            pk=int(request.POST.get('team_id') or 0),
            site=site,
            enabled=True,
        ).first()
        result = publish_reviewed_customers_bulk(
            site=site,
            actor=request.user,
            team=team,
            company_ids=company_ids,
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
        )
    except (ValueError, ValidationError, PermissionDenied) as error:
        messages.error(request, _error_text(error))
    else:
        request.session['customer_pool_bulk_publish_result'] = {
            'succeeded': result.succeeded,
            'failed': result.failed,
            'items': [
                {
                    'position': index,
                    'company_id': item.company_id,
                    'company_name': item.company_name,
                    'status': item.status,
                    'route_label': item.route_label,
                    'message': item.message,
                }
                for index, item in enumerate(result.items, start=1)
            ],
        }
        if result.failed:
            messages.warning(
                request,
                f'批量核验发布完成：{result.succeeded} 家已发布，{result.failed} 家未发布。',
            )
        else:
            messages.success(request, f'已核验并发布 {result.succeeded} 家客户到公海。')
    return redirect(_safe_next(request))


@login_required
@require_POST
def customer_pool_assign(request, company_id: int):
    site, _locale = default_site_locale(None)
    target_user = get_user_model().objects.filter(
        pk=int(request.POST.get('target_user_id') or 0)
    ).first() if str(request.POST.get('target_user_id') or '').isdigit() else None
    try:
        result = assign_customer(
            company_id=company_id,
            site=site,
            actor=request.user,
            target_user=target_user,
            expected_version=int(request.POST.get('expected_version') or 0),
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
            reason=str(request.POST.get('reason') or ''),
        )
    except (
        ValueError,
        ValidationError,
        PermissionDenied,
        Company.DoesNotExist,
        CompanyPoolState.DoesNotExist,
    ) as error:
        messages.error(request, _error_text(error))
    else:
        action = '分配' if result.replayed is False else '确认'
        messages.success(
            request,
            f'已{action}“{result.company.name}”给 {result.company.owner_user.get_username()}。',
        )
    return redirect('console:customer_pool_detail', company_id=company_id)


@login_required
@require_POST
def customer_pool_review(request, company_id: int):
    site, _locale = default_site_locale(None)
    try:
        team = SalesTeam.objects.get(
            pk=int(request.POST.get('team_id') or 0),
            site=site,
            enabled=True,
        )
        result = publish_reviewed_customer(
            company_id=company_id,
            site=site,
            actor=request.user,
            expected_version=int(request.POST.get('expected_version') or 0),
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
            contact_point_id=int(request.POST.get('contact_point_id') or 0),
            team=team,
        )
    except (
        ValueError,
        ValidationError,
        PermissionDenied,
        Company.DoesNotExist,
        CompanyPoolState.DoesNotExist,
        SalesTeam.DoesNotExist,
    ) as error:
        messages.error(request, _error_text(error))
    else:
        messages.success(request, f'已核验“{result.company.name}”并发布到客户公海。')
    return redirect('console:customer_pool_detail', company_id=company_id)


@login_required
@require_POST
def customer_pool_archive(request, company_id: int):
    site, _locale = default_site_locale(None)
    try:
        result = archive_reviewed_customer(
            company_id=company_id,
            site=site,
            actor=request.user,
            expected_version=int(request.POST.get('expected_version') or 0),
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
            reason=str(request.POST.get('reason') or ''),
        )
    except (
        ValueError,
        ValidationError,
        PermissionDenied,
        Company.DoesNotExist,
        CompanyPoolState.DoesNotExist,
    ) as error:
        messages.error(request, _error_text(error))
    else:
        messages.success(request, f'已归档“{result.company.name}”。')
    return redirect('console:customer_pool_detail', company_id=company_id)


@login_required
@require_POST
def customer_pool_restore(request, company_id: int):
    site, _locale = default_site_locale(None)
    try:
        result = restore_archived_customer(
            company_id=company_id,
            site=site,
            actor=request.user,
            expected_version=int(request.POST.get('expected_version') or 0),
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
            reason=str(request.POST.get('reason') or ''),
        )
    except (
        ValueError,
        ValidationError,
        PermissionDenied,
        Company.DoesNotExist,
        CompanyPoolState.DoesNotExist,
    ) as error:
        messages.error(request, _error_text(error))
    else:
        messages.success(request, f'已将“{result.company.name}”恢复到待核验。')
    return redirect('console:customer_pool_detail', company_id=company_id)


@login_required
@require_POST
def customer_pool_release(request, company_id: int):
    site, _locale = default_site_locale(None)
    try:
        result = release_customer(
            company_id=company_id,
            site=site,
            actor=request.user,
            expected_version=int(request.POST.get('expected_version') or 0),
            idempotency_token=str(request.POST.get('idempotency_token') or ''),
            reason=str(request.POST.get('reason') or ''),
        )
    except (ValueError, ValidationError, PermissionDenied, Company.DoesNotExist, CompanyPoolState.DoesNotExist) as error:
        messages.error(request, _error_text(error))
    else:
        messages.success(request, f'已将“{result.company.name}”释放到客户公海。')
    return redirect(_safe_next(request))


def _customer_template_workbook() -> Workbook:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = '填写说明'
    instructions.append(['Vorntek 客户导入模板', '版本 1'])
    instructions.append(['规则', '“企业”与“联系方式”通过企业外部键关联；手机号、邮箱、证件号按文本填写。'])
    instructions.append(['数据来源', 'manual / research / website_form / meta_native / other'])
    instructions.append(['禁止', '不要填写密码、Token、Cookie 或其他密钥。'])
    companies = workbook.create_sheet('企业')
    company_headers = ['企业外部键*', '企业名称*', '国家/地区', '城市', '行业', '网站', '数据来源*', '来源详情', '备注']
    companies.append(company_headers)
    companies.append(['example-company-001', 'Vorntek 虚构工业客户', '中国', '苏州', '循环水控制系统', 'https://example.invalid', 'research', '合成来源示例，非真实研究结果', '仅为格式示例，禁止联系'])
    contacts = workbook.create_sheet('联系方式')
    contact_headers = ['企业外部键*', '联系人姓名', '职位', '联系方式类型*', '联系方式值*', '分机', '用途', '使用状态', '证据说明']
    contacts.append(contact_headers)
    contacts.append(['example-company-001', '示例联系人', '采购经理', 'email', 'person@example.invalid', '', 'business', 'unknown', '仅为格式示例'])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F5B99')
            cell.alignment = Alignment(vertical='center')
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(14, max(len(str(cell.value or '')) for cell in column) + 2))
    source_validation = DataValidation(type='list', formula1='"manual,research,website_form,meta_native,other"', allow_blank=False)
    companies.add_data_validation(source_validation)
    source_validation.add('G2:G1001')
    channel_validation = DataValidation(type='list', formula1='"email,phone,whatsapp,website,other"', allow_blank=False)
    contacts.add_data_validation(channel_validation)
    channel_validation.add('D2:D3001')
    for sheet in (companies, contacts):
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.number_format = '@'
    return workbook


@login_required
@require_GET
def customer_pool_template(request):
    require_sales(request.user, SalesCapability.POOL_IMPORT)
    output = BytesIO()
    _customer_template_workbook().save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="vorntek-customer-import-template-v1.xlsx"'
    response['X-Content-Type-Options'] = 'nosniff'
    response['Cache-Control'] = 'private, no-store'
    return response


@login_required
@require_http_methods(['GET', 'POST'])
def customer_pool_import(request):
    require_sales(request.user, SalesCapability.POOL_IMPORT)
    site, locale = default_site_locale(None)
    selected_batch = None
    selected_import_profile = str(request.POST.get('import_profile') or 'template')
    if selected_import_profile not in {'template', 'standard21', *RESEARCH_SNAPSHOT_PROFILES}:
        selected_import_profile = 'template'
    if request.method == 'POST':
        try:
            if request.POST.get('action') == 'commit':
                selected_batch = commit_customer_import(
                    batch_id=int(request.POST.get('batch_id') or 0),
                    site=site,
                    actor=request.user,
                )
                messages.success(request, '导入批次已执行；逐行结果已保留。')
            else:
                upload = request.FILES.get('file')
                max_bytes = (
                    MAX_IMPORT_BYTES
                    if selected_import_profile == 'template'
                    else MAX_RESEARCH_IMPORT_BYTES
                )
                if upload and upload.size > max_bytes:
                    raise ValidationError('文件超过所选导入类型的大小限制。')
                file_bytes = upload.read(max_bytes + 1) if upload else b''
                original_name = upload.name if upload else ''
                if selected_import_profile == 'standard21':
                    from leads.customer_standard21_import import preview_standard21_import

                    result = preview_standard21_import(
                        site=site,
                        actor=request.user,
                        file_bytes=file_bytes,
                        original_name=original_name,
                    )
                elif selected_import_profile == 'template':
                    result = preview_customer_import(
                        site=site,
                        actor=request.user,
                        file_bytes=file_bytes,
                        original_name=original_name,
                    )
                else:
                    result = preview_research_snapshot_import(
                        site=site,
                        actor=request.user,
                        file_bytes=file_bytes,
                        original_name=original_name,
                        profile_code=selected_import_profile,
                    )
                selected_batch = result.batch
                messages.info(request, '该文件已存在，显示原预览结果。' if result.replayed else '预览完成，尚未写入客户主数据。')
        except (ValueError, ValidationError, PermissionDenied, CustomerImportBatch.DoesNotExist) as error:
            messages.error(request, _error_text(error))
    elif request.GET.get('batch'):
        selected_batch = CustomerImportBatch.objects.filter(
            pk=request.GET.get('batch'), site=site, created_by=request.user
        ).first()
    count_labels = {
        'source_rows': '判定行',
        'company_rows': '企业源行',
        'contact_rows': '联系方式源行',
        'accounts': '账户数',
        'source_file_rows': '源文件行数',
        'pool_rows': '公海行',
        'new': '新增',
        'supplement': '补充',
        'unchanged': '不变',
        'quarantined': '待人工核验',
        'rejected': '拒绝',
        'imported': '已写入',
        'failed': '失败',
    }
    status_labels = {
        'preview': '待确认',
        'processing': '处理中',
        'partial': '部分完成',
        'succeeded': '已完成',
        'new': '新增',
        'supplement': '补充',
        'unchanged': '不变',
        'quarantined': '待人工核验',
        'rejected': '拒绝',
        'imported': '已写入',
        'failed': '失败',
    }
    batch_queryset = CustomerImportBatch.objects.filter(site=site).select_related('created_by').order_by('-created_at')[:20]
    batches = [
        {
            'id': batch.id,
            'original_name': batch.original_name,
            'created_by': batch.created_by,
            'created_at': batch.created_at,
            'status_label': status_labels.get(batch.status, batch.status),
        }
        for batch in batch_queryset
    ]
    selected_rows = []
    selected_summary = []
    selected_batch_status_label = ''
    if selected_batch:
        selected_rows = [
            {
                'row_number': row.row_number,
                'payload_json': row.payload_json,
                'status': row.status,
                'status_label': status_labels.get(row.status, row.status),
                'message': row.message,
            }
            for row in selected_batch.rows.order_by('row_number', 'id')[:500]
        ]
        selected_summary = [
            {'label': count_labels.get(key, key), 'value': value}
            for key, value in selected_batch.counts_json.items()
        ]
        selected_batch_status_label = status_labels.get(selected_batch.status, selected_batch.status)
    workspace = {
        'batches': batches,
        'selected_batch': selected_batch,
        'selected_batch_status_label': selected_batch_status_label,
        'selected_summary': selected_summary,
        'selected_rows': selected_rows,
        'import_profiles': [
            {'code': profile.code, 'label': profile.label}
            for profile in RESEARCH_SNAPSHOT_PROFILES.values()
        ] + [
            {'code': 'standard21', 'label': '标准 21 列客户池（CSV，10 MB 以内）'},
        ],
        'selected_import_profile': selected_import_profile,
        'template_url': reverse('console:customer_pool_template'),
        'pool_url': reverse('console:customer_pool'),
    }
    return render(request, 'console/v2/pages/customer_pool_import.html', _context(request, workspace=workspace, locale=locale))
