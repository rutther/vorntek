from __future__ import annotations

import csv
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
import shutil
import tempfile
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.db import connection
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from console.customer_pool_exports import _private_path, process_customer_export_job
from leads.customer_pool_services import create_manual_customer, publish_reviewed_customer
from leads.customer_standard21_import import STANDARD21_HEADERS
from leads.models import (
    Activity,
    Company,
    CompanyContactPoint,
    CompanyPoolState,
    Contact,
    CrmMutationReceipt,
    CustomerExportJob,
    CustomerImportBatch,
    CustomerImportRow,
    CustomerPoolRow,
    CustomerSource,
    LeadFormDefinition,
    LeadConversion,
    LeadSubmission,
    Opportunity,
    SalesTeam,
    SalesTeamMember,
    Task,
)
from marketing.models import CanonicalEvent
from sitecore.models import Cta, PageRoute, Site, SiteLocale

from .access import GROUP_BY_ROLE, ROLE_SALES, ROLE_SALES_MANAGER, ROLE_SYSTEM_ADMIN
from .customer_pool_views import DENSE_COLUMNS


def standard21_file(rows: list[list[str]], headers) -> bytes:
    buffer = StringIO()
    writer = csv.writer(buffer, lineterminator='\r\n')
    writer.writerow(list(headers))
    writer.writerows(rows)
    return '\ufeff'.encode('utf-8') + buffer.getvalue().encode('utf-8')


def standard21_row(headers, **overrides) -> list[str]:
    values = {name: '' for name in headers}
    values.update(overrides)
    return [values[name] for name in headers]


class CustomerPoolViewTests(TestCase):
    unmanaged_models = (
        Site,
        SiteLocale,
        PageRoute,
        CanonicalEvent,
        Cta,
        LeadFormDefinition,
        SalesTeam,
        SalesTeamMember,
        LeadSubmission,
        Company,
        Contact,
        Opportunity,
        LeadConversion,
        Activity,
        Task,
        CrmMutationReceipt,
        CompanyPoolState,
        CompanyContactPoint,
        CustomerSource,
        CustomerImportBatch,
        CustomerImportRow,
        CustomerPoolRow,
        CustomerExportJob,
    )

    @classmethod
    def setUpClass(cls):
        existing = set(connection.introspection.table_names())
        cls.created_models = []
        try:
            with connection.schema_editor() as editor:
                for model in cls.unmanaged_models:
                    if model._meta.db_table not in existing:
                        editor.create_model(model)
                        cls.created_models.append(model)
                        existing.add(model._meta.db_table)
            super().setUpClass()
        except Exception:
            with connection.schema_editor() as editor:
                for model in reversed(cls.created_models):
                    editor.delete_model(model)
            raise

    @classmethod
    def tearDownClass(cls):
        try:
            super().tearDownClass()
        finally:
            with connection.schema_editor() as editor:
                for model in reversed(cls.created_models):
                    editor.delete_model(model)

    def setUp(self):
        self.site = Site.objects.create(
            code='siteos_demo',
            name='SiteOS Demo',
            base_url='https://example.invalid',
            default_locale='en',
            enabled=True,
        )
        SiteLocale.objects.create(
            site=self.site,
            locale_code='en',
            label='English',
            is_default=True,
            enabled=True,
        )
        self.team = SalesTeam.objects.create(site=self.site, code='sales', name='Sales')
        sales_group = Group.objects.create(name=GROUP_BY_ROLE[ROLE_SALES])
        manager_group = Group.objects.create(name=GROUP_BY_ROLE[ROLE_SALES_MANAGER])
        admin_group = Group.objects.create(name=GROUP_BY_ROLE[ROLE_SYSTEM_ADMIN])
        User = get_user_model()
        self.sales = User.objects.create_user(username='pool-view-sales', password='test-password')
        self.sales_two = User.objects.create_user(username='pool-view-sales-two', password='test-password')
        self.admin = User.objects.create_user(username='pool-view-admin', password='test-password')
        self.manager = User.objects.create_user(username='pool-view-manager', password='test-password')
        self.sales.groups.add(sales_group)
        self.sales_two.groups.add(sales_group)
        self.manager.groups.add(manager_group)
        self.admin.groups.add(admin_group)
        SalesTeamMember.objects.create(team=self.team, user=self.sales, membership_role='member')
        SalesTeamMember.objects.create(team=self.team, user=self.sales_two, membership_role='member')
        SalesTeamMember.objects.create(team=self.team, user=self.manager, membership_role='manager')
        self.available = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-create-available-0001',
            company_name='Visible Pool Company',
            source_type='research',
            source_detail='Public company website',
            email='visible@example.invalid',
        )
        self.available = publish_reviewed_customer(
            company_id=self.available.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=self.available.pool_state.version,
            idempotency_token='view-publish-available-0001',
            contact_point_id=self.available.company.contact_points.get().pk,
            team=self.team,
        )

    def import_standard21_fixture(self):
        from leads.customer_import_services import commit_customer_import
        from leads.customer_standard21_import import preview_standard21_import

        rows = [
            standard21_row(
                STANDARD21_HEADERS,
                phone='+254700000001',
                country_name='肯尼亚',
                country='KE',
                person_name='Round Trip Person',
                company='Round Trip Ltd',
                value='31.0',
                route_type='公司总机',
                route_tier='企业总机',
                account_id='AF-KE-8001',
                identity_I='I2',
                evidence_V='V3',
                priority_P='P2',
                source_channel='research_public',
            ),
            standard21_row(
                STANDARD21_HEADERS,
                phone='+254700000002',
                email='Second@Example.INVALID',
                country_name='肯尼亚',
                country='KE',
                person_name='未公开',
                company='Round Trip Ltd',
                value='36.0',
                email_2='backup@example.invalid',
                route_type='公司总机',
                route_tier='企业总机',
                account_id='AF-KE-8001',
                identity_I='I2',
                evidence_V='V3',
                priority_P='P2',
                source_channel='research_public',
            ),
            standard21_row(
                STANDARD21_HEADERS,
                phone='+254700000003',
                email='direct@example.invalid',
                country_name='中国',
                country='CN',
                person_name='Direct Buyer',
                company='Round Trip Ltd',
                value='75.0',
                phone_2='+254700000004',
                route_type='直联手机',
                route_tier='本人直联',
                account_id='AF-KE-8002',
                whatsapp_confirmed='已确认',
                identity_I='I3',
                evidence_V='V1',
                priority_P='P1',
                restriction_note='该号码不导出',
                source_channel='research_public',
            ),
        ]
        payload = standard21_file(rows, STANDARD21_HEADERS)
        preview = preview_standard21_import(
            site=self.site,
            actor=self.admin,
            file_bytes=payload,
            original_name='pool21.csv',
        )
        commit_customer_import(batch_id=preview.batch.pk, site=self.site, actor=self.admin)
        return payload

    def test_sales_sees_consistent_pool_page_without_admin_import_controls(self):
        self.client.force_login(self.sales)

        response = self.client.get(reverse('console:customer_pool'), {'view': 'company'})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Visible Pool Company')
        self.assertContains(response, '研究开发')
        self.assertNotContains(response, '导入客户</a>')

    def test_pool_supports_only_the_documented_page_sizes(self):
        self.client.force_login(self.sales)

        response = self.client.get(
            reverse('console:customer_pool'),
            {'view': 'company', 'page_size': '100'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['workspace']['page'].paginator.per_page, 100)
        self.assertContains(response, 'value="100" selected')
        self.assertContains(response, 'value="500"')
        self.assertContains(response, 'value="1000"')

        response = self.client.get(
            reverse('console:customer_pool'),
            {'view': 'company', 'page_size': '5000'},
        )
        self.assertEqual(response.context['workspace']['page'].paginator.per_page, 100)

    def test_pool_combines_evidence_team_and_import_batch_filters(self):
        source = self.available.company.customer_sources.get()
        source.evidence_json = {'batch_id': 42}
        source.save(update_fields=['evidence_json', 'updated_at'])
        self.client.force_login(self.sales)

        response = self.client.get(
            reverse('console:customer_pool'),
            {
                'view': 'company',
                'evidence': 'verified',
                'team': str(self.team.pk),
                'batch': '42',
            },
        )
        self.assertContains(response, 'Visible Pool Company')
        response = self.client.get(
            reverse('console:customer_pool'),
            {'view': 'company', 'batch': '43'},
        )
        self.assertNotContains(response, 'Visible Pool Company')

    def test_contact_search_uses_a_session_bound_opaque_url_token(self):
        self.client.force_login(self.sales)
        contact_value = 'visible@example.invalid'

        response = self.client.post(
            reverse('console:customer_pool_contact_search'),
            {
                'contact_query': contact_value,
                'next': reverse('console:customer_pool') + '?view=company',
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertIn('contact_search=', response['Location'])
        self.assertNotIn(contact_value, response['Location'])
        results = self.client.get(response['Location'])
        self.assertContains(results, 'Visible Pool Company')
        self.assertTrue(results.context['workspace']['contact_search_active'])
        self.assertTrue(results.context['workspace']['clear_contact_search_url'])
        self.assertNotIn(contact_value, results.request['QUERY_STRING'])

        generic_get = self.client.get(
            reverse('console:customer_pool'),
            {'view': 'company', 'q': contact_value},
        )
        self.assertNotContains(generic_get, 'Visible Pool Company')

    def test_admin_can_download_real_three_sheet_template(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('console:customer_pool_template'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertGreater(len(response.content), 4000)

    def test_import_workspace_uses_human_labels_for_counts_and_statuses(self):
        batch = CustomerImportBatch.objects.create(
            site=self.site,
            created_by=self.admin,
            namespace='customer',
            source_type='other',
            adapter_version='customer-xlsx-v1',
            file_sha256='a' * 64,
            original_name='customers.xlsx',
            status='preview',
            counts_json={
                'source_rows': 1,
                'company_rows': 1,
                'contact_rows': 2,
                'new': 1,
            },
        )
        CustomerImportRow.objects.create(
            batch=batch,
            row_key='b' * 64,
            row_number=2,
            payload_json={'company_name': 'Readable Import Company'},
            status='new',
            message='将创建新的企业及来源记录。',
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse('console:customer_pool_import'), {'batch': batch.pk})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '企业源行')
        self.assertContains(response, '联系方式源行')
        self.assertContains(response, '待确认')
        self.assertNotContains(response, '>contact_rows<')

    def test_admin_without_a_sales_team_is_not_offered_an_impossible_claim(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('console:customer_pool'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '当前账号没有唯一销售团队，只能先进入待核验')
        self.assertNotContains(response, '>领取</button>')
        self.assertContains(response, 'name="assignment_mode"')
        self.assertNotContains(response, 'value="self"')

    def test_failed_manual_create_reopens_modal_and_preserves_safe_input_once(self):
        self.client.force_login(self.sales)
        response = self.client.post(
            reverse('console:customer_pool_manual_create'),
            {
                'idempotency_token': 'view-manual-draft-0001',
                'company_name': 'Preserved Draft Company',
                'country': 'Ghana',
                'source_type': 'research',
                'source_detail': 'Public registry draft',
                'assignment_mode': 'review',
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-pool-reopen')
        self.assertContains(response, 'value="Preserved Draft Company"')
        self.assertContains(response, 'value="research" selected')
        self.assertFalse(Company.objects.filter(name='Preserved Draft Company').exists())

        next_page = self.client.get(reverse('console:customer_pool'))
        self.assertNotContains(next_page, 'data-pool-reopen')

    def test_sales_cannot_download_admin_import_template(self):
        self.client.force_login(self.sales)

        response = self.client.get(reverse('console:customer_pool_template'))

        self.assertEqual(response.status_code, 403)

    def test_admin_import_workspace_lists_standard_and_locked_research_profiles(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('console:customer_pool_import'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '标准客户模板')
        self.assertContains(response, '标准 21 列客户池')
        self.assertContains(response, '非洲第二轮 v345（674 家）')
        self.assertContains(response, '中东第二轮 v013（175 家）')
        self.assertContains(response, '仅接受系统锁定的原始版本')

    def test_admin_can_import_commit_and_export_standard21_csv(self):
        values = {header: '' for header in STANDARD21_HEADERS}
        values.update(
            {
                'phone': '+254700000021',
                'country_name': '肯尼亚',
                'country': 'KE',
                'person_name': 'Synthetic Buyer',
                'company': 'Synthetic Standard 21 Ltd',
                'value': '31.0',
                'route_type': '公司总机',
                'route_tier': '企业总机',
                'account_id': 'SYN-KE-0021',
                'evidence_V': 'V3',
                'identity_I': 'I2',
                'priority_P': 'P2',
                'source_channel': 'research_public',
            }
        )
        buffer = StringIO()
        writer = csv.DictWriter(buffer, fieldnames=STANDARD21_HEADERS, lineterminator='\r\n')
        writer.writeheader()
        writer.writerow(values)
        payload = buffer.getvalue().encode('utf-8')
        self.client.force_login(self.admin)

        preview = self.client.post(
            reverse('console:customer_pool_import'),
            {
                'action': 'preview',
                'import_profile': 'standard21',
                'file': SimpleUploadedFile('standard21.csv', payload, content_type='text/csv'),
            },
        )

        self.assertEqual(preview.status_code, 200)
        batch = CustomerImportBatch.objects.get(namespace='standard21')
        self.assertEqual(batch.counts_json['pool_rows'], 1)
        committed = self.client.post(
            reverse('console:customer_pool_import'),
            {'action': 'commit', 'batch_id': str(batch.pk)},
        )
        self.assertEqual(committed.status_code, 200)
        self.assertEqual(CustomerPoolRow.objects.count(), 1)

        exported = self.client.post(
            reverse('console:customer_pool_export_standard21'),
            {'mode': 'current', 'scope': 'review'},
        )
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(exported.content, b'\xef\xbb\xbf' + payload)
        self.assertIn('vorntek-customer-pool-standard21.csv', exported['Content-Disposition'])

    def test_admin_import_workspace_routes_locked_snapshot_to_selected_adapter(self):
        self.client.force_login(self.admin)
        batch = CustomerImportBatch.objects.create(
            site=self.site,
            created_by=self.admin,
            namespace='research_snapshot',
            source_type='research',
            adapter_version='research-africa-v345',
            file_sha256='c' * 64,
            original_name='africa-v345.xlsx',
            status='preview',
            counts_json={'source_rows': 0},
        )
        with patch(
            'console.customer_pool_views.preview_research_snapshot_import',
            return_value=SimpleNamespace(batch=batch, replayed=False),
        ) as preview:
            response = self.client.post(
                reverse('console:customer_pool_import'),
                {
                    'action': 'preview',
                    'import_profile': 'africa_v345',
                    'file': SimpleUploadedFile(
                        'africa-v345.xlsx', b'synthetic-xlsx-bytes'
                    ),
                },
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'value="africa_v345" selected')
        self.assertContains(response, '预览完成，尚未写入客户主数据。')
        self.assertEqual(preview.call_args.kwargs['profile_code'], 'africa_v345')

    def test_manual_create_enters_review_and_admin_can_publish_it_to_team_pool(self):
        self.client.force_login(self.sales)
        response = self.client.post(
            reverse('console:customer_pool_manual_create'),
            {
                'idempotency_token': 'view-manual-review-0001',
                'company_name': 'Review Queue Company',
                'country': 'Kenya',
                'email': 'review-queue@example.invalid',
                'source_type': 'manual',
                'source_detail': 'Sales supplied record',
                'assignment_mode': 'review',
            },
        )
        company = Company.objects.get(name='Review Queue Company')
        self.assertRedirects(
            response,
            reverse('console:customer_pool_detail', args=[company.pk]),
            fetch_redirect_response=False,
        )
        self.assertEqual(company.pool_state.state, 'review')
        self.assertIsNone(company.owner_user_id)

        self.client.force_login(self.admin)
        review_page = self.client.get(
            reverse('console:customer_pool'),
            {'scope': 'review', 'view': 'company'},
        )
        self.assertContains(review_page, 'Review Queue Company')
        detail = self.client.get(reverse('console:customer_pool_detail', args=[company.pk]))
        self.assertContains(detail, '核验并发布')
        response = self.client.post(
            reverse('console:customer_pool_review', args=[company.pk]),
            {
                'expected_version': company.pool_state.version,
                'idempotency_token': 'view-review-publish-0001',
                'team_id': self.team.pk,
                'contact_point_id': company.contact_points.get().pk,
            },
        )
        self.assertRedirects(
            response,
            reverse('console:customer_pool_detail', args=[company.pk]),
            fetch_redirect_response=False,
        )
        company.refresh_from_db()
        self.assertEqual(company.pool_state.state, 'available')
        self.assertEqual(company.team_id, self.team.pk)
        point = company.contact_points.get()
        self.assertEqual(point.status, 'active')
        self.assertEqual(point.usage_status, 'permitted')

    def test_claim_rejects_external_next_url_and_assigns_the_company(self):
        self.client.force_login(self.sales)

        response = self.client.post(
            reverse('console:customer_pool_claim', args=[self.available.company.pk]),
            {
                'expected_version': self.available.pool_state.version,
                'idempotency_token': 'view-claim-customer-0001',
                'next': 'https://evil.example/redirect',
            },
        )

        self.assertRedirects(response, reverse('console:customer_pool'), fetch_redirect_response=False)
        self.available.company.refresh_from_db()
        self.assertEqual(self.available.company.owner_user_id, self.sales.pk)

    def test_manager_can_assign_from_detail_and_sales_cannot_see_assignment_form(self):
        self.client.force_login(self.sales)
        sales_detail = self.client.get(
            reverse('console:customer_pool_detail', args=[self.available.company.pk])
        )
        self.assertNotContains(sales_detail, '分配/交接原因')

        self.client.force_login(self.manager)
        manager_detail = self.client.get(
            reverse('console:customer_pool_detail', args=[self.available.company.pk])
        )
        self.assertContains(manager_detail, '分配客户')
        self.assertContains(manager_detail, self.sales_two.get_username())
        response = self.client.post(
            reverse('console:customer_pool_assign', args=[self.available.company.pk]),
            {
                'target_user_id': self.sales_two.pk,
                'expected_version': self.available.pool_state.version,
                'idempotency_token': 'view-manager-assign-0001',
                'reason': 'Assign this test customer for regional coverage.',
            },
            follow=True,
        )
        self.assertContains(response, f'已分配“{self.available.company.name}”')
        self.assertContains(response, '进入客户工作台')
        self.available.company.refresh_from_db()
        self.assertEqual(self.available.company.owner_user_id, self.sales_two.pk)

    def test_dense_view_defaults_to_value_order_and_exposes_all_standard_columns(self):
        self.import_standard21_fixture()
        self.client.force_login(self.admin)

        response = self.client.get(reverse('console:customer_pool'), {'scope': 'all'})

        self.assertEqual(response.status_code, 200)
        workspace = response.context['workspace']
        columns = {column['key']: column for column in workspace['dense_columns']}
        self.assertEqual(workspace['view'], 'dense')
        self.assertEqual(workspace['sort'], 'value')
        self.assertEqual(workspace['direction'], 'desc')
        self.assertEqual(set(columns), {key for key, _sort, _label in DENSE_COLUMNS})
        self.assertEqual(len(columns), 21)
        self.assertTrue(columns['value']['active'])
        self.assertEqual(columns['value']['aria_sort'], 'descending')
        self.assertIn('sort=value&dir=asc', columns['value']['sort_url'])
        values = [item['record'].value for item in workspace['records']]
        self.assertEqual(values, [Decimal('75.0'), Decimal('36.0'), Decimal('31.0')])

    def test_dense_view_filters_value_and_multiple_country_codes(self):
        self.import_standard21_fixture()
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse('console:customer_pool'),
            {'scope': 'all', 'country': ['ke', 'CN'], 'value_min': '35', 'value_max': '80'},
        )

        self.assertEqual(response.status_code, 200)
        workspace = response.context['workspace']
        records = [item['record'] for item in workspace['records']]
        self.assertEqual([record.phone for record in records], ['+254700000003', '+254700000002'])
        self.assertEqual(workspace['country_selected'], ['KE', 'CN'])
        options = {item['code']: item for item in workspace['country_options']}
        self.assertEqual(set(options), {'CN', 'KE'})
        self.assertTrue(options['CN']['selected'])
        self.assertTrue(options['KE']['selected'])

    def test_standard21_export_selects_matching_companies_and_requested_columns(self):
        self.import_standard21_fixture()
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse('console:customer_pool_export_standard21'),
            {
                'mode': 'current',
                'scope': 'all',
                'country': ['KE'],
                'value_min': '35',
                'value_max': '40',
                'columns': ['phone', 'email', 'value'],
            },
        )

        self.assertEqual(response.status_code, 200)
        lines = response.content.decode('utf-8-sig').splitlines()
        self.assertEqual(lines[0], 'phone,email,value')
        self.assertEqual(len(lines), 3)
        body = response.content.decode('utf-8-sig')
        self.assertIn('+254700000001', body)
        self.assertIn('+254700000002', body)
        self.assertNotIn('+254700000003', body)

    def test_dense_phone_sort_is_numeric_inside_country_code(self):
        self.import_standard21_fixture()
        CustomerPoolRow.objects.filter(account_id='AF-KE-8001', value=Decimal('31.0')).update(
            phone='+254700000009'
        )
        CustomerPoolRow.objects.filter(account_id='AF-KE-8001', value=Decimal('36.0')).update(
            phone='+2547000000010'
        )
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse('console:customer_pool'),
            {'scope': 'all', 'sort': 'phone', 'dir': 'asc'},
        )

        self.assertEqual(response.status_code, 200)
        phones = [item['record'].phone for item in response.context['workspace']['records']]
        self.assertEqual(phones, ['+254700000003', '+254700000009', '+2547000000010'])

    def test_all_scope_does_not_expose_review_rows_to_sales(self):
        self.import_standard21_fixture()

        self.client.force_login(self.admin)
        admin_page = self.client.get(reverse('console:customer_pool'), {'scope': 'all'})
        self.assertContains(admin_page, '+254700000001')
        self.assertContains(admin_page, '+254700000003')

        self.client.force_login(self.sales)
        sales_page = self.client.get(reverse('console:customer_pool'), {'scope': 'all'})
        self.assertEqual(sales_page.status_code, 200)
        self.assertNotContains(sales_page, '+254700000001')
        self.assertNotContains(sales_page, '+254700000003')

    def test_sales_can_bulk_claim_explicit_rows_and_see_per_company_results(self):
        second = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-bulk-create-0001',
            company_name='Second Bulk View Company',
            email='second-bulk@example.invalid',
        )
        second = publish_reviewed_customer(
            company_id=second.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=second.pool_state.version,
            idempotency_token='view-bulk-publish-0001',
            contact_point_id=second.company.contact_points.get().pk,
            team=self.team,
        )
        self.client.force_login(self.sales)
        page = self.client.get(reverse('console:customer_pool'))
        self.assertContains(page, '领取已选')
        response = self.client.post(
            reverse('console:customer_pool_bulk_claim'),
            {
                'idempotency_token': 'view-bulk-claim-0001',
                'items': [
                    f'{self.available.company.pk}:{self.available.pool_state.version}',
                    f'{second.company.pk}:{second.pool_state.version}',
                ],
                'next': reverse('console:customer_pool'),
            },
            follow=True,
        )
        self.assertContains(response, '批量领取结果')
        self.assertContains(response, '2 家成功，0 家未领取')
        self.available.company.refresh_from_db()
        second.company.refresh_from_db()
        self.assertEqual(self.available.company.owner_user_id, self.sales.pk)
        self.assertEqual(second.company.owner_user_id, self.sales.pk)

    def test_admin_can_bulk_review_selected_companies_and_sales_is_denied(self):
        publishable = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-bulk-review-create-0001',
            company_name='Bulk Review Publishable',
            email='bulk-review-publishable@example.invalid',
            assignment_mode='review',
        )
        restricted = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-bulk-review-create-0002',
            company_name='Bulk Review Restricted',
            email='bulk-review-restricted@example.invalid',
            assignment_mode='review',
        )
        restricted_point = restricted.company.contact_points.get()
        restricted_point.usage_status = 'restricted'
        restricted_point.save(update_fields=['usage_status', 'updated_at'])

        self.client.force_login(self.admin)
        review_page = self.client.get(
            reverse('console:customer_pool'),
            {'scope': 'review', 'view': 'company'},
        )
        self.assertContains(review_page, '核验并发布已选')
        self.assertContains(review_page, 'pool-bulk-review-form')
        response = self.client.post(
            reverse('console:customer_pool_bulk_review'),
            {
                'company_ids': [str(publishable.company.pk), str(restricted.company.pk)],
                'team_id': str(self.team.pk),
                'idempotency_token': 'bulk-review-view-0001',
                'next': reverse('console:customer_pool') + '?scope=review',
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '批量核验发布结果')
        self.assertContains(response, '1 家已发布，1 家未发布')
        self.assertContains(
            response,
            '没有可核验的联系方式：源表要求不导出或联系方式已失效，请人工处理。',
        )
        publishable.company.refresh_from_db()
        restricted.company.refresh_from_db()
        self.assertEqual(
            CompanyPoolState.objects.get(company=publishable.company).state,
            'available',
        )
        self.assertEqual(publishable.company.team_id, self.team.pk)
        self.assertEqual(
            CompanyPoolState.objects.get(company=restricted.company).state,
            'review',
        )
        self.assertEqual(
            Activity.objects.filter(
                company=publishable.company,
                subject='核验并发布到客户公海',
            ).count(),
            1,
        )
        self.assertFalse(
            Activity.objects.filter(
                company=restricted.company,
                subject='核验并发布到客户公海',
            ).exists()
        )

        self.client.force_login(self.sales)
        denied = self.client.post(
            reverse('console:customer_pool_bulk_review'),
            {
                'company_ids': [str(restricted.company.pk)],
                'team_id': str(self.team.pk),
                'idempotency_token': 'bulk-review-view-0002',
            },
            follow=True,
        )
        self.assertEqual(denied.status_code, 200)
        self.assertContains(denied, '当前账号没有执行该销售操作或访问该数据范围的权限。')
        restricted.company.refresh_from_db()
        self.assertEqual(
            CompanyPoolState.objects.get(company=restricted.company).state,
            'review',
        )

    def test_admin_can_archive_a_review_record_and_restore_it(self):
        review = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-archive-create-0001',
            company_name='View Archive Company',
            email='view-archive@example.invalid',
        )
        self.client.force_login(self.admin)
        detail = self.client.get(
            reverse('console:customer_pool_detail', args=[review.company.pk])
        )
        self.assertContains(detail, '归档客户')
        response = self.client.post(
            reverse('console:customer_pool_archive', args=[review.company.pk]),
            {
                'expected_version': review.pool_state.version,
                'idempotency_token': 'view-archive-0001',
                'reason': 'Evidence rejected in local verification',
            },
        )
        self.assertEqual(response.status_code, 302)
        review.pool_state.refresh_from_db()
        self.assertEqual(review.pool_state.state, 'archived')
        archived_page = self.client.get(
            reverse('console:customer_pool'), {'scope': 'archived'}
        )
        self.assertContains(archived_page, 'View Archive Company')
        detail = self.client.get(
            reverse('console:customer_pool_detail', args=[review.company.pk])
        )
        self.assertContains(detail, '恢复到待核验')
        response = self.client.post(
            reverse('console:customer_pool_restore', args=[review.company.pk]),
            {
                'expected_version': review.pool_state.version,
                'idempotency_token': 'view-restore-0001',
                'reason': 'New evidence available',
            },
        )
        self.assertEqual(response.status_code, 302)
        review.pool_state.refresh_from_db()
        self.assertEqual(review.pool_state.state, 'review')

    def test_admin_can_archive_an_available_record_from_its_detail_page(self):
        self.client.force_login(self.admin)
        detail = self.client.get(
            reverse('console:customer_pool_detail', args=[self.available.company.pk])
        )
        self.assertContains(detail, '从公海归档')

        response = self.client.post(
            reverse('console:customer_pool_archive', args=[self.available.company.pk]),
            {
                'expected_version': self.available.pool_state.version,
                'idempotency_token': 'view-available-archive-0001',
                'reason': 'Acceptance fixture cleanup',
            },
            follow=True,
        )

        self.assertContains(response, f'已归档“{self.available.company.name}”')
        self.available.pool_state.refresh_from_db()
        self.assertEqual(self.available.pool_state.state, 'archived')
        self.assertEqual(self.available.pool_state.evidence_status, 'verified')

    def test_sales_export_contains_only_the_explicit_owned_company(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-owned-0001',
            company_name='Owned Export Company',
            source_type='manual',
            source_detail='Local test',
            email='owned-export@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)

        response = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )

        self.assertEqual(response.status_code, 302)
        job = CustomerExportJob.objects.get(created_by=self.sales)
        self.assertEqual(job.status, 'pending')
        self.assertEqual(job.storage_path, '')
        self.assertTrue(process_customer_export_job(job.pk))
        job.refresh_from_db()
        download = self.client.get(
            reverse('console:customer_pool_export_download', args=[job.pk])
        )
        self.assertEqual(download.status_code, 200)
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ['企业', '联系方式', '来源明细'])
        enterprise_rows = list(workbook['企业'].iter_rows(values_only=True))
        self.assertEqual(len(enterprise_rows), 2)
        self.assertEqual(enterprise_rows[1][0], owned.company.pk)
        self.assertEqual(enterprise_rows[1][1], 'Owned Export Company')
        self.assertEqual(job.status, 'ready')
        self.assertEqual(job.record_count, 1)
        self.assertEqual(job.scope_json['mode'], 'selected')
        self.assertEqual(job.scope_json['company_ids'], [owned.company.pk])

        self.assertFalse(Path(job.storage_path).is_absolute())
        private_path = _private_path(job.storage_path)
        # A real generated XLSX still downloads after relocating its private
        # root, without rewriting the database record or weakening ownership.
        with tempfile.TemporaryDirectory(prefix='newCrownExportRestore-') as directory:
            restored_root = Path(directory)
            shutil.copyfile(private_path, restored_root/job.storage_path)
            with patch('console.customer_pool_exports._export_root', return_value=restored_root):
                restored_download = self.client.get(
                    reverse('console:customer_pool_export_download', args=[job.pk])
                )
            self.assertEqual(restored_download.status_code, 200)
            self.assertEqual(restored_download.content, download.content)
        self.assertTrue(private_path.is_file())
        cancel = self.client.post(
            reverse('console:customer_pool_export_cancel', args=[job.pk])
        )
        self.assertEqual(cancel.status_code, 302)
        job.refresh_from_db()
        self.assertEqual(job.status, 'canceled')
        self.assertFalse(private_path.exists())

    def test_export_dialog_and_server_honor_explicit_field_selection(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-fields-owned-0001',
            company_name='Field Selection Export Company',
            email='field-selection@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)
        page = self.client.get(reverse('console:customer_pool'), {'scope': 'mine'})
        self.assertContains(page, '确认导出')
        self.assertContains(page, '数据来源和来源明细')
        self.assertContains(page, 'CSV 暂不提供')

        response = self.client.post(
            reverse('console:customer_pool_export'),
            {
                'mode': 'selected',
                'company_ids': [str(owned.company.pk)],
                'fields': ['company', 'source'],
                'format': 'xlsx',
            },
        )
        self.assertEqual(response.status_code, 302)
        job = CustomerExportJob.objects.get(created_by=self.sales)
        self.assertEqual(job.fields_json, ['company', 'source'])
        self.assertEqual(job.format, 'xlsx')
        self.assertTrue(process_customer_export_job(job.pk))
        job.refresh_from_db()
        download = self.client.get(
            reverse('console:customer_pool_export_download', args=[job.pk])
        )
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ['企业', '来源明细'])
        self.assertNotIn('负责人', [cell.value for cell in workbook['企业'][1]])
        self.assertEqual(workbook['来源明细']['C2'].value, 'manual')
        _private_path(job.storage_path).unlink()

    def test_export_rejects_unknown_fields_and_lossy_format(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-invalid-contract-0001',
            company_name='Invalid Export Contract Company',
            email='invalid-contract@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)
        unknown = self.client.post(
            reverse('console:customer_pool_export'),
            {
                'mode': 'selected',
                'company_ids': [str(owned.company.pk)],
                'fields': ['company', 'password'],
                'format': 'xlsx',
            },
            follow=True,
        )
        self.assertContains(unknown, '导出字段无效')
        lossy = self.client.post(
            reverse('console:customer_pool_export'),
            {
                'mode': 'selected',
                'company_ids': [str(owned.company.pk)],
                'fields': ['company'],
                'format': 'csv',
            },
            follow=True,
        )
        self.assertContains(lossy, '当前只支持 XLSX')
        self.assertFalse(CustomerExportJob.objects.filter(created_by=self.sales).exists())

    def test_pending_export_status_is_pollable_and_cancel_wins_before_worker(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-pending-cancel-0001',
            company_name='Pending Export Company',
            email='pending-export@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)
        queued = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )
        self.assertEqual(queued.status_code, 302)
        job = CustomerExportJob.objects.get(created_by=self.sales)
        status = self.client.get(reverse('console:customer_pool_export_status'))
        self.assertEqual(status.status_code, 200)
        self.assertEqual(status.json()['jobs'][0]['status'], 'pending')

        canceled = self.client.post(
            reverse('console:customer_pool_export_cancel', args=[job.pk])
        )
        self.assertEqual(canceled.status_code, 302)
        self.assertFalse(process_customer_export_job(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'canceled')
        self.assertEqual(job.storage_path, '')

        self.client.logout()
        logged_out = self.client.get(reverse('console:customer_pool_export_status'))
        self.assertEqual(logged_out.status_code, 302)
        self.assertIn('/admin/login/', logged_out['Location'])

    def test_cancel_during_generation_cannot_be_overwritten_by_worker_completion(self):
        from console import customer_pool_exports

        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-concurrent-cancel-0001',
            company_name='Concurrent Cancel Export Company',
            email='concurrent-cancel@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)
        self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )
        job = CustomerExportJob.objects.get(created_by=self.sales)
        generated_paths = []
        original_builder = customer_pool_exports._build_private_export

        def cancel_after_build(companies, fields):
            path, digest = original_builder(companies, fields)
            generated_paths.append(path)
            CustomerExportJob.objects.filter(pk=job.pk).update(
                status='canceled',
                completed_at=timezone.now(),
            )
            return path, digest

        with patch(
            'console.customer_pool_exports._build_private_export',
            side_effect=cancel_after_build,
        ):
            self.assertFalse(process_customer_export_job(job.pk))

        job.refresh_from_db()
        self.assertEqual(job.status, 'canceled')
        self.assertTrue(generated_paths)
        self.assertFalse(generated_paths[0].exists())

    def test_management_worker_processes_a_pending_export(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-worker-0001',
            company_name='Worker Export Company',
            email='worker-export@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)
        self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )
        job = CustomerExportJob.objects.get(created_by=self.sales)
        output = StringIO()

        call_command('process_customer_exports', limit=10, stdout=output)

        job.refresh_from_db()
        self.assertEqual(job.status, 'ready')
        self.assertIn('processed 1/1', output.getvalue())
        _private_path(job.storage_path).unlink()

    def test_expired_export_is_revoked_and_failed_export_can_be_retried(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-lifecycle-0001',
            company_name='Export Lifecycle Company',
            source_type='manual',
            source_detail='Local test',
            email='export-lifecycle@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)
        response = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )
        self.assertEqual(response.status_code, 302)
        job = CustomerExportJob.objects.get(created_by=self.sales)
        self.assertTrue(process_customer_export_job(job.pk))
        job.refresh_from_db()
        private_path = _private_path(job.storage_path)
        job.expires_at = timezone.now() - timedelta(seconds=1)
        job.save(update_fields=['expires_at'])
        expired = self.client.get(
            reverse('console:customer_pool_export_download', args=[job.pk])
        )
        self.assertEqual(expired.status_code, 410)
        job.refresh_from_db()
        self.assertEqual(job.status, 'expired')
        self.assertFalse(private_path.exists())

        failed = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )
        self.assertEqual(failed.status_code, 302)
        queued_failure = CustomerExportJob.objects.filter(
            created_by=self.sales,
            status='pending',
        ).get()
        with patch(
            'console.customer_pool_exports._workbook',
            side_effect=RuntimeError('synthetic export failure'),
        ):
            self.assertFalse(process_customer_export_job(queued_failure.pk))
        failed_job = CustomerExportJob.objects.filter(
            created_by=self.sales,
            status='failed',
        ).get()
        retried = self.client.post(
            reverse('console:customer_pool_export_retry', args=[failed_job.pk]),
            follow=True,
        )
        self.assertContains(retried, '重新进入后台队列')
        failed_job.refresh_from_db()
        self.assertEqual(failed_job.status, 'pending')
        self.assertTrue(process_customer_export_job(failed_job.pk))
        failed_job.refresh_from_db()
        self.assertEqual(failed_job.status, 'ready')
        self.assertTrue(_private_path(failed_job.storage_path).is_file())
        _private_path(failed_job.storage_path).unlink()

    def test_regular_export_excludes_restricted_contact_routes(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-restricted-0001',
            company_name='Restricted Route Company',
            source_type='manual',
            source_detail='Local test',
            email='restricted-route@example.invalid',
            assignment_mode='self',
        )
        point = owned.company.contact_points.get()
        point.usage_status = 'restricted'
        point.save(update_fields=['usage_status', 'updated_at'])
        self.client.force_login(self.sales)
        response = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk)]},
        )
        self.assertEqual(response.status_code, 302)
        job = CustomerExportJob.objects.get(created_by=self.sales)
        self.assertTrue(process_customer_export_job(job.pk))
        job.refresh_from_db()
        response = self.client.get(
            reverse('console:customer_pool_export_download', args=[job.pk])
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(
            list(workbook['联系方式'].iter_rows(values_only=True)),
            [('企业ID', '企业名称', '类型', '联系方式', '分机', '状态', '使用状态')],
        )
        _private_path(job.storage_path).unlink()

    def test_current_filter_export_respects_mine_vs_team_scope(self):
        sales_owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-scope-sales-0001',
            company_name='Sales Owned Scope Company',
            source_type='manual',
            source_detail='Local test',
            email='sales-scope@example.invalid',
            assignment_mode='self',
        )
        manager_owned = create_manual_customer(
            site=self.site,
            actor=self.manager,
            idempotency_token='view-export-scope-manager-0001',
            company_name='Manager Owned Scope Company',
            source_type='manual',
            source_detail='Local test',
            email='manager-scope@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.manager)
        mine = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'current', 'scope': 'mine'},
        )
        self.assertEqual(mine.status_code, 302)
        mine_job = CustomerExportJob.objects.filter(created_by=self.manager).latest('id')
        self.assertTrue(process_customer_export_job(mine_job.pk))
        mine_job.refresh_from_db()
        mine_download = self.client.get(
            reverse('console:customer_pool_export_download', args=[mine_job.pk])
        )
        workbook = load_workbook(BytesIO(mine_download.content), read_only=True, data_only=True)
        mine_ids = [
            row[0]
            for row in list(workbook['企业'].iter_rows(values_only=True))[1:]
        ]
        self.assertEqual(mine_ids, [manager_owned.company.pk])

        team = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'current', 'scope': 'team'},
        )
        self.assertEqual(team.status_code, 302)
        team_job = CustomerExportJob.objects.filter(created_by=self.manager).latest('id')
        self.assertTrue(process_customer_export_job(team_job.pk))
        team_job.refresh_from_db()
        team_download = self.client.get(
            reverse('console:customer_pool_export_download', args=[team_job.pk])
        )
        workbook = load_workbook(BytesIO(team_download.content), read_only=True, data_only=True)
        team_ids = {
            row[0]
            for row in list(workbook['企业'].iter_rows(values_only=True))[1:]
        }
        self.assertEqual(
            team_ids,
            {sales_owned.company.pk, manager_owned.company.pk},
        )
        for job in CustomerExportJob.objects.filter(created_by=self.manager):
            path = _private_path(job.storage_path)
            if path.exists():
                path.unlink()

    def test_mixed_authorized_and_unauthorized_selection_cancels_export(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='view-export-owned-0002',
            company_name='Owned Export Company Two',
            source_type='manual',
            source_detail='Local test',
            email='owned-export-two@example.invalid',
            assignment_mode='self',
        )
        self.client.force_login(self.sales)

        response = self.client.post(
            reverse('console:customer_pool_export'),
            {'mode': 'selected', 'company_ids': [str(owned.company.pk), str(self.available.company.pk)]},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '所选范围包含不存在或无权导出的客户，已取消整个导出。')
        self.assertFalse(CustomerExportJob.objects.filter(created_by=self.sales).exists())
