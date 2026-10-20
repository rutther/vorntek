from types import SimpleNamespace
from unittest.mock import patch

from django.test import RequestFactory, SimpleTestCase
from openpyxl import Workbook, load_workbook

from console.lead_exports import (
    MAX_LEAD_EXPORT_BYTES,
    MAX_LEAD_EXPORT_ROWS,
    LeadExportLimitError,
    build_leads_csv,
    build_leads_xlsx,
)
from console.views import lead_submission_export_csv


def synthetic_submission(index: int = 1):
    return SimpleNamespace(
        id=index,
        submitted_at=None,
        stage='new',
        assignee=None,
        full_name=f'Lead {index}',
        company='Synthetic Company',
        country='Example',
        email=f'lead{index}@example.invalid',
        phone='+10000000000',
        payload_json={'product_category': 'Synthetic', 'capacity': '1'},
        message='Synthetic request',
        source_channel='direct',
        source_detail='test',
        source_url='https://example.invalid/contact/',
        referrer_url='',
        utm_json={},
        locale=None,
        next_follow_up_at=None,
        follow_up_notes='',
        notification_status='pending',
        duplicate_of_id=None,
    )


class FakeQuerySet:
    def __init__(self, rows, *, reported_count=None):
        self.rows = list(rows)
        self.reported_count = len(self.rows) if reported_count is None else reported_count
        self.iterator_called = False

    def count(self):
        return self.reported_count

    def iterator(self, *, chunk_size):
        self.iterator_called = True
        self.chunk_size = chunk_size
        yield from self.rows


class LeadExportResourceBoundaryTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def test_row_limit_rejects_before_reading_export_rows(self):
        queryset = FakeQuerySet([], reported_count=MAX_LEAD_EXPORT_ROWS + 1)
        with self.assertRaisesRegex(LeadExportLimitError, '最多导出'):
            build_leads_csv(queryset)
        self.assertFalse(queryset.iterator_called)

    def test_csv_uses_seekable_spooled_stream_and_enforces_byte_limit(self):
        stream, count = build_leads_csv(FakeQuerySet([synthetic_submission()]), spool_memory_bytes=128)
        try:
            self.assertEqual(count, 1)
            self.assertTrue(stream.seekable())
            self.assertTrue(stream._rolled)
            self.assertEqual(stream.read(3), b'\xef\xbb\xbf')
        finally:
            stream.close()

        oversized = synthetic_submission(2)
        oversized.message = 'x' * 513
        with self.assertRaisesRegex(LeadExportLimitError, '文件大小'):
            build_leads_csv(FakeQuerySet([oversized]), max_bytes=512)

    def test_xlsx_uses_write_only_workbook_and_seekable_spooled_stream(self):
        with patch('console.lead_exports.Workbook', wraps=Workbook) as workbook_factory:
            stream, count = build_leads_xlsx(
                FakeQuerySet([synthetic_submission()]), spool_memory_bytes=128
            )
        try:
            workbook_factory.assert_called_once_with(write_only=True)
            self.assertEqual(count, 1)
            self.assertTrue(stream.seekable())
            self.assertTrue(stream._rolled)
            workbook = load_workbook(stream, read_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(rows[1][0], 1)
            self.assertEqual(rows[1][4], 'Lead 1')
        finally:
            stream.close()

    def test_csv_view_streams_spooled_file_and_rejects_oversized_query(self):
        request = self.factory.get('/admin/sales/leads/export.csv')
        request.user = SimpleNamespace(is_authenticated=True)
        queryset = FakeQuerySet([synthetic_submission()])
        with (
            patch('console.views.default_site_locale', return_value=(object(), object())),
            patch('console.views.requested_locale_code', return_value=None),
            patch('console.views.lead_export_queryset', return_value=queryset),
        ):
            response = lead_submission_export_csv(request)
        try:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response['X-Export-Count'], '1')
            self.assertTrue(b''.join(response.streaming_content).startswith(b'\xef\xbb\xbf'))
        finally:
            response.close()

        oversized = FakeQuerySet([], reported_count=MAX_LEAD_EXPORT_ROWS + 1)
        with (
            patch('console.views.default_site_locale', return_value=(object(), object())),
            patch('console.views.requested_locale_code', return_value=None),
            patch('console.views.lead_export_queryset', return_value=oversized),
        ):
            rejected = lead_submission_export_csv(request)
        self.assertEqual(rejected.status_code, 413)
        self.assertIn('最多导出', rejected.content.decode('utf-8'))
        self.assertFalse(oversized.iterator_called)
