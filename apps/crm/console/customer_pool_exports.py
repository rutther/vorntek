from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from collections.abc import Callable
from datetime import timedelta
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from leads.customer_standard21_import import STANDARD21_HEADERS, pool_row_standard_values
from leads.models import Company, CustomerExportJob, CustomerPoolRow

from .access import crm_owned_queryset_for_user
from .capabilities import (
    DataScope,
    SalesCapability,
    can_sales,
    require_sales,
    sales_capability_scope,
)
from .customer_pool_queries import (
    all_customers_queryset,
    country_condition,
    filter_import_batch,
    filter_pool_state,
    normalize_source_filters,
    parse_country_filter,
    parse_value_bound,
)
from .customer_pool_search import resolve_sensitive_search
from .payloads import default_site_locale


MAX_EXPORT_ROWS = 5000
EXPORT_TTL_HOURS = 24
PROCESSING_STALE_MINUTES = 15
EXPORT_FIELD_GROUPS = ('company', 'ownership', 'source', 'contact_point')
EXPORT_FIELD_LABELS = {
    'company': '企业基本信息',
    'ownership': '负责人和团队',
    'source': '数据来源和来源明细',
    'contact_point': '可用联系方式',
}
EXPORT_FORMATS = ('xlsx',)
MAX_STANDARD21_ROWS = 50000
STANDARD21_FILE_NAME = 'vorntek-customer-pool-standard21.csv'
STANDARD21_CONTENT_TYPE = 'text/csv; charset=utf-8'
STANDARD21_REQUIRED_COLUMNS = ('phone', 'email')
STANDARD21_PHONE_COLUMNS = frozenset({'phone', 'phone_2', 'phone_3'})
SPREADSHEET_FORMULA_PREFIXES = ('=', '+', '-', '@')
PHONE_TEXT_PATTERN = re.compile(r'^[+-]?[0-9(). /-]+$')
STANDARD21_COLUMN_LABELS = {
    'phone': '电话',
    'email': '邮箱',
    'country_name': '国家',
    'country': '国家代码',
    'person_name': '姓名',
    'company': '企业',
    'value': '价值分',
    'email_2': '邮箱2',
    'email_3': '邮箱3',
    'phone_2': '电话2',
    'phone_3': '电话3',
    'route_type': '路线类型',
    'route_tier': '路线层级',
    'account_id': '账户ID',
    'whatsapp_confirmed': 'WhatsApp',
    'evidence_V': '证据V',
    'identity_I': '身份I',
    'tech_T': '技术T',
    'priority_P': '优先级P',
    'restriction_note': '限制注记',
    'source_channel': '来源渠道',
}


def _excel_text(value) -> str:
    text = str(value or '')
    if text.startswith(('=', '+', '-', '@')):
        return "'" + text
    return text


def _base_queryset(*, site, user):
    return (
        crm_owned_queryset_for_user(Company.objects.filter(site=site), user)
        .select_related('owner_user', 'team', 'pool_state')
        .prefetch_related('customer_sources', 'contact_points')
        .order_by('name', 'id')
    )


def _filtered_queryset(request, *, site):
    queryset = _base_queryset(site=site, user=request.user)
    if request.POST.get('mode') == 'selected':
        ids = []
        invalid_selection = False
        for raw in request.POST.getlist('company_ids'):
            try:
                company_id = int(raw)
            except (TypeError, ValueError):
                invalid_selection = True
                continue
            if company_id <= 0:
                invalid_selection = True
                continue
            ids.append(company_id)
        unique_ids = list(dict.fromkeys(ids))
        return queryset.filter(pk__in=unique_ids), {
            'mode': 'selected',
            'requested_ids': len(unique_ids),
            'invalid_selection': invalid_selection or not unique_ids,
        }
    scope = str(request.POST.get('scope') or '')
    if scope == 'mine':
        queryset = queryset.filter(pool_state__state='owned', owner_user=request.user)
    elif scope == 'all' and sales_capability_scope(request.user, SalesCapability.POOL_EXPORT) is DataScope.ALL:
        queryset = queryset.filter(pk__in=all_customers_queryset(site=site, user=request.user).values('pk'))
    elif scope == 'team':
        queryset = queryset.filter(pool_state__state='owned')
    elif scope in {'review', 'archived'} and can_sales(
        request.user, SalesCapability.POOL_REVIEW
    ):
        queryset = queryset.filter(pool_state__state=scope)
    else:
        queryset = queryset.none()
    query = ' '.join(str(request.POST.get('q') or '').split())[:120]
    contact_search_token = str(request.POST.get('contact_search') or '')[:80]
    contact_search = resolve_sensitive_search(request, contact_search_token)
    source_type = str(request.POST.get('source') or 'all')
    intake_method = str(request.POST.get('intake') or 'all')
    source_type, intake_method = normalize_source_filters(source_type, intake_method)
    country_codes, country_terms = parse_country_filter(request.POST.getlist('country'))
    evidence_status = str(request.POST.get('evidence') or 'all')
    owner_id = int(request.POST.get('owner')) if str(request.POST.get('owner') or '').isdigit() else 0
    team_id = int(request.POST.get('team')) if str(request.POST.get('team') or '').isdigit() else 0
    batch_id = int(request.POST.get('batch')) if str(request.POST.get('batch') or '').isdigit() else 0
    value_min = parse_value_bound(request.POST.get('value_min'))
    value_max = parse_value_bound(request.POST.get('value_max'))
    if value_min is not None and value_max is not None and value_min > value_max:
        value_min, value_max = value_max, value_min
    state_filter = str(request.POST.get('state') or 'all')
    queryset = filter_pool_state(queryset, state_filter)
    if query:
        queryset = queryset.filter(
            Q(name__icontains=query)
            | Q(website__icontains=query)
            | Q(industry__icontains=query)
            | Q(customer_sources__source_detail__icontains=query)
        )
    if contact_search:
        queryset = queryset.filter(contact_points__raw_value__icontains=contact_search)
    if source_type != 'all':
        queryset = queryset.filter(customer_sources__source_type=source_type)
    if intake_method != 'all':
        queryset = queryset.filter(customer_sources__intake_method=intake_method)
    if country_codes or country_terms:
        queryset = queryset.filter(
            country_condition(
                codes=country_codes, terms=country_terms,
                code_field='country_code', name_field='country',
            )
        )
    if evidence_status in {'pending', 'declared', 'verified', 'rejected'}:
        queryset = queryset.filter(pool_state__evidence_status=evidence_status)
    if owner_id:
        queryset = queryset.filter(owner_user_id=owner_id)
    if team_id:
        queryset = queryset.filter(team_id=team_id)
    if batch_id:
        queryset = filter_import_batch(queryset, batch_id)
    if value_min is not None:
        queryset = queryset.filter(value__gte=value_min)
    if value_max is not None:
        queryset = queryset.filter(value__lte=value_max)
    return queryset.distinct(), {
        'mode': 'current_filter',
        'scope': scope,
        'state': state_filter,
        'query': query,
        'source': source_type,
        'intake': intake_method,
        'country_codes': country_codes,
        'country_terms': country_terms,
        'contact_search_applied': bool(contact_search),
        'evidence': evidence_status,
        'owner_id': owner_id,
        'team_id': team_id,
        'batch_id': batch_id,
        'value_min': value_min,
        'value_max': value_max,
    }


def _normalized_fields(raw_fields) -> tuple[str, ...] | None:
    if not isinstance(raw_fields, (list, tuple)):
        return None
    fields = tuple(str(value or '').strip() for value in raw_fields)
    if (
        not fields
        or 'company' not in fields
        or len(fields) != len(set(fields))
        or any(field not in EXPORT_FIELD_GROUPS for field in fields)
    ):
        return None
    return tuple(field for field in EXPORT_FIELD_GROUPS if field in fields)


def normalize_standard21_columns(raw_columns) -> tuple[str, ...] | None:
    if not isinstance(raw_columns, (list, tuple)):
        return None
    selected = [str(value or '').strip() for value in raw_columns]
    if not selected or len(selected) != len(set(selected)):
        return None
    if any(column not in STANDARD21_HEADERS for column in selected):
        return None
    if any(column not in selected for column in STANDARD21_REQUIRED_COLUMNS):
        return None
    return tuple(header for header in STANDARD21_HEADERS if header in selected)


def _standard21_spreadsheet_text(header: str, value) -> str:
    """Keep phone text stable while neutralizing active spreadsheet formulas."""

    text = str(value or '')
    if not text.startswith(SPREADSHEET_FORMULA_PREFIXES):
        return text
    if header in STANDARD21_PHONE_COLUMNS and PHONE_TEXT_PATTERN.fullmatch(text):
        return text
    return "'" + text


def standard21_csv_bytes(pool_rows, columns=STANDARD21_HEADERS) -> bytes:
    """Restore persisted rows as deterministic, spreadsheet-safe UTF-8 CSV."""

    normalized_columns = normalize_standard21_columns(columns)
    if normalized_columns is None:
        raise ValueError('invalid standard21 columns')
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator='\r\n')
    writer.writerow(normalized_columns)
    for pool_row in pool_rows:
        by_header = dict(zip(STANDARD21_HEADERS, pool_row_standard_values(pool_row)))
        writer.writerow([
            _standard21_spreadsheet_text(header, by_header[header])
            for header in normalized_columns
        ])
    return '\ufeff'.encode('utf-8') + buffer.getvalue().encode('utf-8')


def _standard21_pool_rows(*, site, company_ids):
    return CustomerPoolRow.objects.filter(
        site=site,
        company_id__in=company_ids,
    ).order_by('import_batch_id', 'row_number', 'row_key')


def _workbook(
    companies: list[Company],
    fields: tuple[str, ...] | list[str] = EXPORT_FIELD_GROUPS,
) -> Workbook:
    normalized_fields = _normalized_fields(fields)
    if normalized_fields is None:
        raise ValueError('invalid export fields')
    enabled = set(normalized_fields)
    workbook = Workbook()
    enterprise = workbook.active
    enterprise.title = '企业'
    enterprise_headers = ['企业ID', '企业名称', '国家/地区', '城市', '行业', '网站']
    if 'ownership' in enabled:
        enterprise_headers.extend(['负责人', '团队', '公海状态'])
    if 'source' in enabled:
        enterprise_headers.extend(['首次数据来源', '首次进入方式', '来源数量'])
    enterprise.append(enterprise_headers)

    contacts = None
    if 'contact_point' in enabled:
        contacts = workbook.create_sheet('联系方式')
        contacts.append(['企业ID', '企业名称', '类型', '联系方式', '分机', '状态', '使用状态'])

    sources_sheet = None
    if 'source' in enabled:
        sources_sheet = workbook.create_sheet('来源明细')
        sources_sheet.append([
            '企业ID', '企业名称', '数据来源', '进入方式', '来源详情',
            '外部记录ID', '发生时间', '接收时间', '证据状态',
        ])

    for company in companies:
        sources = list(company.customer_sources.all())
        first_source = sources[0] if sources else None
        enterprise_row = [
            company.pk,
            _excel_text(company.name),
            _excel_text(company.country),
            _excel_text(company.city),
            _excel_text(company.industry),
            _excel_text(company.website),
        ]
        if 'ownership' in enabled:
            enterprise_row.extend([
                _excel_text(company.owner_user.get_username() if company.owner_user else ''),
                _excel_text(company.team.name if company.team else ''),
                _excel_text(getattr(getattr(company, 'pool_state', None), 'state', '')),
            ])
        if 'source' in enabled:
            enterprise_row.extend([
                _excel_text(getattr(first_source, 'source_type', '')),
                _excel_text(getattr(first_source, 'intake_method', '')),
                len(sources),
            ])
        enterprise.append(enterprise_row)

        if contacts is not None:
            for point in company.contact_points.all():
                if point.status in {'do_not_contact', 'invalid'} or point.usage_status == 'restricted':
                    continue
                contacts.append([
                    company.pk,
                    _excel_text(company.name),
                    _excel_text(point.channel),
                    _excel_text(point.raw_value),
                    _excel_text(point.extension),
                    _excel_text(point.status),
                    _excel_text(point.usage_status),
                ])

        if sources_sheet is not None:
            for source in sources:
                sources_sheet.append([
                    company.pk,
                    _excel_text(company.name),
                    _excel_text(source.source_type),
                    _excel_text(source.intake_method),
                    _excel_text(source.source_detail),
                    _excel_text(source.external_record_id),
                    _excel_text(source.occurred_at.isoformat() if source.occurred_at else ''),
                    _excel_text(source.received_at.isoformat() if source.received_at else ''),
                    _excel_text(source.evidence_status),
                ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F5B99')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                42,
                max(12, max(len(str(cell.value or '')) for cell in column) + 2),
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.number_format = '@'
    return workbook


def _export_root() -> Path:
    return (Path(settings.BASE_DIR) / '.runtime' / 'customer-exports').resolve()


def _private_path(value: str) -> Path | None:
    if not value:
        return None
    root = _export_root()
    supplied = Path(value)
    # New jobs store portable paths relative to the private export root. Legacy
    # absolute paths remain valid only when already inside the current root;
    # never guess a basename mapping for a file from another server.
    if not supplied.is_absolute() and ('\\' in value or ':' in value or '..' in supplied.parts):
        return None
    candidate = (supplied if supplied.is_absolute() else root / supplied).resolve()
    if not candidate.is_relative_to(root) or candidate.suffix.lower() != '.xlsx':
        return None
    return candidate


def _remove_private_file(job: CustomerExportJob) -> None:
    path = _private_path(job.storage_path)
    if path and path.is_file():
        path.unlink()


def _authorized_snapshot(*, job: CustomerExportJob, user) -> list[Company] | None:
    raw_ids = (job.scope_json or {}).get('company_ids')
    if not isinstance(raw_ids, list) or not raw_ids:
        return None
    try:
        ids = [int(value) for value in raw_ids]
    except (TypeError, ValueError):
        return None
    if len(ids) != len(set(ids)) or any(value <= 0 for value in ids):
        return None
    companies = list(_base_queryset(site=job.site, user=user).filter(pk__in=ids))
    if {item.pk for item in companies} != set(ids):
        return None
    return companies


def _build_private_export(
    companies: list[Company],
    fields: tuple[str, ...] | list[str] = EXPORT_FIELD_GROUPS,
) -> tuple[Path, str]:
    output = BytesIO()
    _workbook(companies, fields).save(output)
    payload = output.getvalue()
    digest = hashlib.sha256(payload).hexdigest()
    root = _export_root()
    root.mkdir(parents=True, exist_ok=True)
    private_path = root / f'{uuid.uuid4().hex}.xlsx'
    private_path.write_bytes(payload)
    return private_path, digest


def _fail_processing_job(job_id: int) -> None:
    with transaction.atomic():
        job = CustomerExportJob.objects.select_for_update().filter(pk=job_id).first()
        if job is None or job.status != 'processing':
            return
        _remove_private_file(job)
        job.status = 'failed'
        job.storage_path = ''
        job.sha256 = ''
        job.expires_at = None
        job.completed_at = timezone.now()
        job.save(
            update_fields=[
                'status', 'storage_path', 'sha256', 'expires_at',
                'completed_at', 'updated_at',
            ]
        )


def process_customer_export_job(job_id: int) -> bool:
    """Claim and render one export outside the web request.

    The final state is committed only after a second permission check. A user
    cancellation that wins while the workbook is being generated therefore
    remains authoritative and the temporary file is removed.
    """

    stale_before = timezone.now() - timedelta(minutes=PROCESSING_STALE_MINUTES)
    with transaction.atomic():
        job = (
            CustomerExportJob.objects.select_for_update()
            .select_related('site', 'created_by')
            .filter(pk=job_id)
            .first()
        )
        if job is None:
            return False
        claimable = job.status == 'pending' or (
            job.status == 'processing' and job.updated_at <= stale_before
        )
        if not claimable:
            return False
        if not job.created_by.is_active or not can_sales(
            job.created_by, SalesCapability.POOL_EXPORT
        ):
            job.status = 'failed'
            job.completed_at = timezone.now()
            job.save(update_fields=['status', 'completed_at', 'updated_at'])
            return False
        job.status = 'processing'
        job.completed_at = None
        job.save(update_fields=['status', 'completed_at', 'updated_at'])
        actor = job.created_by

    try:
        job = CustomerExportJob.objects.select_related('site').get(pk=job_id)
        fields = _normalized_fields(job.fields_json)
        if fields is None or job.format not in EXPORT_FORMATS:
            _fail_processing_job(job_id)
            return False
        companies = _authorized_snapshot(job=job, user=actor)
        if companies is None:
            _fail_processing_job(job_id)
            return False
        private_path, digest = _build_private_export(companies, fields)
    except Exception:
        _fail_processing_job(job_id)
        return False

    with transaction.atomic():
        locked = (
            CustomerExportJob.objects.select_for_update()
            .select_related('site', 'created_by')
            .get(pk=job_id)
        )
        if locked.status != 'processing':
            private_path.unlink(missing_ok=True)
            return False
        if not locked.created_by.is_active or not can_sales(
            locked.created_by, SalesCapability.POOL_EXPORT
        ) or _authorized_snapshot(job=locked, user=locked.created_by) is None:
            private_path.unlink(missing_ok=True)
            locked.status = 'failed'
            locked.completed_at = timezone.now()
            locked.save(update_fields=['status', 'completed_at', 'updated_at'])
            return False
        locked.status = 'ready'
        locked.record_count = len(companies)
        locked.storage_path = private_path.relative_to(_export_root()).as_posix()
        locked.sha256 = digest
        locked.expires_at = timezone.now() + timedelta(hours=EXPORT_TTL_HOURS)
        locked.completed_at = timezone.now()
        locked.save(
            update_fields=[
                'status', 'record_count', 'storage_path', 'sha256',
                'expires_at', 'completed_at', 'updated_at',
            ]
        )
    return True


def process_pending_customer_exports(
    *, limit: int = 20, stop_requested: Callable[[], bool] | None = None,
) -> tuple[int, int]:
    if stop_requested and stop_requested():
        return 0, 0
    bounded_limit = max(1, min(int(limit), 100))
    stale_before = timezone.now() - timedelta(minutes=PROCESSING_STALE_MINUTES)
    job_ids = list(
        CustomerExportJob.objects.filter(
            Q(status='pending')
            | Q(status='processing', updated_at__lte=stale_before)
        )
        .order_by('created_at', 'id')
        .values_list('id', flat=True)[:bounded_limit]
    )
    attempted = succeeded = 0
    for job_id in job_ids:
        if stop_requested and stop_requested():
            break
        attempted += 1
        succeeded += bool(process_customer_export_job(job_id))
    return attempted, succeeded


def _download_response(payload: bytes) -> HttpResponse:
    response = HttpResponse(
        payload,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="vorntek-customers.xlsx"'
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


@login_required
@require_GET
def customer_pool_export_status(request):
    require_sales(request.user, SalesCapability.POOL_EXPORT)
    site, _locale = default_site_locale(None)
    now = timezone.now()
    jobs = CustomerExportJob.objects.filter(
        site=site,
        created_by=request.user,
    ).order_by('-created_at', '-id')[:10]
    snapshot = [
        {
            'id': job.pk,
            'status': (
                'expired'
                if job.status == 'ready' and job.expires_at and job.expires_at <= now
                else job.status
            ),
            'updated_at': job.updated_at.isoformat(),
        }
        for job in jobs
    ]
    response = JsonResponse({'jobs': snapshot})
    response['Cache-Control'] = 'private, no-store'
    return response


@login_required
@require_POST
def customer_pool_export(request):
    require_sales(request.user, SalesCapability.POOL_EXPORT)
    site, _locale = default_site_locale(None)
    queryset, export_scope = _filtered_queryset(request, site=site)
    companies = list(queryset[:MAX_EXPORT_ROWS + 1])
    if export_scope.get('mode') == 'selected' and (
        export_scope.get('invalid_selection')
        or len(companies) != export_scope.get('requested_ids')
    ):
        messages.error(request, '所选范围包含不存在或无权导出的客户，已取消整个导出。')
        return redirect('console:customer_pool')
    if not companies:
        messages.error(request, '没有可导出的已授权客户；销售账号只能导出自己负责的数据。')
        return redirect('console:customer_pool')
    if len(companies) > MAX_EXPORT_ROWS:
        messages.error(request, f'单次最多导出 {MAX_EXPORT_ROWS} 条，请缩小筛选范围。')
        return redirect('console:customer_pool')
    requested_fields = request.POST.getlist('fields')
    fields = _normalized_fields(requested_fields or list(EXPORT_FIELD_GROUPS))
    if fields is None:
        messages.error(request, '导出字段无效；企业基本信息为必选项。')
        return redirect('console:customer_pool')
    export_format = str(request.POST.get('format') or 'xlsx').strip().casefold()
    if export_format not in EXPORT_FORMATS:
        messages.error(request, '当前只支持 XLSX；CSV 无法无损表达一家企业的多条联系方式。')
        return redirect('console:customer_pool')
    export_scope['company_ids'] = [item.pk for item in companies]
    job = CustomerExportJob.objects.create(
        site=site,
        created_by=request.user,
        status='pending',
        scope_json=export_scope,
        fields_json=list(fields),
        format=export_format,
        record_count=len(companies),
    )
    messages.success(request, f'导出任务 #{job.pk} 已进入后台队列，页面会自动更新状态。')
    return redirect('console:customer_pool')


@login_required
@require_POST
def customer_pool_export_standard21(request):
    require_sales(request.user, SalesCapability.POOL_EXPORT)
    site, _locale = default_site_locale(None)
    queryset, export_scope = _filtered_queryset(request, site=site)
    companies = list(queryset[:MAX_EXPORT_ROWS + 1])
    if export_scope.get('mode') == 'selected' and (
        export_scope.get('invalid_selection')
        or len(companies) != export_scope.get('requested_ids')
    ):
        messages.error(request, '所选范围包含不存在或无权导出的客户，已取消整个导出。')
        return redirect('console:customer_pool')
    if not companies:
        messages.error(request, '没有可导出的已授权客户。')
        return redirect('console:customer_pool')
    if len(companies) > MAX_EXPORT_ROWS:
        messages.error(request, f'单次最多导出 {MAX_EXPORT_ROWS} 家企业，请缩小筛选范围。')
        return redirect('console:customer_pool')
    requested_columns = request.POST.getlist('columns')
    columns = (
        STANDARD21_HEADERS
        if not requested_columns
        else normalize_standard21_columns(requested_columns)
    )
    if columns is None:
        messages.error(request, '导出列无效：电话和邮箱必须保留，且不能包含重复或未知列。')
        return redirect('console:customer_pool')
    pool_rows = list(
        _standard21_pool_rows(
            site=site,
            company_ids=[company.pk for company in companies],
        )[: MAX_STANDARD21_ROWS + 1]
    )
    if not pool_rows:
        messages.error(request, '所选范围没有标准 21 列客户池记录。')
        return redirect('console:customer_pool')
    if len(pool_rows) > MAX_STANDARD21_ROWS:
        messages.error(request, f'单次最多导出 {MAX_STANDARD21_ROWS} 行，请缩小筛选范围。')
        return redirect('console:customer_pool')
    response = HttpResponse(
        standard21_csv_bytes(pool_rows, columns=columns),
        content_type=STANDARD21_CONTENT_TYPE,
    )
    response['Content-Disposition'] = f'attachment; filename="{STANDARD21_FILE_NAME}"'
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


@login_required
@require_GET
def customer_pool_export_download(request, job_id: int):
    require_sales(request.user, SalesCapability.POOL_EXPORT)
    site, _locale = default_site_locale(None)
    job = CustomerExportJob.objects.filter(
        pk=job_id,
        site=site,
        created_by=request.user,
    ).first()
    if job is None:
        return HttpResponse('导出任务不存在或当前账号无权访问。', status=404)
    if job.status != 'ready':
        return HttpResponse('导出文件当前不可下载。', status=410)
    if not job.expires_at or job.expires_at <= timezone.now():
        _remove_private_file(job)
        job.status = 'expired'
        job.save(update_fields=['status'])
        return HttpResponse('导出文件已过期，请重新生成。', status=410)
    if _authorized_snapshot(job=job, user=request.user) is None:
        _remove_private_file(job)
        job.status = 'failed'
        job.save(update_fields=['status'])
        return HttpResponse('当前权限已变化，原导出已作废，请重新生成。', status=403)
    if _normalized_fields(job.fields_json) is None or job.format not in EXPORT_FORMATS:
        _remove_private_file(job)
        job.status = 'failed'
        job.save(update_fields=['status'])
        return HttpResponse('导出字段或格式已失效，请重新生成。', status=410)
    path = _private_path(job.storage_path)
    if path is None or not path.is_file():
        job.status = 'failed'
        job.save(update_fields=['status'])
        return HttpResponse('导出文件不可用，请重新生成。', status=410)
    payload = path.read_bytes()
    if hashlib.sha256(payload).hexdigest() != job.sha256:
        _remove_private_file(job)
        job.status = 'failed'
        job.save(update_fields=['status'])
        return HttpResponse('导出文件校验失败，请重新生成。', status=410)
    return _download_response(payload)


@login_required
@require_POST
@transaction.atomic
def customer_pool_export_retry(request, job_id: int):
    require_sales(request.user, SalesCapability.POOL_EXPORT)
    site, _locale = default_site_locale(None)
    job = CustomerExportJob.objects.select_for_update().filter(
        pk=job_id,
        site=site,
        created_by=request.user,
    ).first()
    if job is None:
        return HttpResponse('导出任务不存在或当前账号无权访问。', status=404)
    if job.status not in {'failed', 'expired'}:
        messages.error(request, '只有失败或过期的导出任务可以重试。')
        return redirect('console:customer_pool')
    companies = _authorized_snapshot(job=job, user=request.user)
    if (
        companies is None
        or _normalized_fields(job.fields_json) is None
        or job.format not in EXPORT_FORMATS
    ):
        messages.error(request, '当前权限或客户范围已变化，请重新选择导出范围。')
        return redirect('console:customer_pool')
    _remove_private_file(job)
    job.status = 'pending'
    job.storage_path = ''
    job.sha256 = ''
    job.expires_at = None
    job.completed_at = None
    job.save(
        update_fields=['status', 'storage_path', 'sha256', 'expires_at', 'completed_at']
    )
    messages.success(request, '导出重试已重新进入后台队列。')
    return redirect('console:customer_pool')


@login_required
@require_POST
@transaction.atomic
def customer_pool_export_cancel(request, job_id: int):
    require_sales(request.user, SalesCapability.POOL_EXPORT)
    site, _locale = default_site_locale(None)
    job = CustomerExportJob.objects.select_for_update().filter(
        pk=job_id,
        site=site,
        created_by=request.user,
    ).first()
    if job is None:
        return HttpResponse('导出任务不存在或当前账号无权访问。', status=404)
    if job.status not in {'pending', 'processing', 'ready'}:
        messages.error(request, '该导出任务当前不能取消。')
        return redirect('console:customer_pool')
    _remove_private_file(job)
    job.status = 'canceled'
    job.completed_at = job.completed_at or timezone.now()
    job.save(update_fields=['status', 'completed_at'])
    messages.success(request, '导出任务已取消，原文件已撤销。')
    return redirect('console:customer_pool')
