from __future__ import annotations

import hashlib
from dataclasses import replace
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied
from django.db import IntegrityError, connection
from django.test import TestCase
from django.utils import timezone
from io import BytesIO
from openpyxl import Workbook
from unittest.mock import patch
import zipfile

from console.access import (
    GROUP_BY_ROLE,
    ROLE_SALES,
    ROLE_SALES_MANAGER,
    ROLE_SYSTEM_ADMIN,
)
from marketing.models import CanonicalEvent
from sitecore.models import Cta, PageRoute, Site, SiteLocale

from .customer_pool_services import (
    archive_reviewed_customer,
    assign_customer,
    CustomerPoolConflict,
    CustomerPoolError,
    claim_customer,
    claim_customers_bulk,
    create_manual_customer,
    publish_reviewed_customer,
    public_pool_queryset,
    release_customer,
    restore_archived_customer,
)
from .customer_import_services import (
    AFRICA_ACCOUNT_HEADERS,
    AFRICA_CONTACT_HEADERS,
    MIDDLE_EAST_ACCOUNT_HEADERS,
    MIDDLE_EAST_CORRECTION_HEADERS,
    MIDDLE_EAST_EMAIL_HEADERS,
    MIDDLE_EAST_PHONE_HEADERS,
    RESEARCH_SNAPSHOT_PROFILES,
    CustomerImportError,
    commit_customer_import,
    preview_customer_import,
    preview_research_snapshot_import,
)
from .models import (
    Activity,
    Company,
    CompanyContactPoint,
    CompanyPoolState,
    Contact,
    CrmMutationReceipt,
    CustomerSource,
    CustomerImportBatch,
    CustomerImportRow,
    LeadConversion,
    LeadFormDefinition,
    LeadSubmission,
    Opportunity,
    SalesTeam,
    SalesTeamMember,
    Task,
)


class CustomerPoolServiceTests(TestCase):
    unmanaged_models = (
        Site,
        SiteLocale,
        PageRoute,
        CanonicalEvent,
        Cta,
        LeadFormDefinition,
        SalesTeam,
        SalesTeamMember,
        LeadSubmission,
        Company,
        Contact,
        Opportunity,
        LeadConversion,
        Activity,
        Task,
        CrmMutationReceipt,
        CompanyPoolState,
        CompanyContactPoint,
        CustomerSource,
        CustomerImportBatch,
        CustomerImportRow,
    )

    @classmethod
    def setUpClass(cls):
        existing = set(connection.introspection.table_names())
        cls.created_models = []
        try:
            with connection.schema_editor() as editor:
                for model in cls.unmanaged_models:
                    if model._meta.db_table not in existing:
                        editor.create_model(model)
                        cls.created_models.append(model)
                        existing.add(model._meta.db_table)
            super().setUpClass()
        except Exception:
            with connection.schema_editor() as editor:
                for model in reversed(cls.created_models):
                    editor.delete_model(model)
            raise

    @classmethod
    def tearDownClass(cls):
        try:
            super().tearDownClass()
        finally:
            with connection.schema_editor() as editor:
                for model in reversed(cls.created_models):
                    editor.delete_model(model)

    def setUp(self):
        self.site = Site.objects.create(code='pool-test', name='Pool Test', base_url='https://example.invalid', default_locale='en')
        self.team = SalesTeam.objects.create(site=self.site, code='sales', name='Sales')
        group = Group.objects.create(name=GROUP_BY_ROLE[ROLE_SALES])
        self.sales = get_user_model().objects.create_user(username='pool-sales', password='test-password')
        self.sales.groups.add(group)
        self.sales_two = get_user_model().objects.create_user(username='pool-sales-two', password='test-password')
        self.sales_two.groups.add(group)
        manager_group = Group.objects.create(name=GROUP_BY_ROLE[ROLE_SALES_MANAGER])
        self.manager = get_user_model().objects.create_user(username='pool-manager', password='test-password')
        self.manager.groups.add(manager_group)
        admin_group = Group.objects.create(name=GROUP_BY_ROLE[ROLE_SYSTEM_ADMIN])
        self.admin = get_user_model().objects.create_user(username='pool-admin', password='test-password')
        self.admin.groups.add(admin_group)
        SalesTeamMember.objects.create(team=self.team, user=self.sales, membership_role='member')
        SalesTeamMember.objects.create(team=self.team, user=self.sales_two, membership_role='member')
        SalesTeamMember.objects.create(team=self.team, user=self.manager, membership_role='manager')

    def import_workbook_bytes(
        self,
        *,
        formula: bool = False,
        version: str = '版本 1',
        source_type: str = 'research',
        source_detail: str = 'Public website',
    ) -> bytes:
        workbook = Workbook()
        instructions = workbook.active
        instructions.title = '填写说明'
        instructions.append(['New Crown 客户导入模板', version])
        companies = workbook.create_sheet('企业')
        companies.append(['企业外部键*', '企业名称*', '国家/地区', '城市', '行业', '网站', '数据来源*', '来源详情', '备注'])
        companies.append(['company-001', 'Imported Beverage', 'Ghana', 'Accra', 'Beverage', 'https://imported.example.org', source_type, source_detail, ''])
        if formula:
            companies['B2'] = '=CONCAT("Imported", " Beverage")'
        contacts = workbook.create_sheet('联系方式')
        contacts.append(['企业外部键*', '联系人姓名', '职位', '联系方式类型*', '联系方式值*', '分机', '用途', '使用状态', '证据说明'])
        contacts.append(['company-001', '', '', 'email', 'info@imported.example.org', '', 'business', 'unknown', 'Public contact page'])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def africa_research_workbook_bytes(self) -> bytes:
        workbook = Workbook()
        contacts = workbook.active
        contacts.title = '第二轮全量联系方式'
        for _unused in range(3):
            contacts.append([])
        contacts.append(list(AFRICA_CONTACT_HEADERS))
        contacts.append([
            'Ghana', '+233 20 000 0001', 'ignored-on-phone@example.invalid', '+233',
            '企业总机', '+233 20 000 0001', 'Africa Fixture Beverage', 'AF-001',
            'Public Person', 'Public Role', '首通仅核验需求', '当前',
            'https://example.invalid/africa-phone', '2026-09-08', '饮料制造商',
        ])
        contacts.append([
            'Ghana', '', 'sales@example.invalid', '', '企业邮箱', '',
            'Africa Fixture Beverage', 'AF-001', '', '', '业务咨询', '当前',
            'https://example.invalid/africa-email', '2026-09-08', '饮料制造商',
        ])
        accounts = workbook.create_sheet('第二轮账户证据')
        for _unused in range(3):
            accounts.append([])
        accounts.append(list(AFRICA_ACCOUNT_HEADERS))
        accounts.append([
            'Ghana', 'Africa Fixture Beverage', 'AF-001', '1', '1', '饮料制造商',
            'https://example.invalid/africa-company', '2026-09-08', '',
        ])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def middle_east_research_workbook_bytes(self, *, email='sales@example.invalid') -> bytes:
        workbook = Workbook()
        accounts = workbook.active
        accounts.title = '中东账户总表'
        accounts.append(list(MIDDLE_EAST_ACCOUNT_HEADERS))
        accounts.append([
            'Jordan', 'Middle East Fixture Beverage', 'ME-001', '果汁', 'Amman',
            '+962 6 000 0001', email, '完成',
        ])
        phones = workbook.create_sheet('联系方式总表')
        phones.append(list(MIDDLE_EAST_PHONE_HEADERS))
        phones.append([
            'Jordan', '+962 6 000 0001', 'ignored-on-phone@example.invalid',
            'Middle East Fixture Beverage', '企业总机', '业务咨询', 'Public Person',
            'Public Role', '+962', 'Amman', 'https://example.invalid/me-phone', '',
            '2026-09-07', '2026-09-08', '公开来源', 'ME-001', '', '+962 6 000 0001',
            '', '是', '', '', '是', '是',
        ])
        emails = workbook.create_sheet('全量邮箱')
        emails.append(list(MIDDLE_EAST_EMAIL_HEADERS))
        emails.append([
            'Jordan', '', email, 'Middle East Fixture Beverage', '企业邮箱', '业务咨询',
            'https://example.invalid/me-email', '2026-09-07', '2026-09-08',
            '公开来源', 'ME-001', '是', '',
        ])
        corrections = workbook.create_sheet('历史号码纠正')
        corrections.append(list(MIDDLE_EAST_CORRECTION_HEADERS))
        corrections.append([
            'Jordan', 'ME-001', 'Middle East Fixture Beverage', '+962 6 000 0001',
            '企业总机', '排除号码', '不使用', 'https://example.invalid/me-correction',
            '2026-09-08', '合成测试纠正记录',
        ])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_import_rejects_wrong_template_version_and_formulas(self):
        with self.assertRaises(CustomerImportError) as version_error:
            preview_customer_import(
                site=self.site,
                actor=self.admin,
                file_bytes=self.import_workbook_bytes(version='版本 0'),
                original_name='old-template.xlsx',
            )
        self.assertEqual(version_error.exception.workflow_code, 'template_version_invalid')

        with self.assertRaises(CustomerImportError) as formula_error:
            preview_customer_import(
                site=self.site,
                actor=self.admin,
                file_bytes=self.import_workbook_bytes(formula=True),
                original_name='formula-template.xlsx',
            )
        self.assertEqual(formula_error.exception.workflow_code, 'formula_rejected')

    def test_import_rejects_suspicious_archive_compression(self):
        payload = BytesIO()
        with zipfile.ZipFile(payload, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/workbook.xml', b'0' * (1024 * 1024))
        with self.assertRaises(CustomerImportError) as caught:
            preview_customer_import(
                site=self.site,
                actor=self.admin,
                file_bytes=payload.getvalue(),
                original_name='compressed-bomb.xlsx',
            )
        self.assertEqual(caught.exception.workflow_code, 'archive_ratio_invalid')

    def test_research_snapshot_rejects_any_file_not_matching_locked_hash(self):
        with self.assertRaises(CustomerImportError) as caught:
            preview_research_snapshot_import(
                site=self.site,
                actor=self.admin,
                file_bytes=self.africa_research_workbook_bytes(),
                original_name='africa-v345.xlsx',
                profile_code='africa_v345',
            )
        self.assertEqual(caught.exception.workflow_code, 'research_snapshot_hash_mismatch')

    def test_africa_snapshot_adapter_uses_canonical_routes_without_fake_contacts(self):
        payload = self.africa_research_workbook_bytes()
        profile = replace(
            RESEARCH_SNAPSHOT_PROFILES['africa_v345'],
            expected_sha256=hashlib.sha256(payload).hexdigest(),
        )
        with patch.dict(RESEARCH_SNAPSHOT_PROFILES, {'africa_v345': profile}):
            preview = preview_research_snapshot_import(
                site=self.site,
                actor=self.admin,
                file_bytes=payload,
                original_name='africa-v345-fixture.xlsx',
                profile_code='africa_v345',
            )
            replay = preview_research_snapshot_import(
                site=self.site,
                actor=self.admin,
                file_bytes=payload,
                original_name='africa-v345-fixture.xlsx',
                profile_code='africa_v345',
            )
        self.assertTrue(replay.replayed)
        self.assertEqual(replay.batch.pk, preview.batch.pk)
        self.assertEqual(preview.batch.counts_json['new'], 1)
        self.assertEqual(preview.batch.counts_json['contact_rows'], 2)

        commit_customer_import(batch_id=preview.batch.pk, site=self.site, actor=self.admin)

        company = Company.objects.get(name='Africa Fixture Beverage')
        self.assertEqual(Contact.objects.filter(company=company).count(), 0)
        self.assertEqual(
            set(company.contact_points.values_list('channel', 'raw_value')),
            {('phone', '+233 20 000 0001'), ('email', 'sales@example.invalid')},
        )
        source = CustomerSource.objects.get(company=company)
        self.assertEqual(source.external_record_id, 'research:africa_v345:AF-001')
        self.assertEqual(source.evidence_json['snapshot'], 'africa_v345')
        self.assertEqual(company.pool_state.state, 'review')
        self.assertEqual(LeadSubmission.objects.count(), 0)

    def test_middle_east_snapshot_correction_restricts_phone_and_increment_is_additive(self):
        first_payload = self.middle_east_research_workbook_bytes()
        first_profile = replace(
            RESEARCH_SNAPSHOT_PROFILES['middle_east_v013'],
            expected_sha256=hashlib.sha256(first_payload).hexdigest(),
        )
        with patch.dict(RESEARCH_SNAPSHOT_PROFILES, {'middle_east_v013': first_profile}):
            first = preview_research_snapshot_import(
                site=self.site,
                actor=self.admin,
                file_bytes=first_payload,
                original_name='middle-east-v013-fixture.xlsx',
                profile_code='middle_east_v013',
            )
        commit_customer_import(batch_id=first.batch.pk, site=self.site, actor=self.admin)
        company = Company.objects.get(name='Middle East Fixture Beverage')
        self.assertEqual(
            company.contact_points.get(channel='phone').usage_status,
            'restricted',
        )
        self.assertEqual(Contact.objects.filter(company=company).count(), 0)

        second_payload = self.middle_east_research_workbook_bytes(
            email='new-route@example.invalid'
        )
        second_profile = replace(
            RESEARCH_SNAPSHOT_PROFILES['middle_east_v013'],
            expected_sha256=hashlib.sha256(second_payload).hexdigest(),
        )
        with patch.dict(RESEARCH_SNAPSHOT_PROFILES, {'middle_east_v013': second_profile}):
            second = preview_research_snapshot_import(
                site=self.site,
                actor=self.admin,
                file_bytes=second_payload,
                original_name='middle-east-v013-increment-fixture.xlsx',
                profile_code='middle_east_v013',
            )
        self.assertEqual(second.batch.rows.get().status, 'supplement')
        commit_customer_import(batch_id=second.batch.pk, site=self.site, actor=self.admin)
        self.assertEqual(Company.objects.filter(site=self.site).count(), 1)
        self.assertEqual(company.contact_points.filter(channel='email').count(), 2)
        self.assertEqual(CustomerSource.objects.filter(company=company).count(), 1)

    def test_manual_create_keeps_source_and_intake_separate_without_fake_contact(self):
        result = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-create-test-0001',
            company_name='  Example   Beverage  ',
            country='Nigeria',
            website='https://example.invalid',
            email='sales@example.invalid',
            source_type='research',
            source_detail='Public company website',
            assignment_mode='self',
        )

        self.assertEqual(result.company.normalized_name, 'example beverage')
        self.assertEqual(result.pool_state.state, 'owned')
        self.assertEqual(Contact.objects.filter(company=result.company).count(), 0)
        self.assertEqual(CompanyContactPoint.objects.filter(company=result.company).count(), 2)
        source = CustomerSource.objects.get(company=result.company)
        self.assertEqual(source.source_type, 'research')
        self.assertEqual(source.intake_method, 'manual_create')
        self.assertIsNone(source.submission_id)
        self.assertEqual(Activity.objects.filter(company=result.company).count(), 1)

    def test_manual_create_defaults_to_review_requires_a_route_and_blocks_duplicates(self):
        with self.assertRaises(CustomerPoolError) as missing_route:
            create_manual_customer(
                site=self.site,
                actor=self.sales,
                idempotency_token='manual-review-missing-route-0001',
                company_name='Review Beverage',
            )
        self.assertEqual(missing_route.exception.workflow_code, 'contact_route_required')

        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-review-create-0001',
            company_name='Review Beverage',
            country='Ghana',
            email='review@example.invalid',
        )
        self.assertEqual(created.pool_state.state, 'review')
        self.assertIsNone(created.company.owner_user_id)
        self.assertEqual(created.company.team_id, self.team.pk)
        self.assertEqual(public_pool_queryset(site=self.site, actor=self.sales).count(), 0)

        with self.assertRaises(CustomerPoolError) as duplicate:
            create_manual_customer(
                site=self.site,
                actor=self.sales,
                idempotency_token='manual-review-duplicate-0001',
                company_name='  review   beverage ',
                country='Ghana',
                phone='+233 20 000 0000',
            )
        self.assertEqual(duplicate.exception.workflow_code, 'duplicate_candidate')

    def test_manager_handoff_moves_current_work_and_preserves_closed_history(self):
        owned = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='handoff-create-owned-0001',
            company_name='Handoff Beverage',
            email='handoff@example.invalid',
            assignment_mode='self',
        )
        contact = Contact.objects.create(
            site=self.site,
            company=owned.company,
            owner_user=self.sales,
            team=self.team,
            full_name='Handoff Contact',
        )
        form = LeadFormDefinition.objects.create(
            site=self.site,
            code='handoff-form',
            name='Handoff form',
            form_schema={},
            config_json={},
        )
        submission = LeadSubmission.objects.create(
            form=form,
            site=self.site,
            assignee=self.sales,
            team=self.team,
            full_name='Handoff Lead',
            submitted_at=timezone.now(),
            stage_updated_at=timezone.now(),
        )
        active_opportunity = Opportunity.objects.create(
            site=self.site,
            company=owned.company,
            primary_contact=contact,
            source_submission=submission,
            owner_user=self.sales,
            team=self.team,
            name='Active handoff opportunity',
            stage='quotation',
        )
        closed_opportunity = Opportunity.objects.create(
            site=self.site,
            company=owned.company,
            owner_user=self.sales,
            team=self.team,
            name='Historical handoff opportunity',
            stage='closed_won',
        )
        LeadConversion.objects.create(
            submission=submission,
            company=owned.company,
            contact=contact,
            opportunity=active_opportunity,
            converted_by_user=self.sales,
            metadata_json={},
        )
        active_task = Task.objects.create(
            site=self.site,
            company=owned.company,
            owner_user=self.sales,
            team=self.team,
            created_by_user=self.sales,
            title='Open handoff task',
            status='open',
            due_at=timezone.now(),
        )
        completed_task = Task.objects.create(
            site=self.site,
            company=owned.company,
            owner_user=self.sales,
            team=self.team,
            created_by_user=self.sales,
            title='Historical handoff task',
            status='completed',
            due_at=timezone.now(),
            completed_at=timezone.now(),
        )

        result = assign_customer(
            company_id=owned.company.pk,
            site=self.site,
            actor=self.manager,
            target_user=self.sales_two,
            expected_version=owned.pool_state.version,
            idempotency_token='manager-handoff-customer-0001',
            reason='Territory coverage changed for this test account.',
        )

        self.assertFalse(result.replayed)
        for record in (owned.company, contact, submission, active_opportunity, active_task):
            record.refresh_from_db()
        closed_opportunity.refresh_from_db()
        completed_task.refresh_from_db()
        self.assertEqual(owned.company.owner_user_id, self.sales_two.pk)
        self.assertEqual(contact.owner_user_id, self.sales_two.pk)
        self.assertEqual(submission.assignee_id, self.sales_two.pk)
        self.assertEqual(active_opportunity.owner_user_id, self.sales_two.pk)
        self.assertEqual(active_task.owner_user_id, self.sales_two.pk)
        self.assertEqual(closed_opportunity.owner_user_id, self.sales.pk)
        self.assertEqual(completed_task.owner_user_id, self.sales.pk)
        activity = Activity.objects.filter(
            company=owned.company,
            subject='交接客户',
        ).get()
        self.assertEqual(activity.metadata_json['previous_owner_id'], self.sales.pk)
        self.assertEqual(activity.metadata_json['target_owner_id'], self.sales_two.pk)
        self.assertEqual(activity.metadata_json['submissions_reassigned'], 1)

        replay = assign_customer(
            company_id=owned.company.pk,
            site=self.site,
            actor=self.manager,
            target_user=self.sales_two,
            expected_version=owned.pool_state.version,
            idempotency_token='manager-handoff-customer-0001',
            reason='Territory coverage changed for this test account.',
        )
        self.assertTrue(replay.replayed)
        self.assertEqual(
            Activity.objects.filter(company=owned.company, subject='交接客户').count(),
            1,
        )

    def test_assignment_rejects_sales_actor_and_cross_team_target(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='assign-guard-create-0001',
            company_name='Assignment Guard Beverage',
            email='assign-guard@example.invalid',
        )
        published = publish_reviewed_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=created.pool_state.version,
            idempotency_token='assign-guard-publish-0001',
            contact_point_id=created.company.contact_points.get().pk,
            team=self.team,
        )
        with self.assertRaises(PermissionDenied):
            assign_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.sales,
                target_user=self.sales_two,
                expected_version=published.pool_state.version,
                idempotency_token='assign-guard-sales-0001',
                reason='Sales cannot assign a public-pool customer.',
            )

        other_team = SalesTeam.objects.create(site=self.site, code='guard-other', name='Guard Other')
        outsider = get_user_model().objects.create_user(username='pool-outsider', password='test-password')
        outsider.groups.add(Group.objects.get(name=GROUP_BY_ROLE[ROLE_SALES]))
        SalesTeamMember.objects.create(team=other_team, user=outsider, membership_role='member')
        with self.assertRaises(CustomerPoolError) as caught:
            assign_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.manager,
                target_user=outsider,
                expected_version=published.pool_state.version,
                idempotency_token='assign-guard-cross-team-0001',
                reason='Cross-team target must be rejected.',
            )
        self.assertEqual(caught.exception.workflow_code, 'assignee_team_invalid')

    def test_bulk_claim_reports_each_company_and_replays_successes(self):
        published = []
        for index in range(2):
            created = create_manual_customer(
                site=self.site,
                actor=self.sales,
                idempotency_token=f'bulk-create-{index}-0001',
                company_name=f'Bulk Claim Beverage {index}',
                email=f'bulk-{index}@example.invalid',
            )
            published.append(publish_reviewed_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.admin,
                expected_version=created.pool_state.version,
                idempotency_token=f'bulk-publish-{index}-0001',
                contact_point_id=created.company.contact_points.get().pk,
                team=self.team,
            ))
        original_versions = [item.pool_state.version for item in published]
        claim_customer(
            company_id=published[1].company.pk,
            site=self.site,
            actor=self.sales_two,
            expected_version=original_versions[1],
            idempotency_token='bulk-competing-claim-0001',
        )

        result = claim_customers_bulk(
            site=self.site,
            actor=self.sales,
            selections=[
                (published[0].company.pk, original_versions[0]),
                (published[1].company.pk, original_versions[1]),
            ],
            idempotency_token='bulk-claim-batch-0001',
        )
        self.assertEqual(result.succeeded, 1)
        self.assertEqual(result.failed, 1)
        self.assertEqual(result.items[0].status, 'claimed')
        self.assertEqual(result.items[1].workflow_code, 'customer_pool_conflict')

        replay = claim_customers_bulk(
            site=self.site,
            actor=self.sales,
            selections=[
                (published[0].company.pk, original_versions[0]),
                (published[1].company.pk, original_versions[1]),
            ],
            idempotency_token='bulk-claim-batch-0001',
        )
        self.assertTrue(replay.items[0].replayed)
        self.assertEqual(replay.items[1].status, 'failed')
        self.assertEqual(
            Activity.objects.filter(
                company=published[0].company,
                subject='领取公海客户',
            ).count(),
            1,
        )

    def test_public_pool_is_team_scoped_and_direct_cross_team_claim_is_denied(self):
        other_team = SalesTeam.objects.create(site=self.site, code='other', name='Other')
        other_sales = get_user_model().objects.create_user(username='pool-other-sales', password='test-password')
        other_sales.groups.add(Group.objects.get(name=GROUP_BY_ROLE[ROLE_SALES]))
        SalesTeamMember.objects.create(team=other_team, user=other_sales, membership_role='member')
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-team-scope-0001',
            company_name='Team Scoped Beverage',
            email='team-scoped@example.invalid',
        )
        published = publish_reviewed_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=created.pool_state.version,
            idempotency_token='review-team-scope-0001',
            contact_point_id=created.company.contact_points.get().pk,
            team=self.team,
        )
        self.assertEqual(public_pool_queryset(site=self.site, actor=self.sales).count(), 1)
        self.assertEqual(public_pool_queryset(site=self.site, actor=other_sales).count(), 0)
        with self.assertRaises(PermissionDenied):
            claim_customer(
                company_id=published.company.pk,
                site=self.site,
                actor=other_sales,
                expected_version=published.pool_state.version,
                idempotency_token='claim-cross-team-0001',
            )

    def test_review_publish_cannot_override_a_restricted_contact_route(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-restricted-route-0001',
            company_name='Restricted Route Beverage',
            email='restricted-route@example.invalid',
        )
        point = created.company.contact_points.get()
        point.usage_status = 'restricted'
        point.save(update_fields=['usage_status', 'updated_at'])

        with self.assertRaises(CustomerPoolError) as blocked:
            publish_reviewed_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.admin,
                expected_version=created.pool_state.version,
                idempotency_token='review-restricted-route-0001',
                contact_point_id=point.pk,
                team=self.team,
            )

        self.assertEqual(blocked.exception.workflow_code, 'contact_point_invalid')
        point.refresh_from_db()
        created.pool_state.refresh_from_db()
        self.assertEqual(point.usage_status, 'restricted')
        self.assertEqual(created.pool_state.state, 'review')

    def test_manual_create_idempotently_replays_without_duplicate_rows(self):
        values = {
            'site': self.site,
            'actor': self.sales,
            'idempotency_token': 'manual-create-test-0002',
            'company_name': 'Idempotent Drinks',
            'email': 'idempotent@example.invalid',
            'assignment_mode': 'self',
        }
        first = create_manual_customer(**values)
        second = create_manual_customer(**values)

        self.assertFalse(first.replayed)
        self.assertTrue(second.replayed)
        self.assertEqual(first.company.pk, second.company.pk)
        self.assertEqual(Company.objects.filter(site=self.site).count(), 1)
        self.assertEqual(CustomerSource.objects.count(), 1)

    def test_review_record_can_be_archived_and_restored_without_deleting_history(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-create-archive-0001',
            company_name='Archive Review Company',
            email='archive-review@example.invalid',
        )
        source_id = created.company.customer_sources.get().pk
        archived = archive_reviewed_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=created.pool_state.version,
            idempotency_token='review-archive-0001',
            reason='Evidence could not be confirmed',
        )
        self.assertEqual(archived.pool_state.state, 'archived')
        self.assertEqual(archived.pool_state.evidence_status, 'rejected')
        self.assertIsNotNone(archived.pool_state.archived_at)
        self.assertTrue(CustomerSource.objects.filter(pk=source_id).exists())

        restored = restore_archived_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=archived.pool_state.version,
            idempotency_token='archive-restore-0001',
            reason='New primary-source evidence received',
        )
        self.assertEqual(restored.pool_state.state, 'review')
        self.assertEqual(restored.pool_state.evidence_status, 'pending')
        self.assertIsNone(restored.pool_state.archived_at)
        self.assertEqual(
            list(
                Activity.objects.filter(company=created.company)
                .order_by('id')
                .values_list('subject', flat=True)
            )[-2:],
            ['核验未通过并归档', '恢复到待核验'],
        )

    def test_claim_uses_version_lock_and_removes_record_from_public_pool(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-create-test-0003',
            company_name='Available Water',
            email='available@example.invalid',
        )
        created = publish_reviewed_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=created.pool_state.version,
            idempotency_token='review-publish-test-0001',
            contact_point_id=created.company.contact_points.get().pk,
            team=self.team,
        )
        self.assertEqual(public_pool_queryset(site=self.site, actor=self.sales).count(), 1)

        claimed = claim_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.sales,
            expected_version=created.pool_state.version,
            idempotency_token='claim-customer-test-0001',
        )

        self.assertEqual(claimed.pool_state.state, 'owned')
        self.assertEqual(claimed.company.owner_user_id, self.sales.pk)
        self.assertEqual(public_pool_queryset(site=self.site, actor=self.sales).count(), 0)
        with self.assertRaises(CustomerPoolConflict):
            claim_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.sales,
                expected_version=1,
                idempotency_token='claim-customer-test-0002',
            )

    def test_available_record_can_be_archived_safely_and_replayed(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-create-retire-0001',
            company_name='Retired Public Pool Company',
            email='retired-pool@example.invalid',
        )
        published = publish_reviewed_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=created.pool_state.version,
            idempotency_token='review-publish-retire-0001',
            contact_point_id=created.company.contact_points.get().pk,
            team=self.team,
        )
        values = {
            'company_id': created.company.pk,
            'site': self.site,
            'actor': self.admin,
            'expected_version': published.pool_state.version,
            'idempotency_token': 'available-archive-0001',
            'reason': 'Remove completed acceptance fixture from circulation',
        }

        archived = archive_reviewed_customer(**values)
        replayed = archive_reviewed_customer(**values)

        self.assertEqual(archived.pool_state.state, 'archived')
        self.assertEqual(archived.pool_state.evidence_status, 'verified')
        self.assertTrue(replayed.replayed)
        self.assertEqual(
            Activity.objects.filter(company=created.company, subject='从客户公海归档').count(),
            1,
        )

    def test_available_record_with_actionable_work_cannot_be_archived(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-create-retire-blocked-0001',
            company_name='Busy Public Pool Company',
            email='busy-pool@example.invalid',
        )
        published = publish_reviewed_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.admin,
            expected_version=created.pool_state.version,
            idempotency_token='review-publish-retire-blocked-0001',
            contact_point_id=created.company.contact_points.get().pk,
            team=self.team,
        )
        Task.objects.create(
            site=self.site,
            company=created.company,
            owner_user=self.sales,
            team=self.team,
            created_by_user=self.sales,
            title='Still actionable',
            status='open',
            due_at=timezone.now(),
        )

        with self.assertRaises(CustomerPoolError) as caught:
            archive_reviewed_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.admin,
                expected_version=published.pool_state.version,
                idempotency_token='available-archive-blocked-0001',
                reason='Must not archive active work',
            )

        self.assertEqual(caught.exception.workflow_code, 'active_tasks_block_archive')
        published.pool_state.refresh_from_db()
        self.assertEqual(published.pool_state.state, 'available')

    def test_release_is_blocked_by_open_tasks_then_succeeds_after_cancel(self):
        created = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-create-test-0004',
            company_name='Task Bound Customer',
            email='task-bound@example.invalid',
            assignment_mode='self',
        )
        task = Task.objects.create(
            site=self.site,
            company=created.company,
            owner_user=self.sales,
            team=self.team,
            created_by_user=self.sales,
            title='Follow up',
            status='open',
            due_at=timezone.now(),
        )
        with self.assertRaises(CustomerPoolError) as caught:
            release_customer(
                company_id=created.company.pk,
                site=self.site,
                actor=self.sales,
                expected_version=created.pool_state.version,
                idempotency_token='release-customer-test-0001',
                reason='Return for reassignment',
            )
        self.assertEqual(caught.exception.workflow_code, 'active_tasks_block_release')

        task.status = 'canceled'
        task.save(update_fields=['status', 'updated_at'])
        released = release_customer(
            company_id=created.company.pk,
            site=self.site,
            actor=self.sales,
            expected_version=created.pool_state.version,
            idempotency_token='release-customer-test-0002',
            reason='Territory changed',
        )
        self.assertEqual(released.pool_state.state, 'available')
        self.assertIsNone(released.company.owner_user_id)
        self.assertEqual(public_pool_queryset(site=self.site, actor=self.sales).count(), 1)

    def test_import_preview_is_read_only_idempotent_and_commit_is_repeatable(self):
        payload = self.import_workbook_bytes()
        preview = preview_customer_import(
            site=self.site,
            actor=self.admin,
            file_bytes=payload,
            original_name='customers.xlsx',
        )
        self.assertFalse(preview.replayed)
        self.assertEqual(preview.batch.counts_json['source_rows'], 1)
        self.assertEqual(preview.batch.counts_json['company_rows'], 1)
        self.assertEqual(preview.batch.counts_json['contact_rows'], 1)
        self.assertEqual(preview.batch.counts_json['new'], 1)
        self.assertEqual(Company.objects.count(), 0)

        replay = preview_customer_import(
            site=self.site,
            actor=self.admin,
            file_bytes=payload,
            original_name='customers.xlsx',
        )
        self.assertTrue(replay.replayed)
        self.assertEqual(replay.batch.pk, preview.batch.pk)
        self.assertEqual(CustomerImportBatch.objects.count(), 1)

        committed = commit_customer_import(batch_id=preview.batch.pk, site=self.site, actor=self.admin)
        replayed_commit = commit_customer_import(batch_id=preview.batch.pk, site=self.site, actor=self.admin)
        self.assertEqual(committed.status, 'succeeded')
        self.assertEqual(replayed_commit.pk, committed.pk)
        self.assertEqual(Company.objects.count(), 1)
        source = CustomerSource.objects.get()
        self.assertEqual(source.intake_method, 'file_import')
        self.assertEqual(source.external_record_id, 'customer-xlsx-v1:company-001')
        self.assertEqual(CompanyContactPoint.objects.count(), 1)
        self.assertEqual(Contact.objects.count(), 0)
        self.assertEqual(CompanyPoolState.objects.get().state, 'review')

    def test_import_preview_quarantines_conflicting_identity_matches(self):
        Company.objects.create(
            site=self.site,
            name='Imported Beverage',
            normalized_name='imported beverage',
            country='Ghana',
            website='https://different.example.org',
        )
        Company.objects.create(
            site=self.site,
            name='Different Beverage',
            normalized_name='different beverage',
            country='Ghana',
            website='https://imported.example.org',
        )

        preview = preview_customer_import(
            site=self.site,
            actor=self.admin,
            file_bytes=self.import_workbook_bytes(),
            original_name='conflicting-identities.xlsx',
        )

        row = preview.batch.rows.get()
        self.assertEqual(row.status, 'quarantined')
        self.assertIn('不会自动合并', row.message)
        self.assertIsNone(row.company_id)

    def test_import_adds_a_new_source_without_overwriting_existing_source_history(self):
        existing = create_manual_customer(
            site=self.site,
            actor=self.sales,
            idempotency_token='manual-before-import-0001',
            company_name='Imported Beverage',
            country='Ghana',
            website='https://imported.example.org',
            email='existing@example.invalid',
            source_type='manual',
            source_detail='Sales declaration',
        )
        preview = preview_customer_import(
            site=self.site,
            actor=self.admin,
            file_bytes=self.import_workbook_bytes(
                source_type='research',
                source_detail='Public company registry',
            ),
            original_name='additive-source.xlsx',
        )
        self.assertEqual(preview.batch.rows.get().status, 'supplement')

        commit_customer_import(batch_id=preview.batch.pk, site=self.site, actor=self.admin)

        self.assertEqual(Company.objects.filter(site=self.site).count(), 1)
        self.assertEqual(
            set(CustomerSource.objects.filter(company=existing.company).values_list('source_type', flat=True)),
            {'manual', 'research'},
        )
        self.assertEqual(existing.company.contact_points.count(), 3)

    def test_partial_import_can_retry_only_the_failed_rows(self):
        preview = preview_customer_import(
            site=self.site,
            actor=self.admin,
            file_bytes=self.import_workbook_bytes(),
            original_name='retry-customers.xlsx',
        )
        with patch(
            'leads.customer_import_services._commit_row',
            side_effect=IntegrityError('synthetic write failure'),
        ):
            partial = commit_customer_import(
                batch_id=preview.batch.pk,
                site=self.site,
                actor=self.admin,
            )
        self.assertEqual(partial.status, 'partial')
        self.assertEqual(partial.rows.get().status, 'failed')

        completed = commit_customer_import(
            batch_id=preview.batch.pk,
            site=self.site,
            actor=self.admin,
        )
        self.assertEqual(completed.status, 'succeeded')
        self.assertEqual(completed.rows.get().status, 'imported')
        self.assertEqual(Company.objects.count(), 1)
