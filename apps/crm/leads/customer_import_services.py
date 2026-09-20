from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook

from console.capabilities import SalesCapability, require_sales

from .customer_pool_services import normalize_company_name, normalize_email, normalize_phone
from .models import (
    Company,
    CompanyContactPoint,
    CompanyPoolState,
    Contact,
    CustomerImportBatch,
    CustomerImportRow,
    CustomerSource,
)


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_RESEARCH_IMPORT_BYTES = 10 * 1024 * 1024
MAX_SOURCE_ROWS = 10000
MAX_ARCHIVE_ENTRIES = 200
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_RESEARCH_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
ADAPTER_VERSION = 'customer-xlsx-v1'
TEMPLATE_VERSION = '版本 1'
COMPANY_HEADERS = ('企业外部键*', '企业名称*', '国家/地区', '城市', '行业', '网站', '数据来源*', '来源详情', '备注')
CONTACT_HEADERS = ('企业外部键*', '联系人姓名', '职位', '联系方式类型*', '联系方式值*', '分机', '用途', '使用状态', '证据说明')
SOURCE_TYPES = frozenset({'manual', 'research', 'website_form', 'meta_native', 'other'})
CHANNELS = frozenset({'email', 'phone', 'whatsapp', 'website', 'other'})


class CustomerImportError(ValidationError):
    def __init__(self, message: str, *, code: str = 'customer_import_invalid'):
        super().__init__(message, code=code)
        self.workflow_code = code


@dataclass(frozen=True, slots=True)
class ImportPreviewResult:
    batch: CustomerImportBatch
    replayed: bool


@dataclass(frozen=True, slots=True)
class ResearchSnapshotProfile:
    code: str
    label: str
    adapter_version: str
    expected_sha256: str
    required_sheets: tuple[str, ...]


RESEARCH_SNAPSHOT_PROFILES = {
    'africa_v345': ResearchSnapshotProfile(
        code='africa_v345',
        label='非洲第二轮 v345（674 家）',
        adapter_version='research-africa-v345',
        expected_sha256='6f3b3ef07a9e4a80051bec0753e64b6069cc3b30f9cf2fe1c6504a34f032a714',
        required_sheets=('第二轮全量联系方式', '第二轮账户证据'),
    ),
    'middle_east_v013': ResearchSnapshotProfile(
        code='middle_east_v013',
        label='中东第二轮 v013（175 家）',
        adapter_version='research-middle-east-v013',
        expected_sha256='f15feb33b2a1d4d66ba8b5dcb6a3561dd0ab8880450745a1d9ca75ec19141543',
        required_sheets=('中东账户总表', '联系方式总表', '全量邮箱', '历史号码纠正'),
    ),
}

AFRICA_CONTACT_HEADERS = (
    '国家', '电话号码', '邮箱', '国际区号', '号码类型', '原始号码', '企业/主体', '账户 ID',
    '公开姓名', '公开职务', '联系用途/首通边界', '当前性', '来源 URL', '核验日期', '制造/商业关系',
)
AFRICA_ACCOUNT_HEADERS = (
    '国家', '企业/主体', '账户 ID', '电话数', '邮箱数', '已证实制造/商业关系', '关键来源 URL',
    '核验日期', '主要证据缺口',
)
MIDDLE_EAST_ACCOUNT_HEADERS = (
    '国家／地区', '企业／实际制造主体', '账户ID', '制造产品／包装', '工厂／地点', '基线电话',
    '基线邮箱', '第二轮状态',
)
MIDDLE_EAST_PHONE_HEADERS = (
    '国家／地区', '电话号码', '邮箱', '企业／实际制造主体', '电话类型', '用途／部门', '姓名',
    '职务', '电话国际区号', '工厂／地点', '电话来源', '邮箱来源', '来源日期', '查阅日期',
    '证据状态', '账户ID', '备注', '原始号码', '分机', '本轮新增号码', '历史备用', '特殊路线',
    '格式有效', '本轮新增企业关联',
)
MIDDLE_EAST_EMAIL_HEADERS = (
    '国家／地区', '电话号码', '邮箱', '企业／实际制造主体', '邮箱类型', '用途／部门', '邮箱来源',
    '来源日期', '查阅日期', '证据状态', '账户ID', '本轮新增邮箱', '特殊路线',
)
MIDDLE_EAST_CORRECTION_HEADERS = (
    '国家／地区', '账户ID', '企业', '原电话号码', '原分类', '纠正分类', '当前处理', '来源URL',
    '查阅日期', '说明',
)


def _text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _row_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return hashlib.sha256(encoded.encode('utf-8')).hexdigest()


def _inspect_xlsx_archive(
    file_bytes: bytes,
    *,
    max_uncompressed_bytes: int = MAX_UNCOMPRESSED_BYTES,
    max_compression_ratio: int = MAX_COMPRESSION_RATIO,
) -> None:
    """Reject active content and decompression abuse before openpyxl reads it."""

    stream = BytesIO(file_bytes)
    if not zipfile.is_zipfile(stream):
        raise CustomerImportError('文件不是有效的 XLSX 压缩包。', code='workbook_invalid')
    stream.seek(0)
    try:
        with zipfile.ZipFile(stream) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise CustomerImportError('XLSX 内部文件数量超过安全限制。', code='archive_too_many_entries')
            total_size = 0
            for entry in entries:
                normalized_name = entry.filename.replace('\\', '/')
                path = PurePosixPath(normalized_name)
                lowered = normalized_name.casefold()
                if path.is_absolute() or '..' in path.parts:
                    raise CustomerImportError('XLSX 包含不安全的内部路径。', code='archive_path_invalid')
                if entry.flag_bits & 0x1:
                    raise CustomerImportError('不接受加密的 XLSX 内容。', code='archive_encrypted')
                total_size += entry.file_size
                if total_size > max_uncompressed_bytes:
                    raise CustomerImportError('XLSX 解压后内容超过该适配器的安全限制。', code='archive_too_large')
                if entry.file_size and (
                    entry.compress_size == 0
                    or entry.file_size / entry.compress_size > max_compression_ratio
                ):
                    raise CustomerImportError('XLSX 压缩比例异常，已拒绝读取。', code='archive_ratio_invalid')
                if (
                    lowered in {'xl/vbaproject.bin', 'xl/connections.xml'}
                    or lowered.startswith(('xl/externallinks/', 'xl/embeddings/', 'xl/activex/', 'xl/macrosheets/', 'customui/'))
                ):
                    raise CustomerImportError('XLSX 包含宏、外部连接或嵌入对象，已拒绝读取。', code='active_content_rejected')
            for entry in entries:
                if entry.filename.casefold().endswith('.rels'):
                    relationship_xml = archive.read(entry)
                    if b'TargetMode="External"' in relationship_xml or b"TargetMode='External'" in relationship_xml:
                        raise CustomerImportError('XLSX 包含外部链接，已拒绝读取。', code='external_link_rejected')
    except (zipfile.BadZipFile, OSError) as error:
        raise CustomerImportError('无法读取该 XLSX 文件，请使用最新模板。', code='workbook_invalid') from error


def _read_sheet(sheet, expected_headers: tuple[str, ...]) -> list[dict[str, str]]:
    return _read_sheet_at(sheet, expected_headers, header_row=1)


def _read_sheet_at(sheet, expected_headers: tuple[str, ...], *, header_row: int) -> list[dict[str, str]]:
    rows = sheet.iter_rows(min_row=header_row, values_only=False)
    try:
        actual_headers = [_text(cell.value) for cell in next(rows)]
    except StopIteration as error:
        raise CustomerImportError(f'工作表“{sheet.title}”为空。', code='sheet_empty') from error
    while actual_headers and not actual_headers[-1]:
        actual_headers.pop()
    if tuple(actual_headers) != expected_headers:
        raise CustomerImportError(
            f'工作表“{sheet.title}”字段与模板版本不一致，请重新下载模板。',
            code='headers_mismatch',
        )
    result = []
    for row_number, cells in enumerate(rows, start=header_row + 1):
        for cell in cells[:len(expected_headers)]:
            if cell.data_type == 'f':
                raise CustomerImportError(
                    f'工作表“{sheet.title}”第 {row_number} 行包含公式，已拒绝读取。',
                    code='formula_rejected',
                )
        normalized = [_text(cell.value) for cell in cells[:len(expected_headers)]]
        normalized += [''] * (len(expected_headers) - len(normalized))
        if not any(normalized):
            continue
        result.append({'__row_number': row_number, **dict(zip(expected_headers, normalized, strict=True))})
        if len(result) > MAX_SOURCE_ROWS:
            raise CustomerImportError(
                f'单个工作表最多允许 {MAX_SOURCE_ROWS} 条非空数据。',
                code='too_many_rows',
            )
    return result


def _contact_payload(row: dict[str, str]) -> dict[str, str]:
    return {
        'external_key': row['企业外部键*'],
        'contact_name': row['联系人姓名'],
        'job_title': row['职位'],
        'channel': row['联系方式类型*'].casefold(),
        'value': row['联系方式值*'],
        'extension': row['分机'],
        'purpose': row['用途'] or 'business',
        'usage_status': row['使用状态'] or 'unknown',
        'evidence_note': row['证据说明'],
        'row_number': row['__row_number'],
    }


def _company_payload(row: dict[str, str], contacts: list[dict[str, str]]) -> dict[str, Any]:
    external_key = row['企业外部键*']
    source_type = row['数据来源*'].casefold()
    return {
        'external_key': external_key,
        'company_name': row['企业名称*'],
        'country': row['国家/地区'],
        'city': row['城市'],
        'industry': row['行业'],
        'website': row['网站'],
        'source_type': source_type,
        'source_detail': row['来源详情'],
        'source_external_id': f'{ADAPTER_VERSION}:{external_key}' if external_key else '',
        'source_evidence': {'adapter_version': ADAPTER_VERSION},
        'notes': row['备注'],
        'company_row_number': row['__row_number'],
        'contacts': contacts,
    }


def _classify_payload(*, site, payload: dict[str, Any]) -> tuple[str, str, Company | None]:
    if not payload['external_key'] or not payload['company_name']:
        return 'rejected', '企业外部键和企业名称均为必填项。', None
    if payload['source_type'] not in SOURCE_TYPES:
        return 'rejected', '数据来源值不在模板允许范围内。', None
    if payload['external_key'].casefold().startswith('example-') or 'example.invalid' in payload['website'].casefold():
        return 'quarantined', '模板示例行不会被导入。', None
    for contact in payload['contacts']:
        if not contact['external_key'] or contact['channel'] not in CHANNELS or not contact['value']:
            return 'rejected', f'联系方式表第 {contact["row_number"]} 行缺少必填项或类型无效。', None
    if not payload['contacts']:
        return 'quarantined', '该企业在本批次中没有联系方式，保留待人工核验。', None
    source_candidates: list[Company] = []
    if payload.get('source_external_id'):
        source_candidates = list(
            Company.objects.filter(
                site=site,
                customer_sources__source_type=payload['source_type'],
                customer_sources__external_record_id=payload['source_external_id'],
            )
            .distinct()
            .order_by('pk')[:2]
        )
    normalized_name = normalize_company_name(payload['company_name'])
    identity_query = Q(normalized_name=normalized_name, country__iexact=payload['country'])
    if payload['website']:
        identity_query |= Q(website__iexact=payload['website'])
    identity_candidates = list(
        Company.objects.filter(site=site)
        .filter(identity_query)
        .order_by('pk')[:2]
    )
    candidates_by_id = {
        company.pk: company for company in [*source_candidates, *identity_candidates]
    }
    candidates = [candidates_by_id[key] for key in sorted(candidates_by_id)[:2]]
    if len(candidates) > 1:
        return 'quarantined', '企业名称/国家和网站分别命中多个现有企业，需人工核验，不会自动合并。', None
    if not candidates:
        return 'new', '将创建新的企业及来源记录。', None
    existing = candidates[0]
    missing_values = any(
        new_value and not getattr(existing, field)
        for field, new_value in (
            ('website', payload['website']),
            ('industry', payload['industry']),
            ('city', payload['city']),
        )
    )
    normalized_contacts = {
        (point.channel, point.normalized_value, point.extension)
        for point in CompanyContactPoint.objects.filter(company=existing)
    }
    has_new_contact = any(
        (
            contact['channel'],
            normalize_email(contact['value']) if contact['channel'] == 'email' else normalize_phone(contact['value']) if contact['channel'] in {'phone', 'whatsapp'} else contact['value'].casefold(),
            contact['extension'],
        ) not in normalized_contacts
        for contact in payload['contacts']
    )
    has_matching_source = CustomerSource.objects.filter(
        company=existing,
        source_type=payload['source_type'],
        source_detail=payload['source_detail'],
    ).exists()
    if missing_values or has_new_contact or not has_matching_source:
        return 'supplement', '匹配到现有企业；仅补充空字段、来源和新联系方式。', existing
    return 'unchanged', '匹配到现有企业，当前文件没有可安全补充的新信息。', existing


@transaction.atomic
def _persist_preview_payloads(
    *,
    site,
    actor,
    file_hash: str,
    original_name: str,
    namespace: str,
    source_type: str,
    adapter_version: str,
    payloads: list[dict[str, Any]],
    contact_row_count: int,
    rejected_rows: list[dict[str, Any]] | None = None,
) -> ImportPreviewResult:
    # Serialize preview publication per site so concurrent parses of the same
    # file cannot race the unique batch key. Parsing remains outside the lock.
    type(site).objects.select_for_update().only('pk').get(pk=site.pk)
    existing = CustomerImportBatch.objects.filter(
        site=site,
        namespace=namespace,
        file_sha256=file_hash,
        adapter_version=adapter_version,
    ).first()
    if existing is not None:
        return ImportPreviewResult(batch=existing, replayed=True)
    batch = CustomerImportBatch.objects.create(
        site=site,
        created_by=actor,
        namespace=namespace,
        source_type=source_type,
        adapter_version=adapter_version,
        file_sha256=file_hash,
        original_name=str(original_name or '')[:500],
        status='preview',
    )
    rejected_rows = rejected_rows or []
    counts = {
        key: 0
        for key in (
            'source_rows', 'company_rows', 'contact_rows', 'new', 'supplement', 'unchanged',
            'quarantined', 'rejected',
        )
    }
    counts['source_rows'] = len(payloads) + len(rejected_rows)
    counts['company_rows'] = len(payloads)
    counts['contact_rows'] = contact_row_count
    seen_company_keys: set[str] = set()
    for payload in payloads:
        company_key = payload['external_key']
        row_number = int(payload.get('company_row_number') or 1)
        if company_key and company_key in seen_company_keys:
            status, message, company = 'rejected', '同一文件中的企业外部键重复。', None
        else:
            status, message, company = _classify_payload(site=site, payload=payload)
        if company_key:
            seen_company_keys.add(company_key)
        CustomerImportRow.objects.create(
            batch=batch,
            row_key=_row_hash({'source_row': row_number, 'payload': payload}),
            row_number=row_number,
            payload_json=payload,
            status=status,
            company=company,
            message=message,
        )
        counts[status] += 1
    for rejected in rejected_rows:
        payload = dict(rejected.get('payload') or {})
        row_number = int(rejected.get('row_number') or 1)
        CustomerImportRow.objects.create(
            batch=batch,
            row_key=_row_hash({'source_row': row_number, 'payload': payload}),
            row_number=row_number,
            payload_json=payload,
            status='rejected',
            message=str(rejected.get('message') or '源记录无法关联到企业。'),
        )
        counts['rejected'] += 1
    terminal_total = sum(
        counts[key] for key in ('new', 'supplement', 'unchanged', 'quarantined', 'rejected')
    )
    if terminal_total != counts['source_rows']:
        raise CustomerImportError('导入预览行数核算失败，未写入业务数据。', code='terminal_accounting_failed')
    batch.counts_json = counts
    batch.save(update_fields=['counts_json', 'updated_at'])
    return ImportPreviewResult(batch=batch, replayed=False)


def _research_usage_status(*values: str) -> str:
    combined = ' '.join(_text(value).casefold() for value in values)
    restricted_markers = (
        '排除', '删除', '错归属', '不使用', '勿用', '传真', '异常', '无效', '停用', '历史',
        'restricted', 'invalid', 'exclude', 'fax',
    )
    return 'restricted' if any(marker in combined for marker in restricted_markers) else 'unknown'


def _split_emails(value: Any) -> list[str]:
    return [item.strip() for item in re.split(r'[;,；\n]+', _text(value)) if item.strip()]


def _research_contact(
    *,
    external_key: str,
    row_number: int,
    channel: str,
    value: str,
    extension: str = '',
    usage_status: str = 'unknown',
    evidence: dict[str, Any],
) -> dict[str, Any]:
    evidence_note = '；'.join(
        f'{key}={_text(item)}' for key, item in evidence.items() if _text(item)
    )
    return {
        'external_key': external_key,
        'contact_name': '',
        'job_title': '',
        'channel': channel,
        'value': value,
        'extension': extension,
        'purpose': 'unknown',
        'usage_status': usage_status,
        'evidence_note': evidence_note,
        'evidence_json': evidence,
        'row_number': row_number,
    }


def _africa_research_payloads(workbook, profile: ResearchSnapshotProfile) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    accounts = _read_sheet_at(
        workbook['第二轮账户证据'], AFRICA_ACCOUNT_HEADERS, header_row=4
    )
    routes = _read_sheet_at(
        workbook['第二轮全量联系方式'], AFRICA_CONTACT_HEADERS, header_row=4
    )
    contacts_by_account: dict[str, list[dict[str, Any]]] = {}
    rejected_rows: list[dict[str, Any]] = []
    for route in routes:
        account_id = route['账户 ID']
        external_key = f'{profile.code}:{account_id}' if account_id else ''
        phone = route['原始号码'] or route['电话号码']
        evidence = {
            'snapshot': profile.code,
            'source_sheet': '第二轮全量联系方式',
            'source_row': route['__row_number'],
            'phone_type': route['号码类型'],
            'public_name': route['公开姓名'],
            'public_role': route['公开职务'],
            'routing_boundary': route['联系用途/首通边界'],
            'currency': route['当前性'],
            'source_url': route['来源 URL'],
            'verified_on': route['核验日期'],
        }
        produced: list[dict[str, Any]] = []
        if phone:
            produced.append(
                _research_contact(
                    external_key=external_key,
                    row_number=route['__row_number'],
                    channel='phone',
                    value=phone,
                    usage_status=_research_usage_status(
                        route['号码类型'], route['当前性'], route['联系用途/首通边界']
                    ),
                    evidence=evidence,
                )
            )
        else:
            for email in _split_emails(route['邮箱']):
                produced.append(
                    _research_contact(
                        external_key=external_key,
                        row_number=route['__row_number'],
                        channel='email',
                        value=email,
                        usage_status=_research_usage_status(
                            route['当前性'], route['联系用途/首通边界']
                        ),
                        evidence=evidence,
                    )
                )
        if not account_id:
            rejected_rows.append({
                'row_number': route['__row_number'],
                'payload': {'orphan_research_route': produced or evidence},
                'message': '研究联系方式缺少账户 ID，不能自动关联企业。',
            })
            continue
        contacts_by_account.setdefault(account_id, []).extend(produced)

    payloads = []
    known_accounts = set()
    for account in accounts:
        account_id = account['账户 ID']
        known_accounts.add(account_id)
        external_key = f'{profile.code}:{account_id}' if account_id else ''
        payloads.append({
            'external_key': external_key,
            'company_name': account['企业/主体'],
            'country': account['国家'],
            'city': '',
            'industry': account['已证实制造/商业关系'],
            'website': '',
            'source_type': 'research',
            'source_detail': f'{profile.label} / {account_id}',
            'source_external_id': f'research:{profile.code}:{account_id}' if account_id else '',
            'source_evidence': {
                'snapshot': profile.code,
                'expected_sha256': profile.expected_sha256,
                'account_row': account['__row_number'],
                'source_url': account['关键来源 URL'],
                'verified_on': account['核验日期'],
                'evidence_status': 'file_declared_unverified',
            },
            'notes': account['主要证据缺口'],
            'company_row_number': account['__row_number'],
            'contacts': contacts_by_account.get(account_id, []),
        })
    for account_id, contacts in contacts_by_account.items():
        if account_id and account_id not in known_accounts:
            rejected_rows.append({
                'row_number': min(contact['row_number'] for contact in contacts),
                'payload': {'orphan_account_id': account_id, 'contacts': contacts},
                'message': '研究联系方式的账户 ID 不在账户证据表中。',
            })
    return payloads, rejected_rows, len(routes)


def _middle_east_research_payloads(workbook, profile: ResearchSnapshotProfile) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    accounts = _read_sheet_at(
        workbook['中东账户总表'], MIDDLE_EAST_ACCOUNT_HEADERS, header_row=1
    )
    phone_rows = _read_sheet_at(
        workbook['联系方式总表'], MIDDLE_EAST_PHONE_HEADERS, header_row=1
    )
    email_rows = _read_sheet_at(
        workbook['全量邮箱'], MIDDLE_EAST_EMAIL_HEADERS, header_row=1
    )
    correction_rows = _read_sheet_at(
        workbook['历史号码纠正'], MIDDLE_EAST_CORRECTION_HEADERS, header_row=1
    )
    corrections = {
        (row['账户ID'], normalize_phone(row['原电话号码'])): row
        for row in correction_rows
        if row['账户ID'] and row['原电话号码']
    }
    contacts_by_account: dict[str, list[dict[str, Any]]] = {}
    rejected_rows: list[dict[str, Any]] = []
    for route in phone_rows:
        account_id = route['账户ID']
        external_key = f'{profile.code}:{account_id}' if account_id else ''
        phone = route['原始号码'] or route['电话号码']
        correction = corrections.get((account_id, normalize_phone(phone)))
        correction_text = ' '.join(
            _text(correction.get(key))
            for key in ('纠正分类', '当前处理', '说明')
        ) if correction else ''
        evidence = {
            'snapshot': profile.code,
            'source_sheet': '联系方式总表',
            'source_row': route['__row_number'],
            'phone_type': route['电话类型'],
            'public_name': route['姓名'],
            'public_role': route['职务'],
            'routing_boundary': route['用途／部门'],
            'source_url': route['电话来源'],
            'evidence_status': route['证据状态'],
            'correction': correction_text,
        }
        contact = _research_contact(
            external_key=external_key,
            row_number=route['__row_number'],
            channel='phone',
            value=phone,
            extension=route['分机'],
            usage_status=_research_usage_status(
                route['电话类型'], route['历史备用'], route['特殊路线'],
                route['格式有效'], correction_text,
            ),
            evidence=evidence,
        ) if phone else None
        if not account_id:
            rejected_rows.append({
                'row_number': route['__row_number'],
                'payload': {'orphan_research_route': contact or evidence},
                'message': '研究电话缺少账户 ID，不能自动关联企业。',
            })
        elif contact:
            contacts_by_account.setdefault(account_id, []).append(contact)
    for route in email_rows:
        account_id = route['账户ID']
        external_key = f'{profile.code}:{account_id}' if account_id else ''
        for email in _split_emails(route['邮箱']):
            contact = _research_contact(
                external_key=external_key,
                row_number=route['__row_number'],
                channel='email',
                value=email,
                usage_status=_research_usage_status(
                    route['邮箱类型'], route['证据状态'], route['特殊路线']
                ),
                evidence={
                    'snapshot': profile.code,
                    'source_sheet': '全量邮箱',
                    'source_row': route['__row_number'],
                    'email_type': route['邮箱类型'],
                    'routing_boundary': route['用途／部门'],
                    'source_url': route['邮箱来源'],
                    'evidence_status': route['证据状态'],
                },
            )
            if not account_id:
                rejected_rows.append({
                    'row_number': route['__row_number'],
                    'payload': {'orphan_research_route': contact},
                    'message': '研究邮箱缺少账户 ID，不能自动关联企业。',
                })
            else:
                contacts_by_account.setdefault(account_id, []).append(contact)

    payloads = []
    known_accounts = set()
    for account in accounts:
        account_id = account['账户ID']
        known_accounts.add(account_id)
        external_key = f'{profile.code}:{account_id}' if account_id else ''
        payloads.append({
            'external_key': external_key,
            'company_name': account['企业／实际制造主体'],
            'country': account['国家／地区'],
            'city': account['工厂／地点'],
            'industry': account['制造产品／包装'],
            'website': '',
            'source_type': 'research',
            'source_detail': f'{profile.label} / {account_id}',
            'source_external_id': f'research:{profile.code}:{account_id}' if account_id else '',
            'source_evidence': {
                'snapshot': profile.code,
                'expected_sha256': profile.expected_sha256,
                'account_row': account['__row_number'],
                'evidence_status': 'file_declared_unverified',
            },
            'notes': account['第二轮状态'],
            'company_row_number': account['__row_number'],
            'contacts': contacts_by_account.get(account_id, []),
        })
    for account_id, contacts in contacts_by_account.items():
        if account_id and account_id not in known_accounts:
            rejected_rows.append({
                'row_number': min(contact['row_number'] for contact in contacts),
                'payload': {'orphan_account_id': account_id, 'contacts': contacts},
                'message': '研究联系方式的账户 ID 不在账户总表中。',
            })
    return payloads, rejected_rows, len(phone_rows) + len(email_rows)


@transaction.atomic
def preview_customer_import(*, site, actor, file_bytes: bytes, original_name: str) -> ImportPreviewResult:
    require_sales(actor, SalesCapability.POOL_IMPORT)
    if not file_bytes:
        raise CustomerImportError('请选择要导入的 XLSX 文件。', code='file_required')
    if len(file_bytes) > MAX_IMPORT_BYTES:
        raise CustomerImportError('导入文件不能超过 5 MB。', code='file_too_large')
    if not str(original_name or '').lower().endswith('.xlsx'):
        raise CustomerImportError('只接受 XLSX 模板文件。', code='file_type_invalid')
    _inspect_xlsx_archive(file_bytes)
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    existing_batch = CustomerImportBatch.objects.filter(
        site=site,
        namespace='customer',
        file_sha256=file_hash,
        adapter_version=ADAPTER_VERSION,
    ).first()
    if existing_batch is not None:
        return ImportPreviewResult(batch=existing_batch, replayed=True)
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False, keep_links=False)
    except Exception as error:
        raise CustomerImportError('无法读取该 XLSX 文件，请使用最新模板。', code='workbook_invalid') from error
    required_sheets = ['填写说明', '企业', '联系方式']
    if workbook.sheetnames != required_sheets:
        raise CustomerImportError('工作簿必须且只能包含“填写说明”“企业”“联系方式”三个工作表。', code='sheets_invalid')
    if _text(workbook['填写说明']['B1'].value) != TEMPLATE_VERSION:
        raise CustomerImportError('模板版本无法识别，请重新下载最新模板。', code='template_version_invalid')
    company_rows = _read_sheet(workbook['企业'], COMPANY_HEADERS)
    contact_rows = _read_sheet(workbook['联系方式'], CONTACT_HEADERS)
    contacts_by_company: dict[str, list[dict[str, str]]] = {}
    for row in contact_rows:
        contact = _contact_payload(row)
        contacts_by_company.setdefault(contact['external_key'], []).append(contact)
    known_company_keys = {row['企业外部键*'] for row in company_rows}
    orphan_contacts = [row for row in contact_rows if row['企业外部键*'] not in known_company_keys]
    payloads = [
        _company_payload(row, contacts_by_company.get(row['企业外部键*'], []))
        for row in company_rows
    ]
    rejected_rows = [
        {
            'row_number': orphan['__row_number'],
            'payload': {'orphan_contact': _contact_payload(orphan)},
            'message': '联系方式引用的企业外部键不存在。',
        }
        for orphan in orphan_contacts
    ]
    return _persist_preview_payloads(
        site=site,
        actor=actor,
        file_hash=file_hash,
        original_name=original_name,
        namespace='customer',
        source_type='other',
        adapter_version=ADAPTER_VERSION,
        payloads=payloads,
        contact_row_count=len(contact_rows),
        rejected_rows=rejected_rows,
    )


@transaction.atomic
def preview_research_snapshot_import(
    *,
    site,
    actor,
    file_bytes: bytes,
    original_name: str,
    profile_code: str,
) -> ImportPreviewResult:
    require_sales(actor, SalesCapability.POOL_IMPORT)
    profile = RESEARCH_SNAPSHOT_PROFILES.get(str(profile_code or '').strip())
    if profile is None:
        raise CustomerImportError('请选择受支持的研究快照版本。', code='research_profile_invalid')
    if not file_bytes:
        raise CustomerImportError('请选择要核验的研究快照 XLSX。', code='file_required')
    if len(file_bytes) > MAX_RESEARCH_IMPORT_BYTES:
        raise CustomerImportError('研究快照不能超过 10 MB。', code='file_too_large')
    if not str(original_name or '').lower().endswith('.xlsx'):
        raise CustomerImportError('研究快照必须是 XLSX 文件。', code='file_type_invalid')
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    if file_hash != profile.expected_sha256:
        raise CustomerImportError(
            '研究快照哈希与锁定版本不一致；请勿导入同名改写文件。',
            code='research_snapshot_hash_mismatch',
        )
    existing_batch = CustomerImportBatch.objects.filter(
        site=site,
        namespace='research_snapshot',
        file_sha256=file_hash,
        adapter_version=profile.adapter_version,
    ).first()
    if existing_batch is not None:
        return ImportPreviewResult(batch=existing_batch, replayed=True)
    _inspect_xlsx_archive(
        file_bytes,
        max_uncompressed_bytes=MAX_RESEARCH_UNCOMPRESSED_BYTES,
        max_compression_ratio=50,
    )
    try:
        workbook = load_workbook(
            BytesIO(file_bytes), read_only=True, data_only=False, keep_links=False
        )
    except Exception as error:
        raise CustomerImportError('无法读取锁定研究快照。', code='workbook_invalid') from error
    missing_sheets = [sheet for sheet in profile.required_sheets if sheet not in workbook.sheetnames]
    if missing_sheets:
        raise CustomerImportError(
            f'研究快照缺少工作表：{"、".join(missing_sheets)}。',
            code='research_sheets_missing',
        )
    try:
        if profile.code == 'africa_v345':
            payloads, rejected_rows, contact_row_count = _africa_research_payloads(
                workbook, profile
            )
        else:
            payloads, rejected_rows, contact_row_count = _middle_east_research_payloads(
                workbook, profile
            )
    finally:
        workbook.close()
    return _persist_preview_payloads(
        site=site,
        actor=actor,
        file_hash=file_hash,
        original_name=original_name,
        namespace='research_snapshot',
        source_type='research',
        adapter_version=profile.adapter_version,
        payloads=payloads,
        contact_row_count=contact_row_count,
        rejected_rows=rejected_rows,
    )


def _normalized_contact_value(channel: str, value: str) -> str:
    if channel == 'email':
        return normalize_email(value)
    if channel in {'phone', 'whatsapp'}:
        return normalize_phone(value)
    return value.strip().casefold()


def _import_contact(*, site, company: Company, payload: dict[str, str]) -> None:
    contact = None
    if payload['contact_name']:
        contact, _created = Contact.objects.get_or_create(
            site=site,
            company=company,
            full_name=payload['contact_name'],
            defaults={
                'owner_user': company.owner_user,
                'team': company.team,
                'job_title': payload['job_title'],
                'status': 'active',
                'country': company.country,
            },
        )
    CompanyContactPoint.objects.get_or_create(
        site=site,
        company=company,
        channel=payload['channel'],
        normalized_value=_normalized_contact_value(payload['channel'], payload['value']),
        extension=payload['extension'],
        defaults={
            'contact': contact,
            'raw_value': payload['value'],
            'purpose': payload['purpose'] if payload['purpose'] in {'business', 'personal', 'unknown'} else 'unknown',
            'status': 'unverified',
            'usage_status': payload['usage_status'] if payload['usage_status'] in {'permitted', 'restricted', 'unknown'} else 'unknown',
            'evidence_json': {
                **dict(payload.get('evidence_json') or {}),
                'note': payload['evidence_note'],
                'imported': True,
            },
        },
    )


def _commit_row(*, row: CustomerImportRow, batch: CustomerImportBatch, site, actor) -> None:
    payload = row.payload_json
    company = row.company
    if company is None:
        company = Company.objects.create(
            site=site,
            name=payload['company_name'],
            normalized_name=normalize_company_name(payload['company_name']),
            country=payload['country'],
            city=payload['city'],
            industry=payload['industry'],
            website=payload['website'],
            status='prospect',
            source_channel=payload['source_type'],
            notes=payload['notes'],
        )
        CompanyPoolState.objects.create(company=company, state='review', evidence_status='pending')
    else:
        changed = []
        for field in ('website', 'industry', 'city'):
            if not getattr(company, field) and payload[field]:
                setattr(company, field, payload[field])
                changed.append(field)
        if changed:
            company.save(update_fields=[*changed, 'updated_at'])
        CompanyPoolState.objects.get_or_create(
            company=company,
            defaults={
                'state': 'owned' if company.owner_user_id else 'review',
                'evidence_status': 'declared' if company.owner_user_id else 'pending',
                'claimed_at': timezone.now() if company.owner_user_id else None,
            },
        )
    CustomerSource.objects.get_or_create(
        site=site,
        company=company,
        source_type=payload['source_type'],
        external_record_id=(
            payload.get('source_external_id') or f'import:{batch.pk}:{row.row_key}'
        ),
        defaults={
            'intake_method': 'file_import',
            'source_detail': payload['source_detail'],
            'responsible_user': actor,
            'responsible_team': company.team,
            'evidence_status': 'declared',
            'evidence_json': {
                **dict(payload.get('source_evidence') or {}),
                'batch_id': batch.pk,
                'row_number': row.row_number,
                'adapter_version': batch.adapter_version,
                'file_sha256': batch.file_sha256,
            },
        },
    )
    for contact_payload in payload['contacts']:
        _import_contact(site=site, company=company, payload=contact_payload)
    if payload.get('kind') == 'standard21':
        from .customer_standard21_import import commit_standard21_rows

        commit_standard21_rows(row=row, batch=batch, site=site, actor=actor, company=company)
    row.company = company
    row.status = 'imported'
    row.result_json = {'company_id': company.pk}
    row.message = '已写入客户主数据；未创建网站表单收据或 Meta 回传事件。'
    row.save(update_fields=['company', 'status', 'result_json', 'message', 'updated_at'])


@transaction.atomic
def commit_customer_import(*, batch_id: int, site, actor) -> CustomerImportBatch:
    require_sales(actor, SalesCapability.POOL_IMPORT)
    batch = CustomerImportBatch.objects.select_for_update().get(pk=batch_id, site=site)
    if batch.status == 'succeeded':
        return batch
    if batch.status not in {'preview', 'partial'}:
        raise CustomerImportError('该导入批次当前不可执行。', code='batch_not_ready')
    batch.status = 'processing'
    batch.save(update_fields=['status', 'updated_at'])
    counts = dict(batch.counts_json)
    imported = int(counts.get('imported') or 0)
    failed = 0
    for row in batch.rows.select_for_update().order_by('row_number', 'id'):
        if row.status not in {'new', 'supplement', 'failed'}:
            continue
        try:
            with transaction.atomic():
                _commit_row(row=row, batch=batch, site=site, actor=actor)
            imported += 1
        except Exception as error:
            row.status = 'failed'
            row.message = f'写入失败：{type(error).__name__}'
            row.save(update_fields=['status', 'message', 'updated_at'])
            failed += 1
    counts['imported'] = imported
    counts['failed'] = failed
    batch.counts_json = counts
    batch.status = 'partial' if failed else 'succeeded'
    batch.completed_at = timezone.now()
    batch.save(update_fields=['counts_json', 'status', 'completed_at', 'updated_at'])
    return batch
