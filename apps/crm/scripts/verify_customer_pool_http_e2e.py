"""Exercise the customer-pool file workflow through the live local HTTP server.

This verifier is intentionally restricted to a loopback server and a PostgreSQL
database whose name ends in ``_test``.  It creates only synthetic records.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import secrets
import subprocess
import sys
import uuid
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOCAL_HOSTS = {'127.0.0.1', 'localhost', '::1'}
XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument('--base-url', default='http://127.0.0.1:26121/')
    parser.add_argument('--dsn', required=True)
    return parser.parse_args()


def require_safe_local_targets(base_url: str, dsn: str) -> None:
    http_target = urlparse(base_url)
    database_target = urlparse(dsn)
    if http_target.scheme != 'http' or http_target.hostname not in LOCAL_HOSTS:
        raise SystemExit('Refusing to exercise a non-loopback HTTP server.')
    if database_target.scheme not in {'postgres', 'postgresql'}:
        raise SystemExit('The verification database must use PostgreSQL.')
    if database_target.hostname not in LOCAL_HOSTS:
        raise SystemExit('Refusing to inspect a non-local PostgreSQL host.')
    if not database_target.path.lstrip('/').endswith('_test'):
        raise SystemExit('The verification database name must end in _test.')


def csrf_token(html: str) -> str:
    match = re.search(
        r'name=["\']csrfmiddlewaretoken["\'][^>]*value=["\']([^"\']+)',
        html,
        flags=re.IGNORECASE,
    )
    if match is None:
        raise AssertionError('The expected CSRF form token was not rendered.')
    return match.group(1)


def decoded(response: requests.Response) -> str:
    response.encoding = 'utf-8'
    return response.text


def login(base_url: str, *, username: str, password: str) -> requests.Session:
    session = requests.Session()
    login_url = urljoin(base_url, 'admin/login/')
    page = session.get(login_url, timeout=10)
    page.raise_for_status()
    response = session.post(
        login_url,
        data={
            'username': username,
            'password': password,
            'csrfmiddlewaretoken': csrf_token(decoded(page)),
            'next': '/admin/',
        },
        headers={'Referer': login_url},
        timeout=10,
        allow_redirects=True,
    )
    response.raise_for_status()
    if '/admin/login/' in response.url or 'sessionid' not in session.cookies:
        raise AssertionError(f'Local candidate login failed for {username}.')
    return session


def build_synthetic_import(template_bytes: bytes, run_key: str) -> tuple[bytes, dict[str, str]]:
    workbook = load_workbook(BytesIO(template_bytes))
    assert workbook.sheetnames == ['填写说明', '企业', '联系方式']
    companies = workbook['企业']
    contacts = workbook['联系方式']
    if companies.max_row > 1:
        companies.delete_rows(2, companies.max_row - 1)
    if contacts.max_row > 1:
        contacts.delete_rows(2, contacts.max_row - 1)

    external_key = f'qa-http-{run_key}'
    company_name = f'HTTP Import QA Beverage {run_key}'
    website = f'https://{external_key}.test'
    email = f'{external_key}@example.test'
    phone = '+1 202 555 0100'
    companies.append([
        external_key,
        company_name,
        'Ghana',
        'Accra',
        'Beverage manufacturing',
        website,
        'research',
        'Local synthetic HTTP verification',
        'Synthetic record; never contact',
    ])
    contacts.append([
        external_key,
        'Synthetic QA Contact',
        'Verification only',
        'email',
        email,
        '',
        'business',
        'unknown',
        'Generated locally for the HTTP verifier',
    ])
    contacts.append([
        external_key,
        '',
        '',
        'phone',
        phone,
        '',
        'business',
        'unknown',
        'Reserved fictional 555 number',
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), {
        'external_key': external_key,
        'company_name': company_name,
        'website': website,
        'email': email,
        'phone': phone,
    }


def main() -> int:
    args = parse_args()
    require_safe_local_targets(args.base_url, args.dsn)
    password = os.environ.get('SITEOS_CANDIDATE_PASSWORD', '')

    sys.path.insert(0, str(PROJECT_ROOT))
    os.environ['DJANGO_SETTINGS_MODULE'] = 'siteos_admin.settings'
    os.environ['SITEOS_ADMIN_DEBUG'] = '1'
    os.environ['SITEOS_ADMIN_DATABASE_URL'] = args.dsn
    os.environ['SITEOS_ADMIN_DATABASE_SSLMODE'] = 'disable'
    os.environ['SITEOS_ADMIN_SECURE_SSL_REDIRECT'] = '0'

    import django

    django.setup()

    from django.conf import settings
    from django.contrib.auth import get_user_model

    from leads.models import (
        Company,
        CompanyContactPoint,
        CompanyPoolState,
        Contact,
        CustomerExportJob,
        CustomerImportBatch,
        CustomerSource,
        LeadEventOutbox,
        LeadSubmission,
        SalesTeamMember,
    )
    from sitecore.models import Site

    admin_user = get_user_model().objects.get(username='local-admin')
    manager_user = get_user_model().objects.get(username='local-manager')
    sales_user = get_user_model().objects.get(username='local-sales')
    if not password:
        password = secrets.token_urlsafe(24)
        for local_user in (admin_user, manager_user, sales_user):
            local_user.set_password(password)
            local_user.save(update_fields=['password'])
    base_url = args.base_url.rstrip('/') + '/'
    admin = login(base_url, username='local-admin', password=password)
    manager = login(base_url, username='local-manager', password=password)
    sales = login(base_url, username='local-sales', password=password)
    site = Site.objects.get(code='siteos_demo')

    import_url = urljoin(base_url, 'admin/sales/customer-pool/import/')
    template_url = urljoin(base_url, 'admin/sales/customer-pool/template.xlsx')
    export_url = urljoin(base_url, 'admin/sales/customer-pool/export/')
    bulk_claim_url = urljoin(base_url, 'admin/sales/customer-pool/bulk-claim/')
    pool_url = urljoin(base_url, 'admin/sales/customer-pool/')
    sales_team = (
        SalesTeamMember.objects.select_related('team')
        .get(user=sales_user, team__enabled=True)
        .team
    )

    template_response = admin.get(template_url, timeout=10)
    template_response.raise_for_status()
    assert template_response.headers['Content-Type'].startswith(XLSX_CONTENT_TYPE)
    template_workbook = load_workbook(BytesIO(template_response.content), read_only=True, data_only=True)
    assert template_workbook.sheetnames == ['填写说明', '企业', '联系方式']
    assert template_workbook['企业']['A1'].value == '企业外部键*'
    assert template_workbook['联系方式']['D1'].value == '联系方式类型*'

    run_key = uuid.uuid4().hex[:10]
    import_bytes, fixture = build_synthetic_import(template_response.content, run_key)
    file_hash = hashlib.sha256(import_bytes).hexdigest()
    original_name = f'customer-pool-http-qa-{run_key}.xlsx'
    before = {
        'companies': Company.objects.filter(site=site).count(),
        'submissions': LeadSubmission.objects.filter(site=site).count(),
        'outbox': LeadEventOutbox.objects.filter(submission__site=site).count(),
        'batches': CustomerImportBatch.objects.filter(site=site).count(),
    }

    import_page = admin.get(import_url, timeout=10)
    import_page.raise_for_status()
    preview = admin.post(
        import_url,
        data={'action': 'preview', 'csrfmiddlewaretoken': csrf_token(decoded(import_page))},
        files={'file': (original_name, import_bytes, XLSX_CONTENT_TYPE)},
        headers={'Referer': import_url},
        timeout=15,
    )
    preview.raise_for_status()
    preview_html = decoded(preview)
    assert '预览完成，尚未写入客户主数据。' in preview_html
    batch = CustomerImportBatch.objects.get(
        site=site,
        created_by=admin_user,
        file_sha256=file_hash,
    )
    assert batch.status == 'preview'
    assert batch.counts_json['source_rows'] == 1
    assert batch.counts_json['new'] == 1
    assert batch.rows.get().status == 'new'
    assert Company.objects.filter(site=site).count() == before['companies']
    assert LeadSubmission.objects.filter(site=site).count() == before['submissions']
    assert LeadEventOutbox.objects.filter(submission__site=site).count() == before['outbox']

    replay = admin.post(
        import_url,
        data={'action': 'preview', 'csrfmiddlewaretoken': csrf_token(preview_html)},
        files={'file': (original_name, import_bytes, XLSX_CONTENT_TYPE)},
        headers={'Referer': import_url},
        timeout=15,
    )
    replay.raise_for_status()
    replay_html = decoded(replay)
    assert '该文件已存在，显示原预览结果。' in replay_html
    assert CustomerImportBatch.objects.filter(site=site).count() == before['batches'] + 1
    assert Company.objects.filter(site=site).count() == before['companies']

    commit = admin.post(
        import_url,
        data={
            'action': 'commit',
            'batch_id': str(batch.pk),
            'csrfmiddlewaretoken': csrf_token(replay_html),
        },
        headers={'Referer': import_url},
        timeout=15,
    )
    commit.raise_for_status()
    assert '导入批次已执行；逐行结果已保留。' in decoded(commit)
    batch.refresh_from_db()
    assert batch.status == 'succeeded'
    assert batch.counts_json['imported'] == 1
    company = Company.objects.get(site=site, normalized_name=fixture['company_name'].casefold())
    state = CompanyPoolState.objects.get(company=company)
    assert state.state == 'review'
    assert company.owner_user_id is None and company.team_id is None
    source = CustomerSource.objects.get(company=company, source_type='research')
    assert source.intake_method == 'file_import'
    assert source.external_record_id == f'customer-xlsx-v1:{fixture["external_key"]}'
    assert CompanyContactPoint.objects.filter(company=company).count() == 2
    assert Contact.objects.filter(company=company, full_name='Synthetic QA Contact').count() == 1
    assert Company.objects.filter(site=site).count() == before['companies'] + 1
    assert LeadSubmission.objects.filter(site=site).count() == before['submissions']
    assert LeadEventOutbox.objects.filter(submission__site=site).count() == before['outbox']

    detail_url = urljoin(base_url, f'admin/sales/customer-pool/{company.pk}/')
    review_url = urljoin(base_url, f'admin/sales/customer-pool/{company.pk}/review/')
    review_page = admin.get(detail_url, timeout=10)
    review_page.raise_for_status()
    assert '核验并发布到公海' in decoded(review_page)
    selected_point = company.contact_points.filter(channel='email').get()
    review = admin.post(
        review_url,
        data={
            'expected_version': str(state.version),
            'idempotency_token': f'http-review-{run_key}',
            'team_id': str(sales_team.pk),
            'contact_point_id': str(selected_point.pk),
            'csrfmiddlewaretoken': csrf_token(decoded(review_page)),
        },
        headers={'Referer': detail_url},
        timeout=15,
        allow_redirects=True,
    )
    review.raise_for_status()
    assert '并发布到客户公海' in decoded(review)
    company.refresh_from_db()
    state.refresh_from_db()
    selected_point.refresh_from_db()
    assert state.state == 'available'
    assert state.evidence_status == 'verified'
    assert company.owner_user_id is None and company.team_id == sales_team.pk
    assert selected_point.status == 'active'
    assert selected_point.usage_status == 'permitted'
    sales_pool = sales.get(pool_url, timeout=10)
    sales_pool.raise_for_status()
    assert fixture['company_name'] in decoded(sales_pool)

    commit_again = admin.post(
        import_url,
        data={
            'action': 'commit',
            'batch_id': str(batch.pk),
            'csrfmiddlewaretoken': csrf_token(decoded(commit)),
        },
        headers={'Referer': import_url},
        timeout=15,
    )
    commit_again.raise_for_status()
    assert Company.objects.filter(site=site, pk=company.pk).count() == 1
    assert CustomerSource.objects.filter(company=company).count() == 1
    assert CompanyContactPoint.objects.filter(company=company).count() == 2

    owned = Company.objects.filter(site=site, owner_user=sales_user).order_by('id').first()
    assert owned is not None
    owned_state = CompanyPoolState.objects.get(company=owned)
    sales_pool = sales.get(urljoin(pool_url, '?scope=available'), timeout=10)
    sales_pool.raise_for_status()
    bulk_token = f'http-bulk-claim-{run_key}'
    bulk_payload = [
        ('items', f'{company.pk}:{state.version}'),
        ('items', f'{owned.pk}:{owned_state.version}'),
        ('idempotency_token', bulk_token),
        ('next', '/admin/sales/customer-pool/?scope=available'),
        ('csrfmiddlewaretoken', csrf_token(decoded(sales_pool))),
    ]
    bulk_claim = sales.post(
        bulk_claim_url,
        data=bulk_payload,
        headers={'Referer': sales_pool.url},
        timeout=15,
        allow_redirects=True,
    )
    bulk_claim.raise_for_status()
    bulk_html = decoded(bulk_claim)
    assert '批量领取完成：1 家成功，1 家未领取。' in bulk_html
    assert '已领取。' in bulk_html
    assert '该客户已被其他销售领取，请刷新后重试。' in bulk_html
    company.refresh_from_db()
    state.refresh_from_db()
    assert company.owner_user_id == sales_user.pk
    assert state.state == 'owned'
    claimed_version = state.version
    assert Contact.objects.filter(company=company, owner_user=sales_user).count() == 1

    replay_payload = [
        ('items', f'{company.pk}:{state.version - 1}'),
        ('items', f'{owned.pk}:{owned_state.version}'),
        ('idempotency_token', bulk_token),
        ('next', '/admin/sales/customer-pool/?scope=available'),
        ('csrfmiddlewaretoken', csrf_token(bulk_html)),
    ]
    bulk_replay = sales.post(
        bulk_claim_url,
        data=replay_payload,
        headers={'Referer': bulk_claim.url},
        timeout=15,
        allow_redirects=True,
    )
    bulk_replay.raise_for_status()
    replay_html = decoded(bulk_replay)
    assert '批量领取完成：1 家成功，1 家未领取。' in replay_html
    assert '该领取请求已处理。' in replay_html
    state.refresh_from_db()
    assert state.version == claimed_version

    assignment_url = urljoin(
        base_url,
        f'admin/sales/customer-pool/{company.pk}/assign/',
    )
    manager_detail = manager.get(detail_url, timeout=10)
    manager_detail.raise_for_status()
    assignment_token = f'http-assignment-{run_key}'
    assignment = manager.post(
        assignment_url,
        data={
            'target_user_id': str(manager_user.pk),
            'expected_version': str(state.version),
            'idempotency_token': assignment_token,
            'reason': 'Synthetic HTTP verification handoff',
            'csrfmiddlewaretoken': csrf_token(decoded(manager_detail)),
        },
        headers={'Referer': detail_url},
        timeout=15,
        allow_redirects=True,
    )
    assignment.raise_for_status()
    assignment_html = decoded(assignment)
    assert f'已分配“{fixture["company_name"]}”给 local-manager。' in assignment_html
    company.refresh_from_db()
    state.refresh_from_db()
    assigned_version = state.version
    assert company.owner_user_id == manager_user.pk
    assert state.state == 'owned'
    assert Contact.objects.filter(company=company, owner_user=manager_user).count() == 1
    assert LeadSubmission.objects.filter(site=site).count() == before['submissions']
    assert LeadEventOutbox.objects.filter(submission__site=site).count() == before['outbox']

    assignment_replay = manager.post(
        assignment_url,
        data={
            'target_user_id': str(manager_user.pk),
            'expected_version': str(claimed_version),
            'idempotency_token': assignment_token,
            'reason': 'Synthetic HTTP verification handoff',
            'csrfmiddlewaretoken': csrf_token(assignment_html),
        },
        headers={'Referer': detail_url},
        timeout=15,
        allow_redirects=True,
    )
    assignment_replay.raise_for_status()
    assert '已确认' in decoded(assignment_replay)
    state.refresh_from_db()
    assert state.version == assigned_version

    export_jobs_before = CustomerExportJob.objects.filter(site=site).count()
    pool_page = sales.get(pool_url, timeout=10)
    pool_page.raise_for_status()
    export_response = sales.post(
        export_url,
        data=[
            ('mode', 'selected'),
            ('company_ids', str(owned.pk)),
            ('fields', 'company'),
            ('fields', 'source'),
            ('format', 'xlsx'),
            ('csrfmiddlewaretoken', csrf_token(decoded(pool_page))),
        ],
        headers={'Referer': pool_url},
        timeout=15,
    )
    export_response.raise_for_status()
    assert export_response.history or export_response.status_code == 302
    export_job = CustomerExportJob.objects.filter(site=site, created_by=sales_user).latest('id')
    assert export_job.status == 'pending'
    assert not export_job.storage_path
    worker = subprocess.run(
        [sys.executable, 'manage.py', 'process_customer_exports', '--limit', '10'],
        cwd=PROJECT_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    assert 'processed 1/1' in worker.stdout
    export_job.refresh_from_db()
    export_download_url = urljoin(
        base_url,
        f'admin/sales/customer-pool/export/{export_job.pk}/download/',
    )
    downloaded = sales.get(export_download_url, timeout=15)
    downloaded.raise_for_status()
    assert downloaded.headers['Content-Type'].startswith(XLSX_CONTENT_TYPE)
    assert downloaded.headers.get('Cache-Control') == 'private, no-store'
    export_workbook = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=True)
    assert export_workbook.sheetnames == ['企业', '来源明细']
    enterprise_rows = list(export_workbook['企业'].iter_rows(values_only=True))
    assert len(enterprise_rows) == 2
    assert enterprise_rows[1][0] == owned.pk
    assert enterprise_rows[1][1] == owned.name
    assert CustomerExportJob.objects.filter(site=site).count() == export_jobs_before + 1
    assert export_job.record_count == 1
    assert export_job.scope_json['mode'] == 'selected'
    assert export_job.fields_json == ['company', 'source']
    assert export_job.format == 'xlsx'
    assert export_job.sha256 == hashlib.sha256(downloaded.content).hexdigest()
    from console.customer_pool_exports import _private_path
    private_path = _private_path(export_job.storage_path)
    assert private_path is not None
    export_root = (Path(settings.BASE_DIR) / '.runtime' / 'customer-exports').resolve()
    assert private_path.is_relative_to(export_root)
    assert private_path.read_bytes() == downloaded.content
    downloaded_again = sales.get(export_download_url, timeout=15)
    downloaded_again.raise_for_status()
    assert downloaded_again.content == downloaded.content

    invalid_format_page = sales.get(pool_url, timeout=10)
    invalid_format_page.raise_for_status()
    invalid_format = sales.post(
        export_url,
        data=[
            ('mode', 'selected'),
            ('company_ids', str(owned.pk)),
            ('fields', 'company'),
            ('format', 'csv'),
            ('csrfmiddlewaretoken', csrf_token(decoded(invalid_format_page))),
        ],
        headers={'Referer': pool_url},
        timeout=15,
        allow_redirects=True,
    )
    invalid_format.raise_for_status()
    assert '当前只支持 XLSX；CSV 无法无损表达一家企业的多条联系方式。' in decoded(invalid_format)
    assert CustomerExportJob.objects.filter(site=site).count() == export_jobs_before + 1

    mixed_page = sales.get(pool_url, timeout=10)
    mixed_page.raise_for_status()
    mixed = sales.post(
        export_url,
        data=[
            ('mode', 'selected'),
            ('company_ids', str(owned.pk)),
            ('company_ids', str(company.pk)),
            ('fields', 'company'),
            ('format', 'xlsx'),
            ('csrfmiddlewaretoken', csrf_token(decoded(mixed_page))),
        ],
        headers={'Referer': pool_url},
        timeout=15,
        allow_redirects=True,
    )
    mixed.raise_for_status()
    assert '所选范围包含不存在或无权导出的客户，已取消整个导出。' in decoded(mixed)
    assert CustomerExportJob.objects.filter(site=site).count() == export_jobs_before + 1

    print('Customer-pool HTTP E2E passed on local PostgreSQL:')
    print('- downloaded and parsed the live three-sheet template')
    print('- preview wrote one batch and zero customer records')
    print('- identical upload reused the existing batch')
    print('- commit created one synthetic review record, one source, two contact points, and no lead/outbox rows')
    print('- admin review selected a team and usable route before publishing to that team pool')
    print('- repeated commit created no duplicates')
    print('- explicit bulk claim preserved one success beside one conflict and replayed safely')
    print('- manager handoff changed the owner and related contact once, with idempotent replay')
    print('- selected export persisted explicit fields/format, then a separate worker produced one private XLSX')
    print('- the generated workbook contained only the requested enterprise and source sheets')
    print('- unsupported CSV was rejected without creating a task')
    print('- authenticated re-download rechecked the task and returned the same verified file')
    print('- mixed authorized/unauthorized selection was rejected without creating a file job')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
